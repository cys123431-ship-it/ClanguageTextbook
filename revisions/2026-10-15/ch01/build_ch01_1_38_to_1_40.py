from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import subprocess, json, zipfile, hashlib, datetime, re

ROOT=Path(__file__).parent
FLAGS=['-std=c11','-Wall','-Wextra','-Wpedantic','-Werror','-O0']
TEAL=RGBColor(0,111,116); GRAY=RGBColor(66,76,80)
SECTIONS=[
 {'number':'1.38','title':'구조체의 탄생 — 서로 관련된 값을 하나의 기록으로 묶기',
  'why':'이름과 점수처럼 함께 이동해야 하는 값들을 서로 다른 변수로 흩어 놓으면, 복사하거나 함수에 전달할 때 일부를 빠뜨리기 쉽다. 구조체(struct)는 서로 다른 형식의 값들을 한 객체의 멤버로 묶는다. 학생 기록, 좌표, 주문 항목, 센서 측정 결과처럼 서로 관련된 정보를 다룰 때 사용한다.',
  'terms':'struct Student { int id; int score; };는 Student라는 구조체 태그와 두 멤버를 정의한다. 이 선언 자체가 학생 객체를 만드는 것은 아니다. struct Student a = {7, 82};에서 a라는 실제 객체가 생성된다. 점 연산자(.)는 객체 안의 특정 멤버를 선택한다. 구조체 멤버는 선언 순서대로 배치되지만 사이에 패딩이 있을 수 있어 바이트 수를 단순 합산하면 안 된다.',
  'goals':['구조체 형식 정의와 구조체 객체 생성을 구별한다.','멤버 접근·초기화·구조체 대입을 값의 변화로 설명한다.','중첩 구조체와 지정 초기화를 C11 문법으로 안전하게 사용한다.']},
 {'number':'1.39','title':'구조체를 활용하는 함수와 배열 — 기록을 안전하게 전달하기',
  'why':'한 학생의 정보를 묶는 데 성공해도 여러 학생을 관리하고 함수에서 수정할 방법이 없으면 실제 프로그램을 만들기 어렵다. 구조체 배열은 같은 종류의 기록을 연속해서 보관하고, 구조체 값 전달과 포인터 전달은 읽기와 수정의 책임을 명확하게 나눈다. 명단 정리·재고 집계·성적 처리에서 사용한다.',
  'terms':'struct Student records[3]는 구조체 객체 세 개를 담는 배열이다. 함수 인수로 구조체 값을 전달하면 멤버 값이 복사되므로 원본이 자동으로 바뀌지 않는다. struct Student *p는 구조체 객체를 가리키는 포인터이며 p->score는 (*p).score의 간단한 표기다. 원본을 수정할 함수에는 유효한 객체 주소를 전달한다. 읽기 전용 함수에는 const struct Student *를 우선 고려한다.',
  'goals':['구조체 배열을 순회·검색한다.','값 전달과 포인터 전달의 원본 변경 여부를 예측한다.','구조체 배열 정렬을 작은 하위 목표로 나누어 구현한다.']},
 {'number':'1.40','title':'열거형·공용체·typedef·비트 필드 — 자료 표현의 선택',
  'why':'숫자 0·1·2만으로 작업 상태를 표현하면 그 의미를 읽기 어렵고, 서로 다른 종류의 값을 같은 저장 공간에 두는 경우에도 해석 규칙이 필요하다. enum은 의미 있는 정수 상수 이름을 만들고, union은 여러 멤버가 저장 공간을 공유하도록 한다. typedef는 형식 이름에 별칭을 붙인다. 비트 필드는 작은 비트 폭의 정수 필드를 표현할 때 사용할 수 있으나 배치와 이식성에 주의해야 한다.',
  'terms':'enum Stage { READY, RUNNING, DONE };의 열거 상수는 정수 상수이며 기본값은 차례로 0,1,2다. union Value { int number; double real; };은 멤버가 저장 공간을 공유하므로 마지막으로 저장한 멤버를 추적해야 한다. typedef struct Point Point;는 형식을 새로 만드는 것이 아니라 이름을 추가한다. unsigned int flag:1; 같은 비트 필드는 멤버에 지정한 비트 수를 사용한다. 비트 필드의 주소는 취할 수 없고 배치 순서는 구현에 따라 달라질 수 있다.',
  'goals':['열거 상수의 값과 switch 분기를 연결한다.','공용체의 활성 멤버를 추적하며 다른 멤버로 잘못 읽지 않는다.','typedef와 비트 필드의 용도·한계·이식성 문제를 구별한다.']}
]

EXAMPLES=[
 dict(sec='1.38',id='A',name='구조체 객체와 점 연산자',prediction='id=4, score=70으로 시작한 뒤 score에 5를 더한다. 출력할 두 값은 무엇일까?',
 code=r'''#include <stdio.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student s = {4, 70};
    s.score += 5;
    printf("id=%d score=%d\n", s.id, s.score);
    return 0;
}''', expected='id=4 score=75\n',
 explain=['struct Student {...};는 멤버 형식과 이름을 정의한다. 마지막 세미콜론은 형식 선언을 마친다.','struct Student s = {4, 70};는 선언 순서에 맞춰 id=4, score=70으로 초기화한다.','s.score는 s라는 객체의 score 멤버를 선택한다. += 5는 이전 값 70을 읽고 75로 저장한다.','printf의 두 %d는 정수 멤버 값을 차례로 출력한다. printf는 학습용 확인 출력이지 기록을 갱신하는 핵심 로직은 아니다.'],
 trace='객체 생성: s=(id 4, score 70) → s.score += 5 수행 후 (4,75) → 두 멤버를 읽어 출력. 두 멤버는 서로 다른 의미를 가지지만 하나의 구조체 객체에 속한다.',
 mistake='“struct Student라는 형식 선언만으로 s가 만들어진다”는 틀리다. 형식 선언과 실제 객체 선언을 구별한다. 디버깅: s->score는 s가 포인터가 아니므로 잘못된 접근이다. s.score로 고친다.',
 exercise='id를 9, 초기 score를 60으로 바꾸고 score에 12를 더하라. 결과를 먼저 예상한 다음 실행하라.',
 variant=r'''#include <stdio.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student s = {9, 60};
    s.score += 12;
    printf("id=%d score=%d\n", s.id, s.score);
    return 0;
}''', vexpected='id=9 score=72\n',vwhy='id는 그대로 9이고 score만 60→72로 변한다.'),
 dict(sec='1.38',id='B',name='구조체 대입은 멤버 값을 복사한다',prediction='a를 b에 대입한 다음 b.score만 바꾼다. a.score까지 바뀔까?',
 code=r'''#include <stdio.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student a = {2, 80};
    struct Student b = {0, 0};
    b = a;
    b.score = 95;
    printf("a=%d b=%d\n", a.score, b.score);
    return 0;
}''',expected='a=80 b=95\n',
 explain=['같은 구조체 형식의 객체끼리는 b = a처럼 대입할 수 있다. 멤버 값이 b에 복사된다.','대입 직후 b.id=2, b.score=80이지만 a와 b는 별도 객체다.','b.score=95는 b의 멤버만 바꾸며 a.score=80은 유지된다.','구조체 안에 포인터가 있다면 포인터 값도 복사된다. 그 포인터가 가리키는 별도 객체까지 깊게 복사된다는 뜻은 아니다. 이 주제는 동적 메모리 절에서 다시 다룬다.'],
 trace='초기 a=(2,80), b=(0,0) → b=a 뒤 b=(2,80) → b.score=95 뒤 a=(2,80), b=(2,95).',
 mistake='“구조체 대입은 항상 두 객체를 같은 객체로 연결한다”는 오개념이다. 구조체 객체 대입은 값을 복사한다. 디버깅: 서로 다른 구조체 형식은 멤버 모양이 같아도 일반적으로 직접 대입할 수 없다.',
 exercise='b의 id도 5로 수정하고 a.id와 b.id를 출력하라.',
 variant=r'''#include <stdio.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student a = {2, 80};
    struct Student b = a;
    b.id = 5;
    printf("a=%d b=%d\n", a.id, b.id);
    return 0;
}''',vexpected='a=2 b=5\n',vwhy='b.id만 2→5로 변경했으며 a.id는 2다.'),
 dict(sec='1.38',id='C',name='중첩 구조체의 멤버를 따라가기',prediction='위치의 x만 바꾼다면 y와 학생 번호는 유지될까?',
 code=r'''#include <stdio.h>
struct Point { int x; int y; };
struct Marker { int id; struct Point position; };
int main(void)
{
    struct Marker m = {3, {4, 6}};
    m.position.x += 2;
    printf("%d:(%d,%d)\n", m.id, m.position.x, m.position.y);
    return 0;
}''',expected='3:(6,6)\n',
 explain=['struct Point는 좌표 두 값을 묶는다. struct Marker는 정수 id와 struct Point position을 멤버로 가진다.','{3, {4, 6}}의 안쪽 중괄호는 중첩된 position을 초기화한다.','m.position은 Point 구조체 멤버이고, m.position.x는 그 안의 x 멤버다. 점 연산자를 두 번 적용한다.','x에만 2를 더하므로 y와 id는 변하지 않는다. 중첩 구조체는 관련 정보를 단계적으로 묶는 도구다.'],
 trace='m.id=3, m.position=(4,6) → m.position.x=6 → 최종 3:(6,6).',
 mistake='“m.x로도 중첩 멤버에 접근할 수 있다”는 틀리다. x는 m의 직접 멤버가 아니라 position의 멤버다. 디버깅: m.position->x는 position이 포인터가 아니므로 .을 사용한다.',
 exercise='초기 좌표를 (1,8)로 바꾸고 y를 3 감소시켜 출력하라.',
 variant=r'''#include <stdio.h>
struct Point { int x; int y; };
struct Marker { int id; struct Point position; };
int main(void)
{
    struct Marker m = {5, {1, 8}};
    m.position.y -= 3;
    printf("%d:(%d,%d)\n", m.id, m.position.x, m.position.y);
    return 0;
}''',vexpected='5:(1,5)\n',vwhy='y만 8→5로 바뀌며 id=5, x=1은 유지된다.'),
 dict(sec='1.38',id='D',name='지정 초기화로 멤버의 뜻 드러내기',prediction='멤버 선언 순서와 다른 순서로 지정하면 실제로 어떤 멤버에 값이 들어갈까?',
 code=r'''#include <stdio.h>
struct Sensor { int channel; int value; int valid; };
int main(void)
{
    struct Sensor s = {.valid = 1, .channel = 7, .value = 42};
    printf("ch=%d value=%d valid=%d\n", s.channel, s.value, s.valid);
    return 0;
}''',expected='ch=7 value=42 valid=1\n',
 explain=['.valid = 1은 구조체의 valid 멤버를 이름으로 지정해 초기화한다. C99 이후 표준 C에서 지원한다.','지정 초기화는 작성 순서보다 멤버 이름이 중요하다. 선언 순서와 다른 순서로 적어도 대응 멤버에 저장된다.','멤버 이름을 잘못 쓰면 컴파일러가 오류를 발견할 수 있다.','실제 출력은 초기화 결과를 확인하는 역할이며 구조체의 의미는 센서 측정값과 유효성 표시를 묶는 데 있다.'],
 trace='valid에 1, channel에 7, value에 42를 지정 → 출력은 printf 인수 순서 channel, value, valid로 읽는다.',
 mistake='“지정 초기화는 항상 선언 순서대로 적어야 한다”는 오개념이다. 단, 지정한 이름이 실제 멤버인지 확인해야 한다. 디버깅: .values=42는 존재하지 않는 멤버이므로 .value로 수정한다.',
 exercise='channel=2, value=18, valid=0을 멤버 이름으로 지정해 초기화하라.',
 variant=r'''#include <stdio.h>
struct Sensor { int channel; int value; int valid; };
int main(void)
{
    struct Sensor s = {.value = 18, .valid = 0, .channel = 2};
    printf("ch=%d value=%d valid=%d\n", s.channel, s.value, s.valid);
    return 0;
}''',vexpected='ch=2 value=18 valid=0\n',vwhy='초기화 순서는 달라도 지정된 멤버 이름에 따라 저장된다.'),
 dict(sec='1.39',id='A',name='구조체 배열에서 최고 점수 찾기',prediction='세 기록 중 score가 가장 큰 학생의 id는 무엇인가?',
 code=r'''#include <stdio.h>
#include <stddef.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student a[] = {{11, 72}, {12, 91}, {13, 85}};
    size_t count = sizeof a / sizeof a[0];
    size_t best = 0;
    for (size_t i = 1; i < count; ++i) {
        if (a[i].score > a[best].score) best = i;
    }
    printf("best=%d score=%d\n", a[best].id, a[best].score);
    return 0;
}''',expected='best=12 score=91\n',
 explain=['a[]는 구조체 세 개를 보관한다. 각 원소는 id와 score를 함께 가진다.','sizeof a / sizeof a[0]는 현재 함수 안에서 실제 배열 객체의 원소 수를 계산한다.','best=0은 현재 최고 기록의 배열 인덱스다. 반복은 1부터 시작해 아직 확인하지 않은 기록을 비교한다.','a[i].score > a[best].score일 때만 best를 교체한다. 실제 id와 인덱스는 서로 다른 값이다.'],
 trace='처음 best=0(72) → i=1에서 91>72이므로 best=1 → i=2에서 85>91이 거짓 → a[1].id=12 출력.',
 mistake='“best에는 학생 id가 들어 있다”는 틀리다. best는 배열 인덱스다. 디버깅: 반복을 i<=count로 쓰면 마지막에 배열 밖을 읽는다. i<count로 수정한다.',
 exercise='세 점수를 72, 88, 93으로 바꾸고 최고 점수 id를 확인하라.',
 variant=r'''#include <stdio.h>
#include <stddef.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student a[] = {{11, 72}, {12, 88}, {13, 93}};
    size_t count = sizeof a / sizeof a[0];
    size_t best = 0;
    for (size_t i = 1; i < count; ++i) {
        if (a[i].score > a[best].score) best = i;
    }
    printf("best=%d score=%d\n", a[best].id, a[best].score);
    return 0;
}''',vexpected='best=13 score=93\n',vwhy='최고 점수 인덱스가 2가 되어 id=13을 읽는다.'),
 dict(sec='1.39',id='B',name='구조체를 값으로 전달할 때',prediction='함수 안에서 매개변수의 점수를 100으로 바꿔도 호출자의 점수가 바뀔까?',
 code=r'''#include <stdio.h>
struct Student { int id; int score; };
int raised(struct Student copy)
{
    copy.score += 10;
    return copy.score;
}
int main(void)
{
    struct Student original = {4, 70};
    int result = raised(original);
    printf("original=%d result=%d\n", original.score, result);
    return 0;
}''',expected='original=70 result=80\n',
 explain=['raised(struct Student copy)는 구조체 값을 매개변수로 받는다. 호출 시 원본 멤버 값이 copy에 복사된다.','copy.score += 10은 복사본의 score만 70→80으로 변경한다.','return copy.score는 80이라는 정수 값을 돌려준다. 함수가 끝나면 지역 매개변수 copy의 수명도 끝난다.','원본은 바뀌지 않으므로 값을 반환해 호출자가 새 결과를 사용할지 결정해야 한다.'],
 trace='main original.score=70 → 호출 시 copy.score=70 → 함수 내부 copy.score=80 → 반환값 result=80 → original.score=70.',
 mistake='“구조체가 크면 C가 자동으로 주소를 전달한다”는 틀리다. C의 인수 전달은 값에 의한 전달이다. 컴파일러 최적화와 언어 의미론을 혼동하지 않는다.',
 exercise='점수 증가량을 20으로 바꾸고 원본·반환값을 확인하라.',
 variant=r'''#include <stdio.h>
struct Student { int id; int score; };
int raised(struct Student copy)
{
    copy.score += 20;
    return copy.score;
}
int main(void)
{
    struct Student original = {4, 70};
    int result = raised(original);
    printf("original=%d result=%d\n", original.score, result);
    return 0;
}''',vexpected='original=70 result=90\n',vwhy='원본은 70이고 복사본의 결과 90만 반환된다.'),
 dict(sec='1.39',id='C',name='구조체 포인터와 화살표 연산자',prediction='add_bonus(&s, 7) 뒤 s.score는 얼마인가?',
 code=r'''#include <stdio.h>
struct Student { int id; int score; };
void add_bonus(struct Student *s, int bonus)
{
    if (s != NULL) s->score += bonus;
}
int main(void)
{
    struct Student s = {8, 81};
    add_bonus(&s, 7);
    printf("id=%d score=%d\n", s.id, s.score);
    return 0;
}''',expected='id=8 score=88\n',
 explain=['&s는 호출자 구조체 객체의 주소를 구한다. 함수 매개변수 s에는 그 주소 값이 복사된다.','if (s != NULL)은 널 포인터를 역참조하지 않도록 확인한다. 널이 아닌 모든 포인터가 유효하다는 보장은 별도로 필요하다.','s->score는 (*s).score와 같으며 포인터가 가리키는 구조체의 score를 선택한다.','score는 원본 객체의 멤버이므로 81→88로 변한다. void 함수는 결과를 반환하지 않아도 원본을 수정할 수 있다.'],
 trace='main의 s=(8,81) → &s를 전달 → 함수에서 s->score=88 → main으로 돌아온 뒤 같은 객체를 출력.',
 mistake='“s->score에서 ->는 단순히 보기 좋은 점(.)의 다른 표기다”는 틀리다. 왼쪽 피연산자가 구조체 포인터일 때 사용한다. 디버깅: s.score는 포인터 s에 적용할 수 없다.',
 exercise='보너스를 4로 바꾸고 초기 점수를 92로 변경하라.',
 variant=r'''#include <stdio.h>
struct Student { int id; int score; };
void add_bonus(struct Student *s, int bonus)
{
    if (s != NULL) s->score += bonus;
}
int main(void)
{
    struct Student s = {8, 92};
    add_bonus(&s, 4);
    printf("id=%d score=%d\n", s.id, s.score);
    return 0;
}''',vexpected='id=8 score=96\n',vwhy='원본 score에 4를 더해 96이 된다.'),
 dict(sec='1.39',id='D',name='구조체 배열을 점수 순서로 정렬하기',prediction='처음 50,90,70 순서인 배열은 선택 정렬 후 어떤 순서가 될까?',
 code=r'''#include <stdio.h>
#include <stddef.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student a[] = {{1, 50}, {2, 90}, {3, 70}};
    size_t n = sizeof a / sizeof a[0];
    for (size_t i = 0; i < n; ++i) {
        size_t best = i;
        for (size_t j = i + 1; j < n; ++j) {
            if (a[j].score > a[best].score) best = j;
        }
        struct Student temp = a[i];
        a[i] = a[best];
        a[best] = temp;
    }
    for (size_t i = 0; i < n; ++i) {
        printf("%d:%d\n", a[i].id, a[i].score);
    }
    return 0;
}''',expected='2:90\n3:70\n1:50\n',
 explain=['바깥 반복의 i는 이번에 확정할 위치다. 안쪽 반복은 i부터 끝까지 가장 큰 점수의 인덱스 best를 찾는다.','temp는 a[i] 구조체 전체를 보관한다. 멤버 score만 바꾸면 id와 score가 어긋날 수 있으므로 구조체 전체를 교환한다.','a[i]=a[best], a[best]=temp는 두 구조체 객체의 멤버 값을 교환한다.','선택 정렬은 비교를 반복한다. 이 코드는 학습용이며 대량 데이터에서는 표준 qsort나 더 효율적인 정렬 알고리즘을 검토한다.'],
 trace='초기 (1:50,2:90,3:70) → i=0에서 최고 2:90과 교환 → (2:90,1:50,3:70) → i=1에서 3:70과 교환 → (2:90,3:70,1:50).',
 mistake='“점수만 교환해도 학생 기록을 정렬한 것이다”는 틀리다. id와 score를 하나의 기록으로 유지해야 한다. 디버깅: temp를 int로 선언하면 구조체 전체를 저장할 수 없다.',
 exercise='점수를 60,40,80으로 바꿔 높은 순서로 출력하라.',
 variant=r'''#include <stdio.h>
#include <stddef.h>
struct Student { int id; int score; };
int main(void)
{
    struct Student a[] = {{1, 60}, {2, 40}, {3, 80}};
    size_t n = sizeof a / sizeof a[0];
    for (size_t i = 0; i < n; ++i) {
        size_t best = i;
        for (size_t j = i + 1; j < n; ++j) {
            if (a[j].score > a[best].score) best = j;
        }
        struct Student temp = a[i];
        a[i] = a[best];
        a[best] = temp;
    }
    for (size_t i = 0; i < n; ++i) {
        printf("%d:%d\n", a[i].id, a[i].score);
    }
    return 0;
}''',vexpected='3:80\n1:60\n2:40\n',vwhy='구조체 전체가 교환되므로 점수와 id의 대응 관계가 유지된다.'),
 dict(sec='1.40',id='A',name='열거형 상수와 switch',prediction='RUNNING의 기본 정수 값은 무엇이고 어느 분기로 들어갈까?',
 code=r'''#include <stdio.h>
enum Stage { READY, RUNNING, DONE };
int main(void)
{
    enum Stage state = RUNNING;
    switch (state) {
        case READY: puts("ready"); break;
        case RUNNING: puts("running"); break;
        case DONE: puts("done"); break;
        default: puts("unknown"); break;
    }
    printf("code=%d\n", (int)state);
    return 0;
}''',expected='running\ncode=1\n',
 explain=['enum Stage는 READY, RUNNING, DONE이라는 이름을 가진 정수 상수를 정의한다. 기본값은 0,1,2다.','state=RUNNING으로 저장하므로 switch의 RUNNING 레이블로 이동한다.','break는 다음 case로 실행이 이어지는 것을 막는다.','(int)state는 여기서 열거형 값을 정수로 명시적으로 변환한다. 열거형이 임의의 문자열을 자동으로 저장하는 것은 아니다.'],
 trace='READY=0, RUNNING=1, DONE=2 → state=RUNNING(1) → running 출력 → code=1 출력.',
 mistake='“enum은 문자열 목록이다”는 틀리다. 열거 상수는 정수 상수다. 디버깅: case RUNNING 뒤 break를 빼면 다음 case의 문장까지 실행될 수 있다.',
 exercise='state를 DONE으로 바꾸고 결과를 예측하라.',
 variant=r'''#include <stdio.h>
enum Stage { READY, RUNNING, DONE };
int main(void)
{
    enum Stage state = DONE;
    switch (state) {
        case READY: puts("ready"); break;
        case RUNNING: puts("running"); break;
        case DONE: puts("done"); break;
        default: puts("unknown"); break;
    }
    printf("code=%d\n", (int)state);
    return 0;
}''',vexpected='done\ncode=2\n',vwhy='DONE은 기본값 2이며 해당 case만 실행한다.'),
 dict(sec='1.40',id='B',name='열거 상수의 명시적 값과 이어지는 값',prediction='LOW=2, MEDIUM을 생략, HIGH=8이면 MEDIUM은 몇인가?',
 code=r'''#include <stdio.h>
enum Level { LOW = 2, MEDIUM, HIGH = 8 };
int main(void)
{
    printf("%d %d %d\n", LOW, MEDIUM, HIGH);
    return 0;
}''',expected='2 3 8\n',
 explain=['LOW=2는 열거 상수의 값을 직접 지정한다.','값을 생략한 MEDIUM은 바로 앞 열거 상수보다 1 큰 3이 된다.','HIGH=8은 값을 다시 명시하므로 8이다. 값이 반드시 0부터 연속적일 필요는 없다.','열거형으로 모든 잘못된 정수 값이 자동 차단되는 것은 아니다. 유효한 상태인지는 프로그램에서 별도로 검사할 수 있다.'],
 trace='LOW를 2로 설정 → MEDIUM은 3 → HIGH는 명시적으로 8 → 출력 2 3 8.',
 mistake='“모든 enum 멤버 값은 0,1,2로 고정된다”는 틀리다. 명시적 값을 지정할 수 있다. 디버깅: 서로 다른 의미에 같은 값을 배정하면 switch에서 중복 case 레이블 오류가 날 수 있다.',
 exercise='HIGH=10으로 바꾸고 세 값을 출력하라.',
 variant=r'''#include <stdio.h>
enum Level { LOW = 2, MEDIUM, HIGH = 10 };
int main(void)
{
    printf("%d %d %d\n", LOW, MEDIUM, HIGH);
    return 0;
}''',vexpected='2 3 10\n',vwhy='MEDIUM은 3으로 유지되고 HIGH만 10이 된다.'),
 dict(sec='1.40',id='C',name='공용체와 활성 멤버를 구별하기',prediction='number에 저장한 뒤 real에 다시 저장하면 어느 멤버 값을 읽어야 할까?',
 code=r'''#include <stdio.h>
union Value { int number; double real; };
int main(void)
{
    union Value v;
    v.number = 12;
    printf("number=%d\n", v.number);
    v.real = 2.5;
    printf("real=%.1f\n", v.real);
    return 0;
}''',expected='number=12\nreal=2.5\n',
 explain=['union Value는 int와 double 멤버가 저장 공간을 공유하는 공용체 형식이다. 구조체와 달리 두 값을 독립적으로 동시에 보관하는 용도가 아니다.','v.number=12 뒤에는 number를 읽어 출력한다.','v.real=2.5는 같은 공용체의 다른 멤버에 값을 저장한다. 이후 real을 읽는다.','이 예제는 마지막으로 저장한 멤버를 읽으므로 이식성 있는 기초 사용법을 보여 준다. 다른 멤버의 비트 표현을 임의로 해석하는 예제는 피한다.'],
 trace='number에 12 저장·읽기 → real에 2.5 저장·읽기. 두 값을 동시에 보존하는 상황이 아니다.',
 mistake='“공용체는 구조체처럼 모든 멤버의 값을 동시에 따로 보관한다”는 틀리다. 디버깅: real에 저장한 뒤 number를 이전 값 12라고 가정하지 않는다.',
 exercise='number=30을 출력한 다음 real=1.25를 저장·출력하라.',
 variant=r'''#include <stdio.h>
union Value { int number; double real; };
int main(void)
{
    union Value v;
    v.number = 30;
    printf("number=%d\n", v.number);
    v.real = 1.25;
    printf("real=%.2f\n", v.real);
    return 0;
}''',vexpected='number=30\nreal=1.25\n',vwhy='각 시점에서 마지막으로 저장한 멤버만 읽는다.'),
 dict(sec='1.40',id='D',name='typedef와 비트 필드의 정확한 역할',prediction='3비트 unsigned int 필드에 5를 저장하면 출력은 무엇일까?',
 code=r'''#include <stdio.h>
typedef struct Flags {
    unsigned int active : 1;
    unsigned int level : 3;
} Flags;
int main(void)
{
    Flags f = {1, 5};
    printf("active=%u level=%u\n", f.active, f.level);
    return 0;
}''',expected='active=1 level=5\n',
 explain=['typedef ... Flags는 struct Flags 형식에 Flags라는 별칭을 부여한다. 별도의 새 저장 객체를 만들지 않는다.','active:1은 부호 없는 1비트 필드로 0 또는 1을 표현한다. level:3은 0부터 7까지를 표현할 수 있다.','Flags f={1,5}는 두 필드를 순서대로 초기화한다. 1과 5는 각 필드의 범위 안에 있다.','비트 필드의 배치 방향·패딩·구조체 전체 크기는 구현에 따라 달라질 수 있다. 통신 파일의 정확한 비트 배치를 위해서는 비트 마스크와 명시적 직렬화를 별도로 설계한다.'],
 trace='f.active=1, f.level=5 → 두 멤버를 %u로 출력. printf 인수 전달 시 정수 승격이 일어나며 이 예제의 값은 int로 표현 가능하다. 엄격한 서식 일치를 위해 (unsigned int)로 캐스팅한 코드를 권장하며 아래 코드는 수정해 적용한다.',
 mistake='“비트 필드에 주소 연산자 &를 적용할 수 있다”는 틀리다. 비트 필드는 주소를 직접 취할 수 없다. 디버깅: level=8은 3비트 범위를 벗어나므로 0~7 값만 사용한다.',
 exercise='active=0, level=7로 바꾸고 출력하라.',
 variant=r'''#include <stdio.h>
typedef struct Flags {
    unsigned int active : 1;
    unsigned int level : 3;
} Flags;
int main(void)
{
    Flags f = {0, 7};
    printf("active=%u level=%u\n", (unsigned int)f.active, (unsigned int)f.level);
    return 0;
}''',vexpected='active=0 level=7\n',vwhy='두 값 모두 지정한 비트 폭 안에 있으며 그대로 출력된다.')
]
# Format correction: integer promotions on bit fields mean %u needs explicit unsigned int cast.
EXAMPLES[-1]['code']=EXAMPLES[-1]['code'].replace('f.active, f.level','(unsigned int)f.active, (unsigned int)f.level')
EXAMPLES[-1]['trace']='f.active=1, f.level=5 → 명시적으로 unsigned int로 변환한 두 값을 %u로 출력한다. 비트 필드의 저장 폭과 printf에 전달되는 인수 형식은 구별해야 한다.'

assert len(EXAMPLES)==12

def p(doc,text='',style=None):
    para=doc.add_paragraph(style=style)
    para.add_run(text)
    return para

def heading(doc,text,level=1):
    para=doc.add_heading(text,level=level)
    para.paragraph_format.keep_with_next=True
    return para

def code(doc,s):
    para=doc.add_paragraph(style='CodeBlock')
    para.paragraph_format.keep_together=True
    para.add_run(s)
    return para

def make_doc():
    d=Document();sec=d.sections[0]
    sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.35);sec.right_margin=Cm(2.25)
    styles=d.styles
    normal=styles['Normal'];normal.font.name='Noto Sans CJK KR';normal.font.size=Pt(9.6);normal.font.color.rgb=GRAY
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.19
    for nm,size in [('Title',22),('Heading 1',16),('Heading 2',12.3),('Heading 3',10.2)]:
        st=styles[nm];st.font.name='Noto Sans CJK KR';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=TEAL
        st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10)
    if 'CodeBlock' not in styles:
        st=styles.add_style('CodeBlock',WD_STYLE_TYPE.PARAGRAPH)
    st=styles['CodeBlock'];st.font.name='DejaVu Sans Mono';st.font.size=Pt(8.05);st.font.color.rgb=RGBColor(28,43,51)
    st.paragraph_format.left_indent=Cm(.38);st.paragraph_format.right_indent=Cm(.2)
    st.paragraph_format.space_before=Pt(4);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.03
    st.paragraph_format.keep_together=True
    title=d.add_paragraph(style='Title');title.add_run('C언어 종합교과서')
    p(d,'이해 중심 정식 개정판 · 제1장 1.38~1.40절 · 독립 집필 원고')
    p(d,'기준: ISO C11 / GCC·Clang 엄격 경고 검사. 이전 거부 원고는 재사용하지 않았으며, 원서 두 권의 모든 문제·LAB 학습목표 전수 대조는 아직 완료되지 않았다.')
    p(d,'읽는 흐름: 필요성 → 새 용어 → 예측 → 전체 코드 → 문장별 해설 → 실행 추적과 실제 출력 → 오개념·디버깅 → 바로 인접한 변형 실습과 정답 → 도움 없는 독립 작성.')
    for secdata in SECTIONS:
        heading(d,secdata['number']+' '+secdata['title'],1)
        heading(d,'왜 필요한가 · 실제 사용 장면',2);p(d,secdata['why'])
        heading(d,'새 용어와 표현을 먼저 이해하기',2);p(d,secdata['terms'])
        heading(d,'이 절을 마치면',2)
        for g in secdata['goals']:p(d,'• '+g)
        for e in [x for x in EXAMPLES if x['sec']==secdata['number']]:
            heading(d,f"{e['sec']}-{e['id']} {e['name']}",2)
            p(d,'작은 목표: ① 다룰 객체·형식을 확인한다. ② 실행을 먼저 예측한다. ③ 멤버 값과 제어 흐름을 추적한다. ④ 조건을 바꾸어 직접 검증한다.')
            heading(d,'실행 전에 예측하기',3);p(d,e['prediction'])
            heading(d,'완성된 전체 C11 코드',3);code(d,e['code'])
            heading(d,'문장과 식을 차례대로 이해하기',3)
            for j,t in enumerate(e['explain'],1):p(d,f'{j}. {t}')
            heading(d,'실행 중 값의 변화',3);p(d,e['trace'])
            heading(d,'실제로 검증한 출력',3);code(d,e['expected'].rstrip('\n'))
            heading(d,'오개념 진단과 단계별 디버깅',3);p(d,e['mistake'])
            heading(d,'예제를 바꾸며 실력 다지기',3);p(d,e['exercise'])
            heading(d,'바로 이어지는 변형 정답 코드',3);code(d,e['variant'])
            heading(d,'변형 출력과 정답 근거',3);code(d,e['vexpected'].rstrip('\n'));p(d,e['vwhy'])
        heading(d,'힌트를 줄이며 독립적으로 풀기',2)
        if secdata['number']=='1.38':
            p(d,'부분 완성 활동: struct Book { int code; int pages; };를 선언한 다음 Book이 아닌 struct Book 형식의 객체를 만들고 pages를 수정하는 코드를 직접 완성하라. 새 개념인 typedef를 아직 배우지 않았으므로 struct 키워드를 생략하지 않는다.')
            p(d,'독립 작성: 제목 대신 정수 코드와 페이지 수를 갖는 책 기록 두 개를 만들고, 두 기록의 페이지 수 합계를 출력하라. 정답의 핵심: 각각의 .pages를 읽어 더하며, 구조체 객체는 별개로 생성한다.')
        elif secdata['number']=='1.39':
            p(d,'부분 완성 활동: struct Student a[]={{1,70},{2,90}};에서 점수 합계를 구하는 반복문의 시작값·종료 조건·멤버 접근을 직접 채워 보라. 먼저 인덱스가 0,1만 유효함을 말로 설명한다.')
            p(d,'독립 작성: 학생 세 명의 기록에서 점수 80 이상인 학생 수를 세는 프로그램을 처음부터 작성하라. 정답의 핵심: 배열을 순회하며 a[i].score >= 80일 때 지역 카운터를 증가시킨다.')
        else:
            p(d,'부분 완성 활동: enum Light { OFF=0, ON=1 };의 값으로 switch 분기를 작성하고, OFF/ON 이외의 정수가 들어왔을 때 default에서 처리하도록 한다.')
            p(d,'독립 작성: union에 int와 double 멤버를 선언하고 각각을 저장한 직후 같은 멤버를 출력하라. 이전 멤버 값이 동시에 보존된다고 가정하지 않는다.')
        heading(d,'누적 복습 · 스스로 인출하기',2)
        if secdata['number']=='1.38':
            p(d,'질문: (1) 배열 원소와 구조체 멤버의 접근 연산자는 각각 무엇인가? (2) 구조체 복사 뒤 한쪽 멤버를 바꾸면 다른 객체가 자동으로 변하는가?')
            p(d,'정답: (1) 배열은 [], 구조체 객체는 .을 쓴다. (2) 아니다. 구조체 대입은 멤버 값을 복사하며 객체가 합쳐지지 않는다.')
        elif secdata['number']=='1.39':
            p(d,'질문: (1) 구조체 포인터 p의 score를 선택하는 표현 두 가지는? (2) 배열 경계 조건은 i<n과 i<=n 중 무엇인가? (3) 선택 정렬에서 점수만 교환하면 왜 위험한가?')
            p(d,'정답: (1) p->score와 (*p).score. (2) i<n. (3) id와 score의 대응이 깨지므로 구조체 전체를 교환한다.')
        else:
            p(d,'질문: (1) typedef는 메모리를 할당하는가? (2) 공용체의 두 멤버가 독립 저장 공간을 가지는가? (3) 3비트 unsigned 비트 필드에 8을 저장해도 되는가?')
            p(d,'정답: (1) 아니다. 형식 별칭을 만든다. (2) 아니다. 저장 공간을 공유한다. (3) 8은 3비트 범위 0~7 밖이다.')
    heading(d,'선행 개념과 표준 경계 확인',1)
    p(d,'이 원고에서 처음 정식 소개한 문법: struct 태그·멤버 접근(.)·중첩 구조체·지정 초기화(1.38), 구조체 값 전달·구조체 포인터와 ->(1.39), enum·union·typedef·비트 필드(1.40). 앞 절에서 배운 배열, 함수, 포인터, switch, sizeof, size_t, 반복문을 재사용한다.')
    p(d,'구조체와 공용체의 실제 바이트 배치·정렬·패딩, 비트 필드의 배치 순서, 특정 ABI는 구현에 따라 달라질 수 있다. 본문의 예제는 그 배치를 전제하지 않는다. C23의 열거형 관련 확장은 이 절의 C11 코드와 구분해 이후 표준 비교 절에서 설명한다.')
    p(d,'학습목표 출처 확인: K&R 6장 구조체 분야 및 Express 13장 구조체 분야와 관련성이 있으나, 개별 원서 예제·LAB·Exercise 전수 대조는 미완료이다. 원서별 항목을 완료로 표시하지 않는다.')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('C언어 종합교과서 · 이해 중심 정식 개정판')
    out=ROOT/'C언어_종합교과서_이해중심_정식개정판_제1장_1.38-1.40.docx';d.save(out)
    return out

def verify():
    src=ROOT/'sources';src.mkdir(exist_ok=True)
    checks=[]
    for e in EXAMPLES:
        for variant in (False,True):
            name=f"ch01_{e['sec'].replace('.','_')}_{e['id']}_{'variant' if variant else 'main'}"
            content=e['variant'] if variant else e['code'];expected=e['vexpected'] if variant else e['expected']
            source=src/(name+'.c');source.write_text(content+'\n',encoding='utf-8')
            for compiler in ('gcc','clang'):
                exe=src/(name+'_'+compiler)
                c=subprocess.run([compiler,*FLAGS,str(source),'-o',str(exe)],capture_output=True,text=True)
                if c.returncode:raise RuntimeError(f'{name} {compiler}: {c.stderr}')
                run=subprocess.run([str(exe)],capture_output=True,text=True,timeout=5)
                if run.returncode or run.stdout!=expected:raise RuntimeError(f'{name} {compiler}: output {run.stdout!r} vs {expected!r}, rc={run.returncode}')
                checks.append(dict(name=name,compiler=compiler,compile_ok=True,run_ok=True,expected=expected,stdout=run.stdout))
                exe.unlink()
    z=ROOT/'official_ch01_1_38-1_40_sources.zip'
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as f:
        for file in sorted(src.glob('*.c')):f.write(file,arcname='sources/'+file.name)
    return checks,z

if __name__=='__main__':
    doc=make_doc();checks,z=verify()
    from docx import Document as ReadDoc
    paragraphs=[p.text for p in ReadDoc(doc).paragraphs]
    for e in EXAMPLES:
        assert e['code'] in paragraphs and e['variant'] in paragraphs
    stamp=datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=9))).isoformat(timespec='seconds')
    info={'run':'chapter01_1.38-1.40','generated_at':stamp,'standard':'ISO C11','compilers':['gcc','clang'],'flags':FLAGS,'example_programs':12,'variant_programs':12,'compilation_and_execution_checks':len(checks),'passed':len(checks),'docx_source_match':24,'source_audit_complete':False,'tests':checks}
    (ROOT/'official_ch01_1_38-1_40_c11_verification.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'official_rewrite_progress_1_1-1_40.json').write_text(json.dumps({'last_completed_section':'1.40','completed_sections':[f'1.{i}' for i in range(1,41)],'overall_percent':None,'source_objective_audit':'pending','timestamp':stamp},ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'official_rewrite_objective_tracking_1_1-1_40.json').write_text(json.dumps({'full_audit_complete':False,'primary_sources':{'Express':'original scanned PDF; individual objectives not exhaustively inspected','K&R':'original scanned PDF; individual objectives not exhaustively inspected','K&R OCR':'TOC verified; chapter 6 relevant'},'new_sections':[{'section':s['number'],'topics':s['title'],'status':'original new content compiled and tested; source-by-source objective mapping pending'} for s in SECTIONS]},ensure_ascii=False,indent=2),encoding='utf-8')
    files=[doc,z,ROOT/'official_ch01_1_38-1_40_c11_verification.json',ROOT/'official_rewrite_progress_1_1-1_40.json',ROOT/'official_rewrite_objective_tracking_1_1-1_40.json',Path(__file__)]
    manifest={'run':'chapter01_1.38-1.40','generated_at':stamp,'files':[{'name':f.name,'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'size':f.stat().st_size} for f in files],'github_status':'pending','library_status':'pending'}
    (ROOT/'official_rewrite_manifest_1_38-1_40.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'docx':str(doc),'sources':str(z),'tests':len(checks),'paragraphs':len(paragraphs)},ensure_ascii=False))