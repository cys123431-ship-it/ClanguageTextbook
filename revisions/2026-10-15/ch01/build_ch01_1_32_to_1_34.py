#!/usr/bin/env python3
"""Original Korean C textbook section 1.32-1.34. C11 validated examples."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import subprocess, json, zipfile, hashlib, datetime, re
ROOT=Path('/mnt/data/c_textbook_run10'); ROOT.mkdir(exist_ok=True)
FLAGS=['-std=c11','-Wall','-Wextra','-Wpedantic','-Werror','-O0']
sections=[
{'number':'1.32','title':'주소와 포인터 — 값이 저장된 곳을 명시적으로 가리키기',
 'why':'변수 이름만으로는 함수 밖의 값을 직접 바꾸는 동작이나 배열 안의 위치를 표현하기 어렵다. 포인터는 어떤 객체가 놓인 메모리 위치를 가리키는 값이다. C에서는 포인터를 무작정 숫자 주소처럼 취급하지 않고, 어떤 형식의 객체를 가리키는지를 함께 다룬다. 그래서 int *는 int 객체를 가리키는 포인터 형식이다.',
 'before':'앞 절의 지역 변수·배열·함수 호출을 기억하자. 이번에는 새 연산자 &와 *를 배운다. 선언의 *는 포인터 형식을 만들고, 식의 *는 포인터가 가리키는 객체에 접근한다. 같은 기호지만 문맥과 역할이 다르다. 실제 숫자 주소는 실행 환경마다 달라지므로 이 절은 주소 숫자 대신 값의 변화로 확인한다.',
 'goals':['&로 실제 객체의 주소를 얻고, *로 가리키는 객체를 읽거나 바꾼다.','포인터 복사와 정수 값 복사가 어떻게 다른지 실행 흐름으로 구별한다.','NULL은 유효한 객체를 가리키지 않으며, 검사 없이 역참조하면 안 됨을 설명한다.']},
{'number':'1.33','title':'포인터 매개변수 — 함수가 호출자의 값을 바꾸도록 설계하기',
 'why':'C의 함수 인수는 기본적으로 값으로 전달된다. 따라서 함수 안에서 일반 int 매개변수를 바꿔도 호출자의 원본 int가 바뀌지 않는다. 반면 주소를 전달하면 함수는 그 주소가 가리키는 객체에 접근할 수 있다. 이를 이용하면 둘 이상의 결과를 계산하거나 입력값을 조건에 맞게 수정하는 함수를 설계할 수 있다.',
 'before':'포인터 매개변수도 값으로 전달된다는 점이 핵심이다. 전달되는 값이 주소이기 때문에 호출자와 피호출자가 같은 객체를 가리킬 수 있을 뿐이다. 호출자는 &score처럼 주소를 넘기고 함수는 int *score처럼 포인터를 받는다. 함수 안의 *score는 원본 객체를 뜻한다.',
 'goals':['값 전달과 주소 전달을 실행 추적으로 구별한다.','출력 매개변수와 반환값의 역할을 분리한다.','const를 사용해 포인터를 통해 읽기만 허용하는 의도를 표현한다.']},
{'number':'1.34','title':'배열과 포인터 — 원소의 위치와 유효한 이동 범위',
 'why':'배열의 연속된 원소를 하나씩 검사할 때 포인터는 현재 위치를 나타내는 도구가 된다. 다만 포인터 연산은 배열 객체 내부와 마지막 원소 바로 다음 위치를 기준으로 엄격한 범위 규칙을 갖는다. 이 규칙을 모르면 범위 밖 접근이나 정의되지 않은 동작을 만들기 쉽다.',
 'before':'배열 이름이 대부분의 식에서 첫 원소를 가리키는 포인터로 변환되는 규칙을 이번 절에서 소개한다. 배열 자체와 포인터 변수는 동일한 객체가 아니다. sizeof 배열은 배열 전체 크기이지만 sizeof 포인터는 포인터 객체의 크기이므로 둘을 혼동하면 안 된다. *(p+i)와 p[i]의 대응도 확인한다.',
 'goals':['배열 첫 원소 주소와 포인터의 관계를 설명한다.','포인터를 배열 안에서 이동시키고 마지막 다음 포인터는 역참조하지 않는다.','sizeof 배열과 sizeof 포인터의 차이를 설명한다.']}
]
E=[]
def add(sec,id,title,code,variant,ask,explain,trace,misconception,task,solution,debug):
 E.append(dict(sec=sec,id=id,title=title,code=code.strip()+'\n',variant=variant.strip()+'\n',ask=ask,explain=explain,trace=trace,misconception=misconception,task=task,solution=solution,debug=debug))
add('1.32','a','주소 얻기와 역참조',r'''
#include <stdio.h>
int main(void)
{
    int score = 7;
    int *where = &score;
    printf("before=%d through=%d\n", score, *where);
    *where = 12;
    printf("after=%d through=%d\n", score, *where);
    return 0;
}''',r'''
#include <stdio.h>
int main(void)
{
    int score = 7;
    int *where = &score;
    *where += 5;
    printf("score=%d\n", score);
    return 0;
}''','첫 줄에서 score와 *where의 값이 같은가? *where=12를 실행하면 어느 변수가 달라지는가?',
 ['int score = 7;은 실제 정수 객체를 만들고 7을 넣는다.','int *where = &score;의 &score는 그 객체의 주소다. where는 주소를 담는 포인터 변수이며 score와 별개의 변수다.','printf에서 *where는 where가 가리키는 int를 읽는다. 값이 7이므로 score와 같은 수가 나온다.','*where = 12;는 포인터 변수 where의 주소를 바꾸는 것이 아니라 그 주소에 있는 score 객체의 값을 12로 바꾼다.'],
 '초기 score=7, where→score. 첫 출력 뒤에도 7. *where=12 이후 score=12이며 where는 여전히 score를 가리킨다.',
 '“*where에 값을 대입했으니 where 자체가 12가 된다”는 틀리다. where는 주소를 보관하며 *where가 가리키는 int 객체다.',
 '*where=12를 *where+=5로 바꿔 같은 시작값 7에서 어떤 결과가 나오는지 확인하자.',
 '7에 5를 더하므로 score=12. 복합 대입은 가리키는 원본 score에 적용된다.',
 '선언을 int *where = score;로 바꾸면 정수 값을 주소처럼 쓰려는 잘못된 초기화다. &score로 수정한다.')
add('1.32','b','정수 복사와 포인터를 통한 수정',r'''
#include <stdio.h>
int main(void)
{
    int first = 10;
    int second = first;
    int *ptr = &first;
    *ptr += 5;
    printf("first=%d second=%d\n", first, second);
    return 0;
}''',r'''
#include <stdio.h>
int main(void)
{
    int first = 10;
    int second = first;
    int *ptr = &second;
    *ptr += 5;
    printf("first=%d second=%d\n", first, second);
    return 0;
}''','first를 15로 만들면 second도 15가 되는가? 먼저 두 객체의 관계를 그려 보자.',
 ['int second = first;는 first의 현재 숫자 10을 새 객체 second에 복사한다. 두 정수는 서로 독립적이다.','int *ptr = &first;는 first의 주소를 포인터에 넣는다.','*ptr += 5;는 first만 15로 만든다. second는 복사 당시의 10을 유지한다.'],
 '초기 first=10. 복사 뒤 second=10. ptr→first. 마지막에는 first=15, second=10.',
 '“값을 복사한 두 변수는 계속 연결된다”는 오개념이다. 포인터가 가리키는 원본만 변경되며 단순 정수 복사본은 자동 갱신되지 않는다.',
 'ptr이 first 대신 second를 가리키도록 한 줄을 바꾸어 두 변수의 최종값을 예상하자.',
 'ptr→second이므로 first=10, second=15. 어떤 객체를 가리키는지가 결과를 결정한다.',
 'ptr이 초기화되기 전에 *ptr += 5를 실행하면 정의되지 않은 동작이다. 항상 유효한 객체의 주소로 초기화한 뒤 사용한다.')
add('1.32','c','둘의 포인터가 하나의 객체를 가리킬 때',r'''
#include <stdio.h>
int main(void)
{
    int total = 20;
    int *left = &total;
    int *right = left;
    *right -= 3;
    printf("left=%d right=%d total=%d\n", *left, *right, total);
    return 0;
}''',r'''
#include <stdio.h>
int main(void)
{
    int total = 20;
    int *left = &total;
    int *right = left;
    *left += 2;
    *right -= 3;
    printf("total=%d\n", total);
    return 0;
}''','left와 right는 서로 다른 포인터 변수다. 그럼 *left와 *right는 서로 다른 객체인가?',
 ['left와 right는 주소를 담는 별도 포인터 변수지만 둘의 주소 값은 동일하다.','right = left는 주소 값을 복사하므로 둘 다 total 객체를 가리킨다.','*right -= 3은 total을 20에서 17로 변경한다. *left로 읽어도 같은 total의 현재값을 본다.'],
 'total=20 → left→total → right→total → *right로 3 감소 → total=17. 세 출력값 모두 17.',
 '“포인터 변수가 둘이므로 int 객체도 둘”이라는 결론은 틀리다. 포인터 변수의 수와 가리키는 대상 객체의 수는 별개다.',
 '같은 객체에 *left += 2, *right -= 3을 순서대로 실행하면 최종 total은?',
 '20→22→19이므로 total=19. 두 포인터가 하나의 객체를 공유한다.',
 '서로 다른 객체를 원한다면 int another = total;과 같이 새 객체를 만들고 &another를 사용해야 한다.')
add('1.32','d','널 포인터 검사 후 사용하기',r'''
#include <stdio.h>
#include <stddef.h>
int main(void)
{
    int score = 8;
    int *ptr = NULL;
    if (ptr == NULL) {
        puts("not connected");
    }
    ptr = &score;
    if (ptr != NULL) {
        printf("score=%d\n", *ptr);
    }
    return 0;
}''',r'''
#include <stdio.h>
#include <stddef.h>
int main(void)
{
    int score = 8;
    int *ptr = NULL;
    if (ptr == NULL) {
        puts("empty");
    }
    ptr = &score;
    *ptr += 2;
    printf("score=%d\n", score);
    return 0;
}''','NULL인 ptr에 대해 *ptr을 실행해도 되는가? 검사문은 어떤 잘못을 막는가?',
 ['NULL은 유효한 객체를 가리키지 않는 널 포인터를 표현하는 매크로다. <stddef.h>가 NULL을 제공한다.','첫 if는 연결 대상이 없다는 사실을 출력할 뿐 *ptr을 실행하지 않는다.','ptr = &score 이후에는 유효한 int 객체를 가리키므로 역참조가 가능하다.','널 검사는 필요조건일 뿐 충분조건은 아니다. NULL이 아니어도 수명이 끝난 객체나 범위 밖을 가리킬 수 있다.'],
 'ptr=NULL → not connected 출력 → ptr=&score → *ptr로 8 읽기 → score=8 출력.',
 '“NULL이 0이니 *NULL도 정수 0”이라는 주장은 틀리다. 널 포인터 역참조는 정의되지 않은 동작이다.',
 '검사 출력 문구를 empty로 바꾸고 연결 후 *ptr에 2를 더해 보자.',
 '첫 줄 empty, 다음 줄 score=10. 유효한 주소를 대입한 뒤에만 수정한다.',
 'if(ptr == NULL) { printf("%d", *ptr); }는 검사 방향이 반대다. ptr != NULL일 때만 역참조해야 한다.')
add('1.33','a','교환 함수: 두 객체의 값을 서로 바꾸기',r'''
#include <stdio.h>
static void exchange(int *a, int *b)
{
    int saved = *a;
    *a = *b;
    *b = saved;
}
int main(void)
{
    int x = 3;
    int y = 9;
    exchange(&x, &y);
    printf("x=%d y=%d\n", x, y);
    return 0;
}''',r'''
#include <stdio.h>
static void exchange(int *a, int *b)
{
    int saved = *a;
    *a = *b;
    *b = saved;
}
int main(void)
{
    int x = 14;
    int y = -2;
    exchange(&x, &y);
    printf("x=%d y=%d\n", x, y);
    return 0;
}''','saved를 두지 않고 *a=*b; *b=*a;라고 하면 제대로 교환될까?',
 ['static void는 이 파일에서만 쓰는 보조 함수의 선언이다. 앞 절에서 배운 파일 내부 연결과 함수의 반환값 없음(void)을 함께 적용한다.','a와 b에는 호출 시 넘긴 &x, &y의 주소 값이 각각 복사된다.','saved=*a로 x의 이전 값 3을 보관하지 않으면 첫 대입에서 원래 x 값이 사라진다.','*a=*b로 x=9가 되고 *b=saved로 y=3이 된다.'],
 '호출 전 x=3,y=9. a→x,b→y. saved=3. *a=*b 이후 x=9. *b=saved 이후 y=3.',
 '“포인터 인수이므로 C가 참조 전달을 한다”는 표현은 엄밀하지 않다. C는 포인터 값 자체를 값으로 전달한다. 주소를 통해 호출자 객체를 수정할 수 있을 뿐이다.',
 '초깃값을 x=14, y=-2로 바꾸어 교환 결과를 확인하자.',
 'x=-2 y=14. 교환 함수의 로직은 숫자의 부호와 무관하다.',
 'saved를 생략하고 *a=*b; *b=*a;로 쓰면 둘 다 원래 y가 된다. 임시 저장 공간이 필요하다.')
add('1.33','b','경계 제한 함수: 원본 정수를 안전한 구간으로',r'''
#include <stdio.h>
static void keep_between(int *value, int low, int high)
{
    if (*value < low) {
        *value = low;
    } else if (*value > high) {
        *value = high;
    }
}
int main(void)
{
    int temperature = 42;
    keep_between(&temperature, 0, 35);
    printf("limited=%d\n", temperature);
    return 0;
}''',r'''
#include <stdio.h>
static void keep_between(int *value, int low, int high)
{
    if (*value < low) {
        *value = low;
    } else if (*value > high) {
        *value = high;
    }
}
int main(void)
{
    int temperature = -7;
    keep_between(&temperature, 0, 35);
    printf("limited=%d\n", temperature);
    return 0;
}''','42가 0~35 범위를 벗어나는지, 함수가 값을 어떤 경계로 바꾸는지 예상하자.',
 ['value는 int 객체의 주소를 받는다. low와 high는 일반 정수 값 매개변수다.','*value < low면 원본 객체를 low로 수정하고, *value > high면 high로 수정한다.','범위 안에 이미 들어 있는 값이면 어떤 대입도 하지 않는다.','이 입문 예제는 호출자가 low<=high를 보장한다고 가정한다. 실무 API에서는 경계의 역전도 검사해야 한다.'],
 'temperature=42. 첫 조건 42<0 거짓, 두 번째 42>35 참. *value=35로 원본이 바뀐다.',
 '“함수 내부의 *value가 별도 복사본”이라는 주장은 틀리다. value 자체는 복사된 포인터지만 *value는 호출자 객체다.',
 'temperature의 초깃값을 -7로 바꾸면 어느 경계에 붙는가?',
 '처음 조건 -7<0이 참이므로 limited=0. else if는 실행하지 않는다.',
 'if(*value < low) *value=high;라고 잘못 쓰면 낮은 값이 최고 경계로 튄다. 하한과 상한을 각각 대응시킨다.')
add('1.33','c','const 포인터 매개변수로 읽기 전용 의도 전달',r'''
#include <stdio.h>
static int twice(const int *source)
{
    return *source * 2;
}
int main(void)
{
    int original = 11;
    int answer = twice(&original);
    printf("original=%d twice=%d\n", original, answer);
    return 0;
}''',r'''
#include <stdio.h>
static int triple(const int *source)
{
    return *source * 3;
}
int main(void)
{
    int original = 11;
    printf("original=%d triple=%d\n", original, triple(&original));
    return 0;
}''','함수 매개변수가 const int *라면 원본 original도 반드시 const로 선언해야 하는가?',
 ['const int *source는 source를 통해 가리키는 int를 수정하지 않겠다는 형식상 약속이다.','return *source * 2;는 값을 읽어 계산하고 결과 정수 하나를 반환한다.','호출자 original은 일반 int여도 주소를 const int *에 전달할 수 있다. 이 함수는 원본을 바꾸지 않는다.','const int *source와 int *const source는 의미가 다르다. 후자는 포인터 변수 자체를 다른 주소로 바꿀 수 없다는 뜻이며 별도 심화에서 다룬다.'],
 'original=11 → source→original → *source=11 읽기 → 22 반환 → original은 여전히 11.',
 '“const int *로 전달했으니 original 객체 자체가 영원히 상수”는 틀리다. 이 포인터 경로를 통한 수정이 금지되는 것이다.',
 '두 배 함수를 세 배 함수로 바꾸되 읽기 전용 규칙을 유지하자.',
 'original=11 triple=33. const int *source는 여전히 원본을 읽기만 한다.',
 '*source = 0;을 함수에 넣으면 const int를 통해 수정하려는 제약 위반이다. 읽기 전용 함수에서는 대입을 제거한다.')
add('1.33','d','출력 매개변수로 결과 두 개 만들기',r'''
#include <stdio.h>
static int quotient_remainder(int value, int divisor, int *q, int *r)
{
    if (divisor == 0) {
        return 0;
    }
    *q = value / divisor;
    *r = value % divisor;
    return 1;
}
int main(void)
{
    int q = 0;
    int r = 0;
    if (quotient_remainder(17, 5, &q, &r)) {
        printf("q=%d r=%d\n", q, r);
    }
    return 0;
}''',r'''
#include <stdio.h>
static int quotient_remainder(int value, int divisor, int *q, int *r)
{
    if (divisor == 0) {
        return 0;
    }
    *q = value / divisor;
    *r = value % divisor;
    return 1;
}
int main(void)
{
    int q = 0;
    int r = 0;
    if (quotient_remainder(29, 6, &q, &r)) {
        printf("q=%d r=%d\n", q, r);
    }
    return 0;
}''','함수 반환값 1과 q,r 출력값의 역할은 각각 무엇일까? 17/5의 몫과 나머지를 먼저 계산하자.',
 ['divisor==0은 0으로 나누는 정의되지 않은 동작을 막기 위한 선행 검사다.','q와 r은 호출자 지역 변수의 주소를 받는 출력 매개변수다. *q와 *r에 결과를 저장한다.','return 1은 계산 성공 여부를 전달한다. 결과값 두 개는 포인터 매개변수로 돌려준다.','호출자가 유효한 q,r 주소를 넘긴다는 전제 아래 동작한다. 널 포인터 검사는 심화 API 설계에서 추가한다.'],
 'divisor=5는 0이 아님. 17/5=3, 17%5=2. *q=3, *r=2, 성공 1 반환. main의 if가 참이 되어 출력.',
 '“return이 두 개 필요하다”는 틀리다. 함수는 하나의 상태값을 반환하면서 별도 출력 매개변수로 여러 결과를 제공할 수 있다.',
 '값 29, 나누는 수 6을 전달하도록 바꾸어 몫과 나머지를 구하자.',
 '29/6=4, 29%6=5이므로 q=4 r=5.',
 'divisor==0 검사 없이 나눗셈부터 실행하면 오류다. 0 검사를 먼저 수행해야 한다.')
add('1.34','a','배열 첫 원소를 가리키는 포인터',r'''
#include <stdio.h>
int main(void)
{
    int values[4] = {2, 4, 6, 8};
    int *p = values;
    printf("first=%d third=%d\n", *p, *(p + 2));
    printf("indexed=%d\n", p[3]);
    return 0;
}''',r'''
#include <stdio.h>
int main(void)
{
    int values[4] = {2, 4, 6, 8};
    int *p = values;
    printf("second=%d fourth=%d\n", *(p + 1), p[3]);
    return 0;
}''','p가 첫 원소를 가리킨다면 p+2는 몇 번째 원소를 가리킬까? p[3]의 값도 예상하자.',
 ['대부분의 식에서 values는 첫 원소 values[0]을 가리키는 포인터로 변환된다. int *p=values는 그 주소를 저장한다.','p+2는 바이트 2개가 아니라 int 원소 두 칸만큼 이동한 위치를 가리킨다.','*(p+2)는 values[2]이며 p[3]은 values[3]과 같은 원소다.','배열 자체와 포인터 변수는 동일하지 않다. 배열은 네 정수의 저장 공간이고 p는 별도의 주소 저장 변수다.'],
 'values={2,4,6,8}. p→values[0]=2. p+2→values[2]=6. p[3]=8.',
 '“p+2는 주소를 2바이트 더한다”는 일반적으로 틀리다. C의 포인터 산술은 가리키는 형식의 원소 단위다.',
 '두 번째와 네 번째 원소를 *(p+1), p[3]으로 읽어 보자.',
 'second=4 fourth=8. 두 표현 모두 배열의 유효한 인덱스를 사용한다.',
 '*(p+4)는 배열 끝 다음 위치를 역참조하므로 정의되지 않은 동작이다. 유효한 인덱스 0~3만 읽는다.')
add('1.34','b','포인터로 순회하며 합계 구하기',r'''
#include <stdio.h>
int main(void)
{
    int values[4] = {3, 1, 4, 2};
    int *current = values;
    int *end = values + 4;
    int sum = 0;
    while (current < end) {
        sum += *current;
        ++current;
    }
    printf("sum=%d\n", sum);
    return 0;
}''',r'''
#include <stdio.h>
int main(void)
{
    int values[4] = {3, 1, 4, 2};
    int *current = values;
    int *end = values + 4;
    int count = 0;
    while (current < end) {
        if (*current >= 3) {
            ++count;
        }
        ++current;
    }
    printf("count=%d\n", count);
    return 0;
}''','while이 몇 번 반복되는지, end가 가리키는 원소를 읽어도 되는지 예상하자.',
 ['end=values+4는 마지막 원소 바로 다음(one-past) 위치를 가리킨다. 포인터 값 자체를 만들고 비교할 수 있지만 *end는 허용되지 않는다.','current<end는 같은 배열 안과 끝 다음 위치를 기준으로 비교하므로 이 코드에서 정의된 동작이다.','sum+=*current로 현재 원소를 더한 뒤 ++current로 다음 원소를 가리킨다.','마지막 원소를 더한 뒤 current==end가 되므로 반복 조건이 거짓이 된다.'],
 'sum=0 → 첫 원소 3 더해 3 → 1 더해 4 → 4 더해 8 → 2 더해 10 → current=end, 종료.',
 '“end가 마지막 다음을 가리키므로 *end도 읽을 수 있다”는 틀리다. 위치 계산과 그 위치의 역참조는 다른 행위다.',
 '합계 대신 값이 3 이상인 원소 수를 세어 보자. 현재 원소를 검사하고 포인터를 이동하는 구조는 유지한다.',
 '3과 4 두 원소만 조건을 만족하므로 count=2.',
 '++current를 빼면 current가 이동하지 않아 반복이 끝나지 않는다. 매 반복에서 진행 상태가 변하는지 확인한다.')
add('1.34','c','포인터 차이로 위치 알아내기',r'''
#include <stdio.h>
#include <stddef.h>
int main(void)
{
    int items[5] = {10, 20, 30, 40, 50};
    int *begin = items;
    int *chosen = &items[3];
    ptrdiff_t index = chosen - begin;
    printf("index=%td value=%d\n", index, *chosen);
    return 0;
}''',r'''
#include <stdio.h>
#include <stddef.h>
int main(void)
{
    int items[5] = {10, 20, 30, 40, 50};
    int *begin = items;
    int *chosen = &items[4];
    ptrdiff_t index = chosen - begin;
    printf("index=%td value=%d\n", index, *chosen);
    return 0;
}''','&items[3]에서 items까지의 차이는 3바이트인가, 3원소인가?',
 ['두 포인터가 같은 배열을 가리킬 때 뺄셈 결과는 원소 개수 단위의 거리다.','ptrdiff_t는 포인터 차이를 표현하기 위한 표준 부호 있는 정수 형식이며 <stddef.h>에 정의된다.','%td는 ptrdiff_t 값을 출력하기 위한 C99 이후의 printf 변환 지정자다. C11에서 사용할 수 있다.','chosen-begin은 3이며 chosen은 items[3]=40을 가리킨다.'],
 'begin→items[0], chosen→items[3]. 두 위치의 원소 간격 3. *chosen=40.',
 '“서로 무관한 두 배열의 포인터도 빼서 거리를 계산할 수 있다”는 틀리다. 유효한 포인터 차이는 같은 배열 객체의 범위 규칙을 따른다.',
 'chosen을 마지막 원소 &items[4]로 바꾸어 차이와 값을 예상하자.',
 'index=4 value=50. 원소 위치를 나타내는 차이는 0부터 시작한다.',
 'printf("%d", index)는 ptrdiff_t에 맞는 지정자가 아니다. %td를 사용한다.')
add('1.34','d','sizeof 배열과 포인터를 구별하는 안전한 길이 계산',r'''
#include <stdio.h>
#include <stddef.h>
int main(void)
{
    int samples[5] = {1, 2, 3, 4, 5};
    size_t count = sizeof samples / sizeof samples[0];
    int *cursor = samples;
    int *end = samples + count;
    int sum = 0;
    while (cursor != end) {
        sum += *cursor;
        ++cursor;
    }
    printf("count=%zu sum=%d\n", count, sum);
    return 0;
}''',r'''
#include <stdio.h>
#include <stddef.h>
int main(void)
{
    int samples[5] = {1, 2, 3, 4, 5};
    size_t count = sizeof samples / sizeof samples[0];
    int *cursor = samples;
    int *end = samples + count;
    int even = 0;
    while (cursor != end) {
        if (*cursor % 2 == 0) {
            ++even;
        }
        ++cursor;
    }
    printf("count=%zu even=%d\n", count, even);
    return 0;
}''','count는 sizeof samples / sizeof samples[0]으로 왜 5가 되는가? sizeof cursor로 바꾸면 왜 문제가 생길까?',
 ['sizeof samples는 배열 객체 전체의 바이트 크기다. sizeof samples[0]은 정수 원소 한 개의 바이트 크기다. 나누면 원소 수 5가 된다.','size_t는 객체 크기와 원소 개수를 표현하는 부호 없는 정수 형식이다. %zu는 size_t 출력 지정자다.','cursor는 별도 포인터 변수이므로 sizeof cursor는 포인터 변수의 바이트 크기다. 이를 배열 전체 크기로 착각하면 안 된다.','이 계산법은 배열 선언이 보이는 현재 블록에서 사용한다. 함수 매개변수 int values[]는 실제로 포인터 매개변수이므로 sizeof로 원소 수를 얻을 수 없다.'],
 'count=5. cursor는 원소 1부터 5까지 순회하며 sum이 1→3→6→10→15가 된다. 끝에서 cursor==end로 종료.',
 '“함수 인수로 받은 int arr[]에도 sizeof arr / sizeof arr[0]을 적용하면 길이를 알 수 있다”는 틀리다. 함수 매개변수에서는 배열이 포인터로 조정된다.',
 '합계 대신 짝수 원소 수를 세는 독립 코드를 작성하자. 개수 계산과 포인터 범위는 그대로 유지한다.',
 '짝수는 2,4 두 개이므로 count=5 even=2.',
 'count=sizeof cursor / sizeof cursor[0]은 배열 원소 수 계산이 아니다. 실제 배열 이름 samples를 사용해야 한다.')

# Verify all full programs with both compilers, strict C11, deterministic outputs.
verified=[]; failures=[]
for e in E:
    for kind,code in [('main',e['code']),('variant',e['variant'])]:
        folder=ROOT/'sources'/(e['sec'].replace('.','_')+'_'+e['id'])/kind
        folder.mkdir(parents=True,exist_ok=True)
        src=folder/'main.c'; src.write_text(code,encoding='utf-8')
        outputs=[]
        for cc in ('gcc','clang'):
            exe=folder/f'program_{cc}'
            comp=subprocess.run([cc,*FLAGS,str(src),'-o',str(exe)],capture_output=True,text=True,timeout=15)
            if comp.returncode:
                failures.append({'id':e['id'],'kind':kind,'compiler':cc,'stderr':comp.stderr})
                continue
            run=subprocess.run([str(exe)],capture_output=True,text=True,timeout=15)
            if run.returncode:
                failures.append({'id':e['id'],'kind':kind,'compiler':cc,'stderr':run.stderr,'returncode':run.returncode})
            outputs.append({'compiler':cc,'stdout':run.stdout,'returncode':run.returncode})
        if len(outputs)==2 and outputs[0]['stdout']!=outputs[1]['stdout']:
            failures.append({'id':e['id'],'kind':kind,'mismatch':outputs})
        verified.append({'id':e['id'],'kind':kind,'source':str(src.relative_to(ROOT)),'sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'outputs':outputs})
        e[kind+'_output']=outputs[0]['stdout'] if outputs else '미검증'
if failures:
    raise RuntimeError(json.dumps(failures,ensure_ascii=False,indent=2))

D=Document(); sec=D.sections[0]; sec.top_margin=Cm(2.15);sec.bottom_margin=Cm(2.05);sec.left_margin=Cm(2.45);sec.right_margin=Cm(2.45)
styles=D.styles
base=styles['Normal'];base.font.name='Malgun Gothic';base.font.size=Pt(10);base.paragraph_format.space_after=Pt(6);base.paragraph_format.line_spacing=1.23
for name,size,color in [('Title',19,RGBColor(0,91,92)),('Heading 1',16,RGBColor(0,108,109)),('Heading 2',12,RGBColor(0,92,92)),('Heading 3',10.5,RGBColor(0,100,100))]:
 st=styles[name];st.font.name='Malgun Gothic';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=color;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
if 'Code C' not in styles:
 s=styles.add_style('Code C',WD_STYLE_TYPE.PARAGRAPH)
else:s=styles['Code C']
s.font.name='Consolas';s.font.size=Pt(8.0);s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(0);s.paragraph_format.line_spacing=1.05
s.paragraph_format.keep_together=True
if 'Small Note' not in styles:
 sn=styles.add_style('Small Note',WD_STYLE_TYPE.PARAGRAPH)
else:sn=styles['Small Note']
sn.font.name='Malgun Gothic';sn.font.size=Pt(9.0);sn.font.color.rgb=RGBColor(62,76,84)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('C언어 종합교과서 · 이해 중심 정식 개정판 · 제1장')

def para(t,style=None):
 p=D.add_paragraph(t,style=style);p.paragraph_format.widow_control=True;return p

def codeblock(code):
 lines=code.rstrip('\n').splitlines()
 for i,line in enumerate(lines):
  p=D.add_paragraph(style='Code C');p.add_run(line or ' ')
  p.paragraph_format.keep_with_next=(i<len(lines)-1)
 return len(lines)

def h(t,level=2):D.add_heading(t,level=level)

D.add_heading('C언어 종합교과서',0)
para('이해 중심 정식 개정판 · 제1장 연속 원고 · 1.32~1.34절', 'Small Note')
para('이번 원고의 목표: 포인터를 단순히 주소 숫자로 외우지 않고, 객체·수명·형식·범위 규칙에 연결하여 안전하게 사용한다. 전체 예제와 변형 답안은 모두 독립 실행 가능한 C11 프로그램이다.')
para('학습 경로: 1.32 주소와 역참조 → 1.33 포인터 매개변수 → 1.34 배열 위치와 범위. 1.23~1.25의 지역 객체·저장 기간, 1.17의 배열, 1.18~1.19의 함수, 1.29~1.31의 전처리 기초를 누적 복습한다.')
para('표준 범위: 여기의 모든 실행 코드는 ISO C11만 사용한다. C23 전용 문법, POSIX/Windows API, 특정 컴파일러 확장은 사용하지 않는다. 포인터의 실제 숫자 주소와 sizeof 포인터의 크기는 환경마다 다를 수 있어 정답 출력에 사용하지 않는다.','Small Note')
for s in sections:
 D.add_heading(s['number']+' '+s['title'],level=1)
 h('왜 배워야 하는가 — 사용 상황과 해결할 문제',2);para(s['why']);para(s['before'])
 h('이번 절에서 직접 확인할 목표',2)
 for g in s['goals']:para('• '+g)
 for idx,e in enumerate([z for z in E if z['sec']==s['number']],1):
  h('예제 '+e['id']+' · '+e['title'],2)
  para('상황과 하위 목표: '+('객체를 준비한다 → 주소/포인터의 관계를 확인한다 → 실행 중 바뀌는 값을 추적한다 → 유효한 범위에서 결과를 출력한다.' if s['number']=='1.32' else '입력과 결과를 구분한다 → 함수/배열의 경계를 정한다 → 안전 조건을 검사한다 → 실제 값으로 결과를 확인한다.'))
  h('먼저 실행을 예측해 보기',3);para(e['ask'])
  h('완전한 C11 프로그램',3);codeblock(e['code'])
  h('코드를 문장 단위로 이해하기',3)
  for j,note in enumerate(e['explain'],1):para(f'{j}. {note}')
  h('실행 흐름과 값의 변화',3);para(e['trace'])
  h('실제 출력 — GCC·Clang 확인',3);codeblock(e['main_output'])
  para('출력 확인용 printf/puts는 학습자가 상태를 관찰하도록 돕는 문장이다. 주소 계산, 역참조, 값 수정 또는 반복 조건이 이 예제의 핵심 로직이다. 출력 문장을 지워도 해당 핵심 계산의 의미는 변하지 않는다.','Small Note')
  h('예제를 바꾸며 실력 다지기 — 바로 풀고 확인하기',3);para(e['task'])
  para('도움말: 기존 코드를 먼저 복사한 다음, 바뀌는 조건·값·연산만 찾아 표시한다. 앞선 예제의 동작을 이해했다면 다음에는 코드를 보지 않고 다시 작성해 보자.','Small Note')
  para('변형 정답:');codeblock(e['variant'])
  para('변형 코드 실제 출력:');codeblock(e['variant_output'])
  para('왜 이 답인가: '+e['solution'])
  h('오개념 점검과 오류 수정',3);para('그럴듯하지만 틀린 생각: '+e['misconception']);para('디버깅: '+e['debug'])
  if idx==2:
   h('중간 인출 복습 — 코드를 가리고 설명하기',3)
   para('지금까지 본 예제의 출력만 기억하지 말고, ① 어떤 객체가 만들어지는가 ② 포인터 변수에는 무엇이 저장되는가 ③ 역참조하면 어느 객체가 읽히는가 ④ 함수나 반복문이 끝날 때 값은 무엇인가를 순서대로 종이에 적어 보자. 답의 근거는 바로 앞 두 예제의 실행 추적에 있다.')
  if idx==4:
   h('도움 없는 독립 작성과 누적 복습',3)
   if s['number']=='1.32':
    para('문제: int level=5를 만들고 포인터 하나로 level을 8로 수정한 뒤 출력하는 프로그램을, 위의 코드를 보지 않고 작성하자. 앞서 배운 변수·대입·printf도 함께 사용한다.')
    para('모범 해설: int level=5; int *p=&level; *p=8; printf("%d\\n",level);의 순서가 핵심이다. 주소를 받기 전에 포인터를 역참조하지 않는다.')
   elif s['number']=='1.33':
    para('문제: int point=6의 주소를 받아 값을 1 증가시키는 static void increase(int *p) 함수를 작성하고, main에서 호출 전후 값을 출력하자. 함수의 매개변수는 포인터 값으로 전달된다는 점도 말로 설명하자.')
    para('모범 해설: 함수 안에서 ++*p;를 수행하고 main에서 increase(&point);를 호출하면 point=7이다. ++*p는 포인터가 가리키는 정수에 1을 더한다.')
   else:
    para('문제: int data[3]={5,7,9}를 포인터로 순회해 7보다 큰 원소의 수를 출력하자. end=data+3을 만들어 end를 역참조하지 않는 조건을 사용한다.')
    para('모범 해설: for(int *p=data; p<data+3; ++p) if(*p>7) ++count;와 같은 구조로 count=1이 된다. 포인터는 같은 배열 범위 안에서만 비교한다.')
 h('이번 절의 기억할 규칙과 다음 절로의 연결',2)
 para('포인터 변수는 주소 값을 담고, 역참조는 그 주소가 가리키는 유효한 객체를 읽거나 수정한다. 객체의 수명·형식·배열 경계를 지키는 것이 문법 암기보다 중요하다. 이후에는 문자 포인터, 문자열 함수, 동적 메모리 할당, 구조체 포인터, 함수 포인터로 범위를 넓힌다.')

D.add_heading('이번 묶음의 종합 확인',level=1)
para('① &value와 *pointer의 역할을 서로 다른 문장으로 설명해 보자. ② 함수에 int 값을 넘기는 것과 int *를 넘기는 것 모두 C에서 값 전달인 이유는 무엇인가? ③ 같은 배열의 마지막 다음 포인터를 만들 수 있지만 역참조할 수 없는 이유는 무엇인가? ④ sizeof 배열과 sizeof 포인터를 혼동하면 길이 계산에서 어떤 문제가 생기는가?')
para('정답과 근거: ① &는 객체 주소를 얻고 *는 유효한 포인터가 가리키는 객체에 접근한다. ② 포인터 매개변수에도 주소 값의 복사본이 전달된다. 다만 그 주소를 통해 호출자 객체에 접근할 수 있다. ③ one-past 위치는 원소가 아니므로 읽을 객체가 없다. ④ 배열의 전체 바이트 크기와 포인터 변수의 바이트 크기는 의미가 달라 원소 수가 잘못 계산된다.')
para('표준과 범위 대조 주의: K&R 제5장의 포인터·배열 학습목표, Express의 포인터·배열·함수 학습목표를 독창적 예제로 부분 반영했다. Express 17장 전체와 K&R 8장 및 부록 A~C의 모든 세부 LAB/Exercise/Programming/Mini Project/Advanced Topic 전수 대조는 미완료이며, 확인하지 않은 항목은 완료 처리하지 않는다.','Small Note')

DOCX=ROOT/'C언어_종합교과서_이해중심_정식개정판_제1장_1.32-1.34.docx';D.save(DOCX)
ZIP=ROOT/'official_ch01_1_32-1_34_sources.zip'
with zipfile.ZipFile(ZIP,'w',zipfile.ZIP_DEFLATED) as z:
 for e in E:
  for kind in ('main','variant'):
   p=ROOT/'sources'/(e['sec'].replace('.','_')+'_'+e['id'])/kind/'main.c'
   z.write(p,str(p.relative_to(ROOT)))
verify={'standard':'ISO C11','flags':FLAGS,'compilers':['gcc','clang'],'program_count':len(verified),'compile_count':len(verified)*2,'execution_count':len(verified)*2,'all_passed':not failures,'results':verified,'failures':failures,'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest(),'code_docx_consistency':'Each program and variant inserted from same E source used for compilation','source_global_objects':False}
(ROOT/'official_ch01_1_32-1_34_c11_verification.json').write_text(json.dumps(verify,ensure_ascii=False,indent=2),encoding='utf-8')
progress={'last_verified_section':'1.34','completed_section_range_reported':'1.1-1.34','new_sections':['1.32','1.33','1.34'],'overall_progress_percent':None,'full_source_audit':'incomplete','final_integrated_docx':'incomplete','github_delivery':'pending_verification'}
(ROOT/'official_rewrite_progress_1_1-1_34.json').write_text(json.dumps(progress,ensure_ascii=False,indent=2),encoding='utf-8')
tracking={'coverage_audit_status':'not_exhaustively_audited','original_sources':{'Express':'c언어 압축 (1)(1).pdf — full source per-item audit pending','K&R':'C_Programming.pdf — full source per-item audit pending','K&R OCR':'C Programming Language - 2nd Edition (OCR).pdf — supplementary; full audit pending'},'new_items':[{'section':'1.32','topics':['address-of','indirection','pointer copying','null pointer'],'status':'original textbook content verified; source-to-item audit pending'},{'section':'1.33','topics':['pointer arguments','output parameters','const pointee','value passing'],'status':'original textbook content verified; source-to-item audit pending'},{'section':'1.34','topics':['array-to-pointer conversion','pointer arithmetic','one-past','ptrdiff_t','sizeof array'],'status':'original textbook content verified; source-to-item audit pending'}]}
(ROOT/'official_rewrite_objective_tracking_1_1-1_34.json').write_text(json.dumps(tracking,ensure_ascii=False,indent=2),encoding='utf-8')
manifest={'run':'2026-10-15 ch01 1.32-1.34','sections':['1.32','1.33','1.34'],'docx':DOCX.name,'source_zip':ZIP.name,'verification':'official_ch01_1_32-1_34_c11_verification.json','build_script':Path(__file__).name,'progress':'official_rewrite_progress_1_1-1_34.json','objective_tracking':'official_rewrite_objective_tracking_1_1-1_34.json','verified_programs':len(verified),'verified_compiles':len(verified)*2,'sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [DOCX,ZIP]}}
(ROOT/'official_rewrite_manifest_1_32-1_34.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
# GitHub transfer text is base64 encoded, not deliverable content itself.
import base64,textwrap
for p in [DOCX,ZIP]:
 b=base64.b64encode(p.read_bytes()).decode('ascii')
 (ROOT/('_transfer_'+p.name+'.b64.txt')).write_text('\n'.join(textwrap.wrap(b,240))+'\n',encoding='ascii')
print(json.dumps({'docx':str(DOCX),'programs':len(verified),'compiles':len(verified)*2,'zip':str(ZIP),'files':[p.name for p in ROOT.iterdir() if p.is_file()]},ensure_ascii=False))