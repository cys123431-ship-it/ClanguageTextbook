#!/usr/bin/env python3
"""C 언어 정식 개정판 1.65–1.67 독립 집필·컴파일 검증·DOCX 제작."""
from __future__ import annotations
import os, sys, json, re, hashlib, subprocess, zipfile, textwrap
from pathlib import Path
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
TAG='1_65-1_67'
FLAGS=['-std=c11','-Wall','-Wextra','-Wpedantic','-Werror','-O0']
KST=timezone(timedelta(hours=9))
NOW=datetime.now(KST).isoformat(timespec='seconds')
CODES=OUT/'sources'; CODES.mkdir(exist_ok=True)

SECTIONS=[
{
'id':'1.65','title':'부동소수점 환경과 예외 플래그 — <fenv.h>',
'unit':'22: 수학·실수·복소수와 환경',
'lead':[
'실수 계산에서 결과값만 확인하면 놓치는 정보가 있다. 어떤 연산이 표현 가능한 범위를 넘었는지, 유효하지 않은 수학적 상황을 만났는지, 어떤 반올림 규칙을 적용했는지는 계산의 신뢰도를 바꾼다. C11의 <fenv.h>는 구현이 지원하는 부동소수점 환경의 상태를 읽고, 예외 플래그를 검사하며, 반올림 방향을 제어하는 인터페이스를 제공한다.',
'여기서 “예외”는 C++의 throw처럼 자동으로 함수 밖으로 탈출하는 사건이 아니다. 많은 환경에서 FE_INVALID 같은 플래그는 상태를 기록하는 표시이며, 반드시 프로그램을 종료시키는 것은 아니다. 실제 하드웨어 트랩이나 수치 결과는 구현·컴파일 옵션·연산 종류에 좌우된다. 이 절에서는 특정 수식의 최적화에 의존하지 않도록 표준 함수로 플래그를 직접 올려 관찰한다.',
'FE_INVALID, FE_DIVBYZERO 등의 개별 매크로는 구현이 지원할 때만 정의될 수 있다. 따라서 전처리기 #ifdef로 지원 여부를 확인한다. 예제의 컴파일 검증에는 수학 라이브러리 연결을 위해 -lm을 추가한다. C11 자체는 특정 운영체제의 실수 하드웨어 구조를 보장하지 않는다.'
],
'concepts':[
('부동소수점 환경', '반올림 모드와 예외 플래그 같은 실행 상태의 묶음이다. 정수 errno와 별개이며, 일반적인 함수 호출 성공/실패와도 다르다.'),
('feclearexcept / fetestexcept', '첫 함수는 지정된 예외 플래그를 지우고, 둘째 함수는 어떤 플래그가 켜졌는지 비트 마스크로 돌려준다. 0이 아닌 반환값은 해당 비트가 포함되었다는 뜻이다.'),
('feraiseexcept', '지원되는 예외 플래그를 프로그램이 명시적으로 올린다. 반환값 0은 성공이다. 이것만으로 실제 나눗셈이 일어났다고 주장할 수 없다.'),
('fegetround / fesetround', '현재 반올림 모드를 읽거나 바꾼다. 설정 후에는 반드시 원래 모드로 복원해야 다른 계산에 영향을 남기지 않는다.'),
('fenv_t, feholdexcept, feupdateenv', '환경을 저장하고 예외 상태를 잠시 보류했다가 복구·병합하는 데 사용한다. 모든 구현에서 동일한 하드웨어 효과를 가정하지 않는다.')
],
'examples':[
{
'key':'fenv_flag','title':'예외 플래그를 명시적으로 올리고 검사하기',
'why':'수치 함수가 실패 가능성을 상태 플래그로 알릴 때, “연산 뒤에 무엇을 검사해야 하는가”를 연습한다. 상태를 지우지 않으면 과거 계산에서 남은 플래그를 이번 계산의 문제로 오인할 수 있다.',
'code':r'''#include <fenv.h>
#include <stdio.h>

int main(void)
{
    if (feclearexcept(FE_ALL_EXCEPT) != 0) {
        puts("clear failed");
        return 1;
    }
#ifdef FE_INVALID
    if (feraiseexcept(FE_INVALID) != 0) {
        puts("raise failed");
        return 1;
    }
    printf("invalid=%d\n", fetestexcept(FE_INVALID) != 0);
#else
    puts("FE_INVALID unavailable");
#endif
    return 0;
}''',
'predict':'FE_INVALID를 지원하는 구현에서 화면에 출력되는 0 또는 1을 먼저 예상하자. 이전 계산에서 남은 플래그가 있으면 어떤 혼란이 생길까?',
'walk':['<fenv.h>는 플래그 관련 선언을 제공하고 <stdio.h>는 puts/printf 선언을 제공한다. 두 헤더의 역할은 다르다.','feclearexcept(FE_ALL_EXCEPT)는 지원되는 예외 플래그 전체를 초기화한다. 반환값이 0이 아니면 검증을 계속하지 않는다.','#ifdef FE_INVALID는 런타임 if가 아니라 번역 전에 처리되는 조건부 컴파일이다. 매크로가 없으면 그 블록 자체를 컴파일하지 않는다.','feraiseexcept(FE_INVALID)는 특정 플래그를 올린다. fetestexcept는 비트 마스크를 반환하므로 != 0으로 참/거짓을 출력한다.','printf는 관찰을 위한 학습용 출력이다. 플래그를 초기화하고 검사하는 부분이 실제 기능이다.'],
'flow':['시작: 예외 플래그 상태는 알 수 없으므로 초기화한다.','지원되는 구현: FE_INVALID 플래그를 1로 만든다.','검사: fetestexcept(FE_INVALID) != 0 → 참(1).','출력: invalid=1. 미지원 구현은 별도 안내 문구를 출력한다.'],
'pitfall':'“FE_INVALID가 켜졌으니 방금 0/0을 계산했다”는 결론은 틀리다. 이 예제는 실제 계산 없이 feraiseexcept로 플래그를 올렸다.',
'debug':'feclearexcept를 제거한 뒤 다른 함수 호출에서 플래그가 남았을 가능성을 설명하자. 왜 검사 전 초기화가 테스트 재현성을 높이는가?',
'variation':r'''#include <fenv.h>
#include <stdio.h>

int main(void)
{
    if (feclearexcept(FE_ALL_EXCEPT) != 0) {
        return 1;
    }
#ifdef FE_DIVBYZERO
    if (feraiseexcept(FE_DIVBYZERO) != 0) {
        return 1;
    }
    printf("divzero=%d\n", fetestexcept(FE_DIVBYZERO) != 0);
#else
    puts("FE_DIVBYZERO unavailable");
#endif
    return 0;
}''',
'variation_task':'FE_INVALID 대신 지원 여부를 확인한 FE_DIVBYZERO를 올려 검사한다. 기존 예제에서 수정해야 할 상수와 출력 문구를 먼저 찾아 적은 뒤 코드를 작성하자.',
'variation_reason':'검사할 플래그와 올릴 플래그를 같은 매크로로 바꿔야 한다. 하나만 바꾸면 다른 상태를 읽는 논리 오류가 생긴다.'
},
{
'key':'fenv_round','title':'반올림 모드를 임시 변경하고 반드시 복원하기',
'why':'금융·측정 프로그램의 수치 실험에서는 계산 전후의 환경이 같아야 비교가 가능하다. 원래 모드를 저장하지 않고 변경하면 다른 함수의 결과에도 영향을 줄 수 있다.',
'code':r'''#include <fenv.h>
#include <stdio.h>

int main(void)
{
    int old_mode = fegetround();
    if (old_mode == -1) {
        puts("round mode unavailable");
        return 1;
    }
    if (fesetround(FE_TONEAREST) != 0) {
        puts("set failed");
        return 1;
    }
    printf("nearest=%d\n", fegetround() == FE_TONEAREST);
    if (fesetround(old_mode) != 0) {
        puts("restore failed");
        return 1;
    }
    puts("restored");
    return 0;
}''',
'predict':'현재 모드가 처음부터 FE_TONEAREST였어도 “복원”이 필요한가? nearest= 뒤에는 무엇이 나올까?',
'walk':['fegetround는 현재 반올림 방향을 정수 코드로 돌려주며 실패 시 -1을 반환할 수 있다.','old_mode는 이전 상태를 기억하는 지역 변수다. 전역 객체를 사용하지 않는다.','fesetround(FE_TONEAREST)가 성공하면 fegetround와 비교해 1을 출력한다.','마지막 fesetround(old_mode)는 이 예제의 핵심 자원 복원 단계다. 출력보다 중요하다.'],
'flow':['old_mode에 실행 전 반올림 모드를 저장한다.','현재 모드를 FE_TONEAREST로 바꾼다.','비교 결과가 참이므로 nearest=1을 출력한다.','old_mode로 복원하고 restored를 출력한다.'],
'pitfall':'“반올림 모드를 바꿨으니 모든 상수 계산이 즉시 바뀐다”는 오개념을 버리자. 컴파일 시점 상수 접기, 최적화, 환경 접근 지원에 따라 계산 관찰은 달라질 수 있다.',
'debug':'복원 문장을 지운 프로그램을 상상하자. 같은 프로세스에서 뒤이어 실행하는 다른 수치 함수에 어떤 숨은 의존성이 생기는가?',
'variation':r'''#include <fenv.h>
#include <stdio.h>

int main(void)
{
    int before = fegetround();
    if (before == -1) {
        return 1;
    }
#ifdef FE_DOWNWARD
    if (fesetround(FE_DOWNWARD) != 0) {
        return 1;
    }
    printf("downward=%d\n", fegetround() == FE_DOWNWARD);
#else
    puts("FE_DOWNWARD unavailable");
#endif
    if (fesetround(before) != 0) {
        return 1;
    }
    puts("mode restored");
    return 0;
}''',
'variation_task':'부분 완성 과제: 기존 old_mode를 before로 이름 바꾸고, FE_DOWNWARD 지원 여부를 검사하는 조건부 컴파일을 추가한 뒤 반드시 이전 모드로 복원하자.',
'variation_reason':'지원 여부가 불명확한 상수는 #ifdef로 보호한다. 분기 밖에서 복원하면 지원/미지원 어느 경로에서도 상태 관리 의도를 유지할 수 있다.'
},
{
'key':'fenv_hold','title':'예외 상태를 잠시 보류하고 병합하기',
'why':'여러 단계의 수치 계산에서 임시 구간이 올린 예외를 호출자에게 전달하려면 이전 환경을 저장했다가 합쳐야 한다. 단순한 초기화와 병합은 다른 동작이다.',
'code':r'''#include <fenv.h>
#include <stdio.h>

int main(void)
{
#ifdef FE_INVALID
    fenv_t saved;
    if (feclearexcept(FE_ALL_EXCEPT) != 0) {
        return 1;
    }
    if (feholdexcept(&saved) != 0) {
        return 1;
    }
    if (feraiseexcept(FE_INVALID) != 0) {
        return 1;
    }
    if (feupdateenv(&saved) != 0) {
        return 1;
    }
    printf("merged=%d\n", fetestexcept(FE_INVALID) != 0);
#else
    puts("FE_INVALID unavailable");
#endif
    return 0;
}''',
'predict':'feholdexcept로 저장한 뒤 FE_INVALID를 올리고 feupdateenv를 호출하면 플래그는 사라질까, 남을까?',
'walk':['fenv_t saved는 환경 상태를 담는 객체이며 구체적인 내부 구조는 구현에 맡겨진다.','feholdexcept는 현재 환경을 저장하고 예외 플래그를 지우며 가능한 경우 예외 트랩을 보류하는 모드로 전환한다.','feraiseexcept는 임시 구간에서 FE_INVALID를 올린다.','feupdateenv는 저장한 환경을 복원한 뒤 현재까지 발생한 예외를 다시 올리는 의미를 가진다. 따라서 이번 플래그가 병합된다.'],
'flow':['초기화: 예외 플래그를 비운다.','저장: saved에 환경을 보관하고 임시 구간으로 진입한다.','임시 구간: FE_INVALID를 올린다.','병합: 이전 환경을 복원하면서 발생한 예외를 다시 반영한다.','검사: merged=1을 출력한다.'],
'pitfall':'“feupdateenv는 저장 상태로 완전히 덮어쓰므로 새 예외는 무조건 지워진다”는 잘못된 이해다. fesetenv와 feupdateenv의 목적을 구별하자.',
'debug':'feupdateenv를 fesetenv로 바꾸면 어떤 의미 차이가 생길까? 아래 변형 코드에서 결과를 확인하자.',
'variation':r'''#include <fenv.h>
#include <stdio.h>

int main(void)
{
#ifdef FE_INVALID
    fenv_t saved;
    if (feclearexcept(FE_ALL_EXCEPT) != 0) {
        return 1;
    }
    if (fegetenv(&saved) != 0) {
        return 1;
    }
    if (feraiseexcept(FE_INVALID) != 0) {
        return 1;
    }
    if (fesetenv(&saved) != 0) {
        return 1;
    }
    printf("cleared=%d\n", fetestexcept(FE_INVALID) == 0);
#else
    puts("FE_INVALID unavailable");
#endif
    return 0;
}''',
'variation_task':'feholdexcept/feupdateenv 대신 fegetenv/fesetenv를 사용해 원래 환경으로 되돌리자. 예외가 남는지 사전 예측한 후 실행한다.',
'variation_reason':'fesetenv는 저장된 환경을 설정하는 동작이며 새 예외를 병합하는 feupdateenv와 다르다. cleared=1은 복원된 환경에 FE_INVALID가 없다는 의미다.'
}
]
},
{
'id':'1.66','title':'비지역 제어 흐름과 복구 지점 — <setjmp.h>',
'unit':'25: 오류 진단·제어 이탈·가변 인수',
'lead':[
'함수 호출이 여러 겹 쌓였을 때 깊은 곳에서 치명적인 입력 오류를 발견할 수 있다. 가장 읽기 쉬운 기본 해법은 함수마다 오류 코드를 반환하고 호출자가 처리하는 것이다. 하지만 매우 제한된 상황에서 여러 호출 단계를 건너뛰어 미리 지정한 복구 지점으로 돌아가야 할 때가 있다. <setjmp.h>는 setjmp와 longjmp를 제공한다.',
'이 기능은 일반적인 함수 반환과 다르다. longjmp는 중간 함수를 정상적인 return으로 빠져나오지 않으므로 그 경로에서 수행했어야 할 파일 닫기, 메모리 해제 같은 정리 코드가 실행되지 않을 수 있다. 또한 setjmp 이후 변경된 일부 지역 변수 값은 longjmp 뒤에 믿을 수 없다. 따라서 사용 이유와 제한을 이해한 뒤에만 적용한다.',
'중요한 C 문법 규칙: setjmp 매크로는 아무 표현식 위치에서나 호출할 수 없다. if·switch의 제어식 또는 허용된 비교식 등 표준이 정한 문맥에서만 사용한다. int x = setjmp(env);처럼 초기화식에 놓는 예제는 만들지 않는다.'
],
'concepts':[
('jmp_buf', '실행 환경을 저장하는 데 필요한 정보를 담는 자료형이다. 일반 배열이나 포인터의 내부 배치처럼 해석해서는 안 된다.'),
('setjmp', '처음 실행하면 0을 돌려준다. 같은 환경으로 longjmp가 돌아오면 지정한 0이 아닌 값처럼 다시 평가된다.'),
('longjmp', '활성 상태의 setjmp 복구 지점으로 제어를 옮긴다. 이미 끝난 함수의 복구 지점으로 점프하면 정의되지 않은 동작이다.'),
('지역 변수의 값', 'setjmp 호출을 포함한 함수의 비volatile 자동 지역 변수가 setjmp 뒤에 변경되었다면 longjmp 뒤 그 값은 불확정이다. 의도한 복구 데이터는 별도 설계가 필요하다.'),
('자원 정리', 'longjmp는 중간 함수의 일반 반환을 생략한다. 파일·동적 메모리 등은 복구 경로에서 명시적으로 정리할 수 있게 설계해야 한다.')
],
'examples':[
{
'key':'jump_basic','title':'첫 실행과 복구 실행을 구별하기',
'why':'정상 실행 경로와 복구 경로를 구분하는 가장 작은 예제로, longjmp가 “함수를 한 번 더 호출하는 것”이 아님을 확인한다.',
'code':r'''#include <setjmp.h>
#include <stdio.h>

int main(void)
{
    jmp_buf point;
    if (setjmp(point) == 0) {
        puts("first pass");
        longjmp(point, 5);
    } else {
        puts("recovered");
    }
    puts("finished");
    return 0;
}''',
'predict':'first pass가 두 번 출력될까? longjmp 이후 프로그램은 어느 if 분기로 진입할까?',
'walk':['jmp_buf point는 복구 위치에 필요한 상태를 보관한다. 전역 객체가 아니다.','if (setjmp(point) == 0)은 허용된 setjmp 사용 형태다. 첫 실행에서 참이 된다.','longjmp(point, 5)는 point를 저장한 setjmp 지점으로 되돌아가게 한다. 이때 setjmp는 5처럼 평가된다.','이제 조건식이 거짓이므로 else가 실행된다. 이어서 if문 다음 문장으로 진행한다.'],
'flow':['처음 setjmp: 0 → first pass 출력.','longjmp(point, 5): 복구 지점으로 제어 이동.','복구 setjmp: 5 → else 선택.','recovered, finished 순서로 출력.'],
'pitfall':'“longjmp가 호출되면 main을 처음부터 다시 시작한다”는 틀리다. 지정된 setjmp 지점에서 다시 이어지고, 이미 실행된 문장이 자동으로 재실행되는 것은 아니다.',
'debug':'longjmp를 if 바깥으로 옮겨 반복 호출하면 복구 루프가 생길 수 있다. 제어 흐름을 화살표로 그려 위험을 설명하자.',
'variation':r'''#include <setjmp.h>
#include <stdio.h>

int main(void)
{
    jmp_buf point;
    if (setjmp(point) == 0) {
        puts("begin");
        longjmp(point, 9);
    } else {
        puts("fallback path");
    }
    puts("end");
    return 0;
}''',
'variation_task':'복구값을 5에서 9로 바꾸고 출력 문구도 변경하자. 실행 경로가 달라지는지 설명한다.',
'variation_reason':'0이 아닌 복구값은 여전히 if 조건을 거짓으로 만든다. 값의 크기 자체가 실행 경로를 바꾸지는 않는다.'
},
{
'key':'jump_zero','title':'longjmp의 두 번째 인수가 0이면 어떻게 될까?',
'why':'setjmp의 반환값 0은 “첫 통과”를 뜻한다. 복구 호출에서도 0을 그대로 반환하면 첫 통과와 구별할 수 없으므로 표준은 특별한 규칙을 둔다.',
'code':r'''#include <setjmp.h>
#include <stdio.h>

int main(void)
{
    jmp_buf point;
    switch (setjmp(point)) {
        case 0:
            puts("initial");
            longjmp(point, 0);
            break;
        case 1:
            puts("zero becomes one");
            break;
        default:
            puts("other value");
            break;
    }
    return 0;
}''',
'predict':'longjmp(point, 0) 뒤 switch가 다시 0번 case를 실행해 무한 반복할까? 아니면 1번 case를 선택할까?',
'walk':['switch의 제어식에 setjmp를 직접 놓는 것은 허용된 형태다.','첫 실행에서 반환값 0이므로 case 0으로 들어간다.','longjmp의 값 인수가 0이면 복구 후 setjmp는 1처럼 평가된다.','따라서 case 1이 실행된다. break는 switch 밖으로 나가며 main이 정상 종료한다.'],
'flow':['첫 통과: switch(0) → initial.','longjmp(point, 0) 호출: 지정값 0을 1로 정규화.','복구: switch(1) → zero becomes one.'],
'pitfall':'“0을 전달했으니 setjmp도 반드시 0을 반환한다”는 오개념이다. longjmp의 값 0은 특별 취급된다.',
'debug':'case 1을 지우고 default만 남기면 출력이 어떻게 바뀔까? 반환값을 0으로 오해한 테스트가 어떤 버그를 만들까?',
'variation':r'''#include <setjmp.h>
#include <stdio.h>

int main(void)
{
    jmp_buf point;
    switch (setjmp(point)) {
        case 0:
            puts("initial");
            longjmp(point, 4);
            break;
        case 4:
            puts("recovered as four");
            break;
        default:
            puts("unexpected");
            break;
    }
    return 0;
}''',
'variation_task':'빈칸 채우기: longjmp(point, 4)에 맞춰 case __를 작성하고 복구 출력 문구를 바꾸자.',
'variation_reason':'longjmp의 0이 아닌 값은 setjmp의 복구 반환값이 된다. switch는 해당 정수 case로 분기한다.'
},
{
'key':'jump_helper','title':'깊은 함수에서 오류를 발견하고 호출자에게 복귀하기',
'why':'입력 검증 함수를 따로 분리하면 복잡한 프로그램의 하위 목표를 나눌 수 있다. 단, 이 방식은 오류 코드 반환이 더 적합한 일반적인 설계를 대체하지 않는다.',
'code':r'''#include <setjmp.h>
#include <stdio.h>

static void require_positive(long value, jmp_buf point)
{
    if (value <= 0) {
        longjmp(point, 2);
    }
    printf("accepted=%ld\n", value);
}

int main(void)
{
    jmp_buf point;
    switch (setjmp(point)) {
        case 0:
            require_positive(-3, point);
            puts("success");
            break;
        case 2:
            puts("invalid input");
            break;
        default:
            puts("unexpected");
            break;
    }
    puts("end");
    return 0;
}''',
'predict':'require_positive(-3, point) 다음의 puts("success")가 실행될까? 마지막 end는 출력될까?',
'walk':['static은 여기서 함수 이름을 현재 번역 단위 내부에만 보이게 한다. 전역 객체를 만들지 않는다.','jmp_buf를 매개변수로 전달하면 함수 안에서도 같은 복구 환경을 사용할 수 있다.','value <= 0이면 longjmp(point, 2)가 실행되므로 printf와 호출자 success 출력은 건너뛴다.','복구 지점의 case 2에서 오류 메시지를 출력하고 switch를 빠져나와 end를 출력한다.'],
'flow':['main: setjmp 최초 0 → require_positive(-3, point) 호출.','검증: -3 <= 0 → 참 → longjmp(point, 2).','복구: switch(2) → invalid input.','공통 종료: end. success는 실행되지 않는다.'],
'pitfall':'“longjmp를 사용하면 호출한 모든 함수가 자동으로 자원을 정리한다”는 틀리다. 중간 return 경로가 생략되므로 자원 소유를 따로 설계해야 한다.',
'debug':'require_positive에서 동적 메모리를 할당한 뒤 longjmp하면 누가 free할까? 자원 정리 책임이 모호하면 오류 코드 반환 방식으로 바꾸는 것이 더 안전하다.',
'variation':r'''#include <setjmp.h>
#include <stdio.h>

static void require_positive(long value, jmp_buf point)
{
    if (value <= 0) {
        longjmp(point, 2);
    }
    printf("accepted=%ld\n", value);
}

int main(void)
{
    jmp_buf point;
    switch (setjmp(point)) {
        case 0:
            require_positive(8, point);
            puts("success");
            break;
        case 2:
            puts("invalid input");
            break;
        default:
            puts("unexpected");
            break;
    }
    puts("end");
    return 0;
}''',
'variation_task':'독립 수정: -3을 8로 바꾼 뒤 어떤 문장이 실행되고 어떤 분기는 실행되지 않는지 적는다.',
'variation_reason':'양수는 오류 분기를 건너뛰므로 longjmp가 호출되지 않는다. 일반 함수 반환 후 success와 end가 차례로 출력된다.'
}
]
},
{
'id':'1.67','title':'신호 처리의 경계와 안전한 최소 실습 — <signal.h>',
'unit':'25: 오류 진단·제어 이탈·가변 인수',
'lead':[
'프로그램이 실행되는 동안 외부 사건이나 실행 환경의 통지를 받는 경우가 있다. C의 <signal.h>는 SIGINT, SIGTERM 같은 신호 이름과 signal, raise를 제공한다. 운영체제의 풍부한 신호 모델(POSIX sigaction, 프로세스 간 전달 등)과 ISO C 표준의 보장 범위는 다르므로 혼동하지 않는다.',
'신호 처리 함수는 일반 함수처럼 아무 작업이나 안전하게 수행할 수 없다. 특히 비동기 신호 처리 도중 printf, malloc 같은 함수를 호출하거나 임의의 객체에 접근하는 것은 이식 가능한 ISO C 프로그램에서 위험하다. 이번 실습에서는 신호 처리 함수가 매개변수를 사용하지 않고 즉시 반환하도록 하여 위험한 동작을 피한다.',
'전역 객체를 사용하지 않는 실습 원칙 때문에 전역 volatile sig_atomic_t 플래그를 쓰는 예제는 이 절에 넣지 않는다. 실제 사건을 주 프로그램에 안전하게 알리는 설계에서는 제한된 종류의 정적 저장 객체가 필요할 수 있으며, 그 경우에도 비동기 안전성·원자성·플랫폼 차이를 별도로 검토해야 한다.'
],
'concepts':[
('signal(signum, handler)', '특정 신호에 대한 처리 동작을 등록한다. 실패하면 SIG_ERR을 반환한다. 등록 뒤 신호 처리 동작의 유지 방식은 구현에 따라 달라질 수 있다.'),
('raise(signum)', '프로그램 자신에게 신호를 발생시킨다. 0이면 성공이며, 실제 운영체제의 터미널 단축키와는 별개로 동작을 시험할 수 있다.'),
('SIG_DFL / SIG_IGN', '기본 처리 동작과 무시 동작을 나타내는 특수 값이다. 기본 동작을 함부로 재현하면 프로그램이 종료될 수 있다.'),
('신호 처리 함수', '매개변수는 int이고 반환형은 void이다. 여기서는 신호 번호를 사용하지 않고 즉시 반환한다. 표준이 허용하지 않는 함수를 핸들러 안에서 호출하지 않는다.'),
('ISO C vs POSIX', 'ISO C가 제공하는 신호 목록과 전달·재설정 의미는 제한적이다. sigaction 등은 POSIX 인터페이스이며 표준 C만으로 모든 플랫폼 동작을 동일하게 만들 수 없다.')
],
'examples':[
{
'key':'signal_basic','title':'신호 처리 함수를 등록하고 직접 신호 발생시키기',
'why':'터미널이나 운영체제에 의존한 외부 신호를 기다리지 않고, 표준 raise 함수로 가장 작은 신호 흐름을 확인한다.',
'code':r'''#include <signal.h>
#include <stdio.h>

static void quiet_handler(int number)
{
    (void)number;
}

int main(void)
{
    if (signal(SIGINT, quiet_handler) == SIG_ERR) {
        puts("install failed");
        return 1;
    }
    if (raise(SIGINT) != 0) {
        puts("raise failed");
        return 1;
    }
    puts("raise returned");
    return 0;
}''',
'predict':'핸들러 안에서 printf를 호출하지 않아도 raise 이후 프로그램이 계속될까? 출력은 몇 줄일까?',
'walk':['quiet_handler는 함수이고 전역 객체가 아니다. static은 외부 링크를 막는 함수 선언의 속성이다.','(void)number는 이 예제에서 사용하지 않는 매개변수에 대한 명시적 표시다.','signal(SIGINT, quiet_handler)는 신호 처리 동작을 등록한다. 반환값 SIG_ERR은 실패를 뜻한다.','raise(SIGINT)는 프로그램 자신에게 신호를 발생시킨다. 처리 함수가 돌아온 뒤 프로그램은 계속된다.','핸들러는 출력·메모리 할당을 하지 않는다. puts는 핸들러 밖인 main에서만 호출한다.'],
'flow':['quiet_handler 등록에 성공한다.','raise(SIGINT)가 등록된 처리 함수를 호출한다.','quiet_handler는 아무 부작용 없이 반환한다.','main으로 돌아와 raise returned를 출력한다.'],
'pitfall':'“핸들러도 보통 함수니까 printf를 넣어도 언제나 안전하다”는 틀리다. 신호 처리 중 허용되는 연산에는 엄격한 제약이 있다.',
'debug':'quiet_handler 안에 puts를 추가하려는 코드를 보고 문제점을 설명하자. 관찰 출력은 왜 main에 배치하는가?',
'variation':r'''#include <signal.h>
#include <stdio.h>

static void quiet_handler(int number)
{
    (void)number;
}

int main(void)
{
    if (signal(SIGTERM, quiet_handler) == SIG_ERR) {
        return 1;
    }
    if (raise(SIGTERM) != 0) {
        return 1;
    }
    puts("term signal returned");
    return 0;
}''',
'variation_task':'신호 종류를 SIGINT에서 SIGTERM으로 바꾸고, 등록하는 신호와 발생시키는 신호를 같은 값으로 맞춘다.',
'variation_reason':'등록한 신호와 raise 인수가 다르면 의도한 핸들러가 호출되지 않을 수 있다. SIGTERM의 기본 동작을 실수로 실행하지 않도록 주의한다.'
},
{
'key':'signal_restore','title':'기존 신호 처리 동작을 저장하고 복원하기',
'why':'라이브러리 코드가 신호 동작을 임시로 변경할 때는 호출자의 설정을 임의로 남겨서는 안 된다. 원래 동작을 복원하는 책임을 연습한다.',
'code':r'''#include <signal.h>
#include <stdio.h>

static void quiet_handler(int number)
{
    (void)number;
}

int main(void)
{
    void (*previous)(int) = signal(SIGTERM, quiet_handler);
    if (previous == SIG_ERR) {
        puts("install failed");
        return 1;
    }
    if (raise(SIGTERM) != 0) {
        puts("raise failed");
        return 1;
    }
    if (signal(SIGTERM, previous) == SIG_ERR) {
        puts("restore failed");
        return 1;
    }
    puts("handler restored");
    return 0;
}''',
'predict':'previous는 정수일까 함수 포인터일까? 원래 동작을 복원하는 호출은 어느 문장인가?',
'walk':['void (*previous)(int)는 int 인수 하나를 받고 void를 반환하는 함수를 가리키는 포인터 선언이다.','signal은 이전 처리 동작을 반환하므로 previous에 저장한다. 실패 시 SIG_ERR과 비교한다.','raise(SIGTERM)는 임시 핸들러가 등록된 동안 한 번만 실행한다.','signal(SIGTERM, previous)는 원래 동작을 되돌리는 호출이다. 이후 새 신호를 발생시키지 않는다.','핸들러의 재설정 규칙은 구현 차이가 있으므로 반복 신호 처리 설계는 별도 플랫폼 지식이 필요하다.'],
'flow':['previous ← 이전 SIGTERM 처리 동작.','임시 핸들러 등록 상태에서 SIGTERM을 한 번 발생.','signal(SIGTERM, previous)로 원래 동작 복원.','handler restored 출력.'],
'pitfall':'“signal을 호출하면 반드시 현재 핸들러를 영구 유지한다”는 오개념이다. 신호 발생 뒤 처리 동작이 어떻게 유지되는지는 구현 차이가 있다.',
'debug':'복원 직후 raise(SIGTERM)을 한 번 더 넣는 코드를 상상하자. 원래 동작이 SIG_DFL이면 어떤 위험이 있는가?',
'variation':r'''#include <signal.h>
#include <stdio.h>

static void quiet_handler(int number)
{
    (void)number;
}

int main(void)
{
    void (*before)(int) = signal(SIGINT, quiet_handler);
    if (before == SIG_ERR) {
        return 1;
    }
    if (raise(SIGINT) != 0) {
        return 1;
    }
    if (signal(SIGINT, before) == SIG_ERR) {
        return 1;
    }
    puts("int handler restored");
    return 0;
}''',
'variation_task':'기존 함수를 복사하지 말고 SIGINT용으로 재구성하자. previous 이름을 before로 바꾼 뒤 신호를 한 번만 발생시키고 복원한다.',
'variation_reason':'변수 이름은 기능에 맞게 자유롭게 바꿀 수 있지만, 이전 동작을 저장하고 복원하는 순서는 지켜야 한다.'
},
{
'key':'signal_ignore','title':'기본 동작과 무시 동작을 구별하기',
'why':'실험 프로그램이 신호를 받았을 때 “등록된 처리 함수 실행”과 “아예 무시”가 서로 다른 선택이라는 사실을 확인한다.',
'code':r'''#include <signal.h>
#include <stdio.h>

int main(void)
{
    if (signal(SIGINT, SIG_IGN) == SIG_ERR) {
        puts("ignore unavailable");
        return 1;
    }
    if (raise(SIGINT) != 0) {
        puts("raise failed");
        return 1;
    }
    puts("signal ignored");
    return 0;
}''',
'predict':'SIG_IGN을 설정하면 사용자 정의 함수가 호출될까? 프로그램은 어디까지 진행할까?',
'walk':['SIG_IGN은 신호 처리 함수의 이름이 아니라 무시 동작을 나타내는 특수 매크로다.','signal(SIGINT, SIG_IGN) 호출이 성공하면 SIGINT는 무시하도록 설정된다.','raise(SIGINT)는 신호를 발생시키지만 무시 동작이므로 사용자 핸들러를 호출하지 않는다.','main은 signal ignored를 출력한다. SIG_DFL로 바꾸면 같은 결과라고 가정하면 안 된다.'],
'flow':['SIGINT 처리 동작을 SIG_IGN으로 설정.','raise(SIGINT) 호출.','무시 동작이 적용되어 프로그램 흐름 유지.','signal ignored 출력.'],
'pitfall':'“SIG_DFL과 SIG_IGN은 둘 다 별도 함수가 없다는 뜻이므로 같다”는 틀리다. 기본 동작은 프로그램 종료 등일 수 있고 무시와 다르다.',
'debug':'SIG_IGN을 SIG_DFL로 무심코 바꾼 프로그램을 직접 실행하지 말고, 종료 가능성에 관해 설명하자.',
'variation':r'''#include <signal.h>
#include <stdio.h>

int main(void)
{
    if (signal(SIGTERM, SIG_IGN) == SIG_ERR) {
        return 1;
    }
    if (raise(SIGTERM) != 0) {
        return 1;
    }
    puts("term ignored");
    return 0;
}''',
'variation_task':'신호 종류를 SIGTERM으로 바꿔 무시 동작을 확인한다. 안전을 위해 SIG_DFL로 변경해 실행하지 않는다.',
'variation_reason':'무시 동작이 설정된 상태에서 발생한 신호는 프로그램 흐름을 종료시키지 않는다. 이 실습은 플랫폼 고유 신호 처리 API를 사용하지 않는다.'
}
]
}
]

INDEPENDENT=[
('1.65','실전 독립 작성: 예외 플래그를 두 번 검사하는 도구',r'''#include <fenv.h>
#include <stdio.h>

int main(void)
{
#ifdef FE_INVALID
    if (feclearexcept(FE_ALL_EXCEPT) != 0) {
        return 1;
    }
    printf("before=%d\n", fetestexcept(FE_INVALID) != 0);
    if (feraiseexcept(FE_INVALID) != 0) {
        return 1;
    }
    printf("after=%d\n", fetestexcept(FE_INVALID) != 0);
#else
    puts("FE_INVALID unavailable");
#endif
    return 0;
}''','처음에는 플래그가 없고, 명시적으로 올린 뒤에는 존재한다. 지원되는 구현에서는 before=0, after=1이다.'),
('1.66','실전 독립 작성: 검사 함수에서 복구 코드 3 전달',r'''#include <setjmp.h>
#include <stdio.h>

static void check_code(int code, jmp_buf point)
{
    if (code != 0) {
        longjmp(point, 3);
    }
    puts("valid");
}

int main(void)
{
    jmp_buf point;
    switch (setjmp(point)) {
        case 0:
            check_code(1, point);
            break;
        case 3:
            puts("recovery code 3");
            break;
        default:
            puts("unexpected");
            break;
    }
    return 0;
}''','check_code(1)는 오류 분기를 선택하고 longjmp로 복구 코드 3을 전달한다. valid는 출력되지 않는다.'),
('1.67','실전 독립 작성: 무시 동작 설정 후 한 번만 raise',r'''#include <signal.h>
#include <stdio.h>

int main(void)
{
    if (signal(SIGINT, SIG_IGN) == SIG_ERR) {
        return 1;
    }
    if (raise(SIGINT) != 0) {
        return 1;
    }
    puts("continued safely");
    return 0;
}''','SIG_IGN을 등록한 뒤 raise(SIGINT)를 한 번 호출한다. 무시가 성공하면 main의 출력 문장까지 진행한다.')
]

# Test all complete programs with both C11 compilers; use identical outputs.
entries=[]
for sec in SECTIONS:
    for idx,ex in enumerate(sec['examples'],1):
        for kind,code in [('example',ex['code']),('variation',ex['variation'])]:
            name=f"ch01_{sec['id'].replace('.','_')}_{idx:02d}_{kind}_{ex['key']}.c"
            entries.append({'section':sec['id'],'key':ex['key'],'kind':kind,'name':name,'code':code})
for sec,title,code,reason in INDEPENDENT:
    entries.append({'section':sec,'key':'independent','kind':'independent','name':f"ch01_{sec.replace('.','_')}_independent.c",'code':code})

def run_test(e):
    path=CODES/e['name'];path.write_text(e['code'].rstrip()+'\n',encoding='utf-8')
    record={'name':e['name'],'section':e['section'],'kind':e['kind'],'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'compilers':{}}
    for cc in ('gcc','clang'):
        exe=CODES/(e['name'].removesuffix('.c')+'_'+cc)
        cmd=[cc,*FLAGS,str(path),'-lm','-o',str(exe)]
        comp=subprocess.run(cmd,capture_output=True,text=True,timeout=20)
        if comp.returncode:
            raise RuntimeError(f"compile {e['name']} {cc}: {comp.stderr}")
        proc=subprocess.run([str(exe)],capture_output=True,text=True,timeout=8)
        if proc.returncode:
            raise RuntimeError(f"run {e['name']} {cc}: exit={proc.returncode} stderr={proc.stderr}")
        record['compilers'][cc]={'compile':'pass','run':'pass','stdout':proc.stdout,'stderr':proc.stderr,'exit':proc.returncode}
        exe.unlink(missing_ok=True)
    if record['compilers']['gcc']['stdout']!=record['compilers']['clang']['stdout']:
        raise RuntimeError(f"output differs {e['name']}")
    record['stdout']=record['compilers']['gcc']['stdout']
    return record

verified=[run_test(e) for e in entries]
rec_by_name={r['name']:r for r in verified}
for e in entries:e['actual']=rec_by_name[e['name']]['stdout']

# Build document
D=Document();sec=D.sections[0];sec.top_margin=Inches(.77);sec.bottom_margin=Inches(.74);sec.left_margin=Inches(.86);sec.right_margin=Inches(.85)
sec.header_distance=Inches(.34);sec.footer_distance=Inches(.39)
styles=D.styles
normal=styles['Normal'];normal.font.name='Malgun Gothic';normal.font.size=Pt(10.1);normal.font.color.rgb=RGBColor(36,47,54)
normal.paragraph_format.space_after=Pt(5.2);normal.paragraph_format.line_spacing=1.26
for st,size,before,after in [('Title',20,0,10),('Heading 1',15,13,8),('Heading 2',12,10,5),('Heading 3',10.5,8,4)]:
    s=styles[st];s.font.name='Malgun Gothic';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,112,117)
    s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.keep_with_next=True
if 'Code C' not in styles:
    cs=styles.add_style('Code C',WD_STYLE_TYPE.PARAGRAPH)
else:cs=styles['Code C']
cs.font.name='Consolas';cs.font.size=Pt(8.05);cs.font.color.rgb=RGBColor(29,52,64)
cs.paragraph_format.space_after=Pt(0);cs.paragraph_format.line_spacing=1.06
if 'Small Note' not in styles:
    ns=styles.add_style('Small Note',WD_STYLE_TYPE.PARAGRAPH)
else:ns=styles['Small Note']
ns.font.name='Malgun Gothic';ns.font.size=Pt(9.2);ns.font.color.rgb=RGBColor(62,80,90)
ns.paragraph_format.space_after=Pt(4)

head=sec.header.paragraphs[0];head.text='C언어 종합교과서  |  이해 중심 정식 개정판  |  제1장';head.style='Small Note';head.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('제1장 · 1.65–1.67  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)

def para(t,style=None):
    p=D.add_paragraph(style=style);p.add_run(t);return p

def line(label,text):
    p=D.add_paragraph();p.paragraph_format.keep_together=True
    r=p.add_run(label+'  ');r.bold=True;r.font.color.rgb=RGBColor(0,112,117)
    p.add_run(text);return p

def codeblock(code):
    # One editable Word paragraph per C listing; prevent source blocks splitting mid-function.
    p=D.add_paragraph(style='Code C')
    p.paragraph_format.left_indent=Inches(.19)
    p.paragraph_format.right_indent=Inches(.1)
    p.paragraph_format.keep_together=True
    p.paragraph_format.space_after=Pt(7)
    lines=code.strip('\n').splitlines()
    p.add_run('\n'.join(ln.replace('    ','\u00a0\u00a0\u00a0\u00a0') for ln in lines))


def output_block(text):
    p=D.add_paragraph(style='Small Note');p.add_run('실제 출력 (GCC·Clang 공통, 현재 검증 환경)').bold=True
    for ln in text.rstrip('\n').splitlines():
        p=D.add_paragraph(style='Code C');p.paragraph_format.left_indent=Inches(.25);p.add_run(ln)

D.add_heading('C언어 종합교과서',0)
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('이해 중심 정식 개정판 · 제1장 1.65~1.67절');r.bold=True;r.font.size=Pt(14);r.font.color.rgb=RGBColor(0,112,117)
para('이 회차의 범위: 부동소수점 환경 · 비지역 복구 · 신호 처리')
para('읽는 순서: 필요성 → 개념 → 전체 코드 → 실행 전 예측 → 문장별 해설 → 값·제어 흐름 → 실제 출력 → 오개념 교정 → 인접 변형 실습과 정답 → 독립 작성·누적 복습.')
para('학습상의 경계: 모든 프로그램은 ISO C11 문법을 기준으로 작성했다. 실제 부동소수점 플래그·신호 동작은 구현 지원 범위를 확인해야 하며, 표준 C와 POSIX 확장을 구분한다. 실행 예시는 GCC·Clang으로 검증했지만 MSVC 등 모든 구현의 동작을 보장하지 않는다.')
para('새 개정판 1.1~1.64절은 이전 회차에서 완료 보고된 범위다. 이번 문서는 해당 내용을 복제하지 않고 그 다음 절을 독립 집필한 산출물이다.')

for s in SECTIONS:
    D.add_heading(s['id']+'  '+s['title'],1)
    for p in s['lead']:para(p)
    D.add_heading('핵심 개념을 먼저 정리하기',2)
    for term,definition in s['concepts']:line(term,definition)
    D.add_heading('학습 목표와 작은 단계',2)
    para('① 인터페이스가 필요한 문제를 파악한다. ② 상태·제어 흐름을 확인하는 최소 코드를 만든다. ③ 실패 경로를 포함해 결과를 예측한다. ④ 실제 컴파일·실행 결과로 가설을 검증한다. ⑤ 코드를 바꾸어도 안전한 이유를 설명한다.')
    for idx,e in enumerate(s['examples'],1):
        D.add_heading(f"{s['id']}.{idx}  {e['title']}",2)
        line('왜 필요한가 · 언제 쓰는가',e['why'])
        D.add_heading('완성 예제 — 전체 코드',3)
        codeblock(e['code'])
        line('실행 전에 예측하기',e['predict'])
        D.add_heading('문장별 해설',3)
        for j,t in enumerate(e['walk'],1):line(f'{j:02d}',t)
        D.add_heading('값과 제어 흐름 추적',3)
        for j,t in enumerate(e['flow'],1):line(f'단계 {j}',t)
        output_block(next(x['actual'] for x in entries if x['section']==s['id'] and x['key']==e['key'] and x['kind']=='example'))
        line('흔한 실수 · 오개념 교정',e['pitfall'])
        line('단계별 디버깅',e['debug'])
        D.add_heading('예제를 바꾸며 실력 다지기 — 바로 옆에서 풀기',3)
        line('변형 실습',e['variation_task'])
        D.add_heading('변형 정답 — 전체 C 코드',3)
        codeblock(e['variation'])
        output_block(next(x['actual'] for x in entries if x['section']==s['id'] and x['key']==e['key'] and x['kind']=='variation'))
        line('정답의 근거',e['variation_reason'])
        if idx==2:
            line('도움 줄이기', '앞의 완성 예제를 다시 보지 말고 이번 변형의 조건·호출 순서를 스스로 먼저 구성해 본 뒤 아래 정답과 대조하자.')
    D.add_heading('도움 없이 직접 작성하기 — 독립 실습',2)
    secid,title,code,reason=next(z for z in INDEPENDENT if z[0]==s['id'])
    para(title+'. 기존 예제를 복사하지 않고 먼저 요구 기능을 작은 단계로 나눠 직접 작성하자. 예상 출력을 종이에 적은 뒤 코드를 완성한다.')
    D.add_heading('독립 실습 정답과 검증 근거',3)
    codeblock(code)
    output_block(next(x['actual'] for x in entries if x['section']==s['id'] and x['kind']=='independent'))
    line('왜 이 답이 되는가',reason)
    D.add_heading('인출 연습 · 누적 복습',2)
    if s['id']=='1.65':
        line('문제','1. 이전에 배운 errno와 부동소수점 예외 플래그가 같은 저장 장소인지 설명하자. 2. 함수가 반환한 0과 fetestexcept의 0은 각각 무엇을 뜻하는가?')
        line('정답·근거','1. 별개의 오류 진단 기제다. errno는 라이브러리 오류 상태를 나타내는 정수 객체이며, fenv 예외 플래그는 별도의 부동소수점 환경 상태다. 2. feclearexcept/feraiseexcept의 0은 성공이고 fetestexcept의 0은 지정 플래그가 켜지지 않았음을 뜻한다.')
    elif s['id']=='1.66':
        line('문제','1. 이전 절의 fegetround 복원과 longjmp를 결합한다면 무엇을 먼저 설계해야 할까? 2. 일반 return과 longjmp의 자원 정리 경로가 어떻게 다른가?')
        line('정답·근거','1. longjmp가 건너뛰는 경로에서도 반올림 모드를 복원하도록 복구 지점의 정리 책임을 설계해야 한다. 2. return은 일반 호출 반환 경로를 따르지만 longjmp는 중간 반환을 건너뛰므로 파일 닫기·free 등 명시적 정리 경로가 필요하다.')
    else:
        line('문제','1. setjmp 복구 경로와 신호 처리 함수가 모두 평범한 함수 호출과 같다고 볼 수 없는 이유는? 2. volatile이 붙은 변수라면 무조건 스레드 안전한가?')
        line('정답·근거','1. setjmp/longjmp는 비지역 제어 이동이고 신호 처리 함수는 실행 중 허용 연산이 제한된다. 2. 아니다. volatile은 원자성·동기화를 제공하지 않는다. 이후 <stdatomic.h>/<threads.h>에서 데이터 경쟁을 별도로 다룬다.')
    D.add_heading('절 종료 점검',2)
    para('새 용어를 정의할 수 있는가? 각 함수의 성공/실패 반환값을 구별하는가? 예측과 실제 출력이 달랐던 원인을 설명할 수 있는가? 학습용 출력문을 제거해도 남아야 할 필수 상태 관리 로직은 무엇인가?')

D.add_heading('이번 회차 기술 검증과 이식성 메모',1)
para('검증 대상은 완성 예제 9개, 인접 변형 정답 9개, 독립 작성 정답 3개로 총 21개 C11 소스다. GCC·Clang에서 각각 -std=c11 -Wall -Wextra -Wpedantic -Werror -O0을 적용하고, <fenv.h> 관련 연결을 위해 -lm을 추가했다. 모든 소스의 실제 stdout을 두 컴파일러에서 비교했다.')
para('운영체제와 C 구현의 차이를 주의한다. FE_INVALID 등의 개별 매크로 지원, 반올림 모드 지원, signal의 핸들러 재설정 동작은 구현 조건을 따른다. 신호 처리 함수 내부의 I/O를 금지하고, 신호 처리 후 정상적인 main에서만 관찰 출력을 수행했다.')
para('다음 권장 범위: <stdatomic.h>와 <threads.h>를 사용한 원자 연산·스레드 동기화. 이후에는 C23 변경점, 전처리기 심화, 다중 파일 설계 및 양 원서의 LAB·Exercise·Programming·Mini Project·Advanced Topic 개별 목표를 전수 대조하여 누락을 보충한다.')
para('저작권 준수: 원서의 문장·문제·그림·코드를 대량 복제하지 않았다. 원서 전수 대조가 완료되었다는 주장은 하지 않는다.')

docx=OUT/'C언어_종합교과서_이해중심_정식개정판_제1장_1.65-1.67.docx';D.save(docx)
zip_path=OUT/'official_ch01_1_65-1_67_sources.zip'
with zipfile.ZipFile(zip_path,'w',compression=zipfile.ZIP_DEFLATED) as z:
    for e in entries:z.write(CODES/e['name'],arcname=e['name'])
verification={'run':'ch01_1.65-1.67','timestamp':NOW,'flags':FLAGS,'link_flags':['-lm'],'programs':len(verified),'compiler_runs':2*len(verified),'all_passed':True,'results':verified}
(OUT/'official_ch01_1_65-1_67_c11_verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
# check exact C code in DOCX via paragraph grouping (code lines); compare each source by line containment in document text
from docx import Document as OpenDocument
R=OpenDocument(docx)
alltext='\n'.join(p.text.replace('\u00a0',' ') for p in R.paragraphs)
code_checks=[]
for e in entries:
    normalized=e['code'].strip()
    found=normalized in alltext
    code_checks.append({'name':e['name'],'match':found})
if not all(x['match'] for x in code_checks):
    raise RuntimeError('DOCX code/source mismatch: '+str([x for x in code_checks if not x['match']]))
(OUT/'official_ch01_1_65-1_67_code_match.json').write_text(json.dumps({'all_match':True,'count':len(code_checks),'files':code_checks},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(OUT/'official_ch01_1_65-1_67_pedagogy_qa.json').write_text(json.dumps({'six_methods':['완성→변형→독립','간격·인출 복습','실행 전 예측','단계별 디버깅','하위 목표 분해','오개념 진단·교정'],'in_place_variations':9,'in_place_answers':9,'global_objects':0,'new_sections':['1.65','1.66','1.67'],'manual_review':'code and content checked; page rendering pending'},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'docx':str(docx),'docx_bytes':docx.stat().st_size,'zip':str(zip_path),'programs':len(entries),'compiler_runs':len(entries)*2,'code_match':True},ensure_ascii=False))