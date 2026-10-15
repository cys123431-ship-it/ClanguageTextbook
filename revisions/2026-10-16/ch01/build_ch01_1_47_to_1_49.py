from pathlib import Path
import subprocess, json, hashlib, zipfile, datetime, re, textwrap
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/mnt/data/c_textbook_run15'); ROOT.mkdir(exist_ok=True)
TITLE='C언어_종합교과서_이해중심_정식개정판_제1장_1.47-1.49.docx'
S=[
('1.47','표준 스트림의 오류·종료 상태를 구별하기','입력 함수가 더 읽을 것이 없다고 알려 주었을 때, 그것이 정상적인 파일 끝인지 저장장치나 스트림의 오류인지 구별해야 한다. 또한 오류를 일반 결과와 섞어 출력하면 파이프나 다른 프로그램이 결과를 받아 처리하기 어렵다.','stdin은 표준 입력, stdout은 정상 결과를 위한 표준 출력, stderr는 진단 메시지를 위한 표준 오류 스트림이다. EOF는 파일 끝 자체가 아니라 입력 함수의 실패 또는 종료를 알리는 정수 반환값이다. feof와 ferror는 마지막 입력 연산 이후의 상태 표시를 검사한다. clearerr는 두 표시를 지우지만 이미 실패한 입출력을 되돌리지 않는다. ungetc는 다음 문자 입력 전에 한 문자를 되돌려 넣는 기능이다.'),
('1.48','형식 문자열을 안전하게 읽고 쓰기','같은 숫자라도 사용 목적에 따라 앞에 0을 채우거나 자릿수를 정렬해야 한다. 하지만 sprintf처럼 크기 제한이 없는 함수를 잘못 쓰면 배열 범위를 넘어설 수 있다. 입력에서도 필드 폭 없이 %s를 쓰면 버퍼에 들어갈 길이를 제어하기 어렵다.','printf 계열의 %d는 int, %s는 널 종료된 문자 배열, %zu는 size_t에 대응한다. 서식의 폭과 정밀도는 값의 표현을 제어하지만 출력 버퍼의 안전을 자동으로 보장하지 않는다. snprintf는 목적지 크기를 인수로 받고, 정상적인 경우 널 종료를 제외하고 필요했을 문자 수를 반환한다. sscanf의 %Ns는 최대 N개의 문자를 읽고 널 문자를 덧붙이므로 목적지에는 N+1칸이 필요하다. 변환 성공 개수를 반드시 검사한다.'),
('1.49','문자열과 바이트 배열의 길이·복사·검색','문자열은 문자 배열 안에 저장되지만 배열의 모든 칸이 문자열에 포함되는 것은 아니다. 널 문자를 기준으로 문자열 길이를 계산하고, 복사할 때는 겹치는 영역과 목적지 크기를 따져야 한다. 실제 데이터에서는 줄 끝 제거, 구분자 찾기, 바이트 이동을 자주 수행한다.','strlen은 널 문자 앞의 바이트 수를 세며 널 종료되지 않은 버퍼에 호출하면 안전하지 않다. sizeof는 배열이 선언된 현재 범위에서는 배열 전체 바이트 수를 반환하지만 함수 매개변수로 포인터를 받은 곳에서는 포인터 크기가 된다. memcpy는 겹치지 않는 메모리 영역을 복사할 때 사용하며, 겹치는 영역은 memmove를 사용한다. strchr·strstr는 발견한 위치를 가리키는 포인터를 돌려주고 실패하면 NULL을 돌려준다. UTF-8 한글 글자 수와 strlen의 바이트 수는 다를 수 있다.')]
E=[]
def add(sec,key,title,why,code,expected,explain,trace,predict,misconception,debug,exercise,variation,variant_output,solution,review):
    E.append(dict(sec=sec,key=key,title=title,why=why,code=textwrap.dedent(code).strip()+'\n',expected=expected,explain=explain,trace=trace,predict=predict,misconception=misconception,debug=debug,exercise=exercise,variation=textwrap.dedent(variation).strip()+'\n',variant_output=variant_output,solution=solution,review=review))

add('1.47','A','stdout과 stderr를 분리하고 반환값으로 실패를 알리기',
'실제 자동화 도구는 표준 출력만 읽어 다음 단계로 넘기기도 한다. 따라서 실패를 일반 데이터처럼 stdout에 출력하면 결과 파일을 오염시킬 수 있다. 이 예제는 실행 인수의 개수를 먼저 확인하고 성공한 경우에만 결과를 출력한다.',
'''#include <stdio.h>
int main(int argc, char *argv[])
{
    if (argc != 2) {
        fprintf(stderr, "사용법: 이름 하나를 전달하세요\\n");
        return 1;
    }
    printf("안녕하세요, %s\\n", argv[1]);
    return 0;
}''','안녕하세요, C언어\n',
['#include <stdio.h>는 fprintf와 printf의 선언 및 stderr를 제공한다.','argc는 프로그램 이름까지 포함한 문자열 개수다. 사용자 인수 하나를 요구하므로 2인지 검사한다.','fprintf(stderr, ...)는 오류를 표준 오류 스트림으로 보낸다. stdout에 섞지 않는 것이 핵심이다.','return 1은 이 예제에서 실패를 알리는 비영(非零) 종료 상태다. 환경에 따라 구체적인 상태 해석은 달라질 수 있다.','argv[1]은 argc 검사를 통과했을 때만 사용한다.','성공한 경우의 printf는 실제 결과이며, 잘못된 인수에 대한 사용법 출력은 진단 메시지다.'],
'argc=2 → 조건 argc!=2는 거짓 → argv[1]의 문자열 C언어를 %s에 대입 → 한 줄 출력 → main에서 0 반환.',
'실행 인수로 C언어를 전달했을 때 어느 스트림으로 어떤 문장이 나올까? 인수가 하나도 없다면 성공 경로가 실행될까?',
'"stderr로 출력하면 자동으로 프로그램이 종료된다"는 틀리다. 출력과 종료는 별개다. 여기서는 명시적인 return 1이 흐름을 끝낸다.',
'단계 1: argc 검사를 제거하면 argv[1]이 NULL일 때 %s에 전달될 수 있다. 단계 2: 오류 메시지를 printf로 바꾸면 정상 결과와 진단을 구별하기 어려워진다.',
'점진적 실습: 이름 대신 과목명 하나를 받아 "학습 과목: ..."으로 출력하라. 오류 조건은 유지하고 과목명은 C11을 사용한다.',
'''#include <stdio.h>
int main(int argc, char *argv[])
{
    if (argc != 2) {
        fprintf(stderr, "사용법: 과목 하나를 전달하세요\\n");
        return 1;
    }
    printf("학습 과목: %s\\n", argv[1]);
    return 0;
}''','학습 과목: C11\n',
'argc 조건은 그대로이며 정상 출력의 형식 문자열과 오류 안내만 변경했다. 오류가 아닌 출력이므로 stdout에 남긴다.',
'1.44절의 argc·argv와 1.32절의 널 포인터 검사를 떠올리자. argc 검사 없이 argv[1]을 읽는 것이 왜 위험한가? 답: 인수가 없으면 argv[1]은 널 포인터일 수 있어 문자열로 사용하면 안 된다.')

add('1.47','B','파일 끝과 입력 오류를 구별하여 읽은 문자 수 계산하기',
'파일에서 한 글자씩 읽는 반복문은 EOF가 나오면 끝나지만, 파일 끝과 오류는 서로 다른 상태다. 프로그램이 결과를 믿어도 되는지 판단하려면 루프가 끝난 뒤 ferror를 확인한다.',
'''#include <stdio.h>
int main(void)
{
    FILE *fp = tmpfile();
    if (fp == NULL) return 1;
    if (fputs("ABC", fp) == EOF) { fclose(fp); return 1; }
    rewind(fp);
    int ch;
    int count = 0;
    while ((ch = fgetc(fp)) != EOF) {
        ++count;
    }
    int bad = ferror(fp);
    int ended = feof(fp);
    if (fclose(fp) == EOF || bad) return 1;
    printf("읽은 문자=%d 파일끝=%d\\n", count, ended != 0);
    return 0;
}''','읽은 문자=3 파일끝=1\n',
['tmpfile은 임시 이진 파일 스트림을 연다. 생성 실패 시 NULL이므로 검사한다.','fputs는 텍스트 문자열 ABC를 스트림에 기록한다. EOF를 반환하면 쓰기에 실패한 것이다.','rewind는 스트림을 파일 시작으로 돌리고 오류 표시를 지운다.','fgetc의 반환형은 int다. 모든 unsigned char 값과 EOF를 구별할 수 있어야 하기 때문이다.','while 조건에서 한 글자를 읽고 EOF가 아니면 count를 증가시킨다. 세 글자 뒤에 EOF를 받는다.','ferror와 feof는 루프가 끝난 뒤 검사한다. 여기서는 파일 끝을 읽어 종료했으므로 feof가 참이다.','fclose는 임시 파일 자원을 정리한다. tmpfile로 생성된 파일은 정상 종료 시 자동 삭제 대상이다.'],
'쓰기 ABC → rewind → fgetc=A(count=1) → B(2) → C(3) → EOF에서 반복 종료 → ferror=0, feof는 0이 아님 → 출력 3과 1.',
'ABC를 기록한 뒤 읽기 시작하면 count는 3인가 4인가? EOF는 읽은 문자 수에 포함되는가?',
'"fgetc가 EOF를 반환하면 반드시 파일 끝이다"는 틀리다. 읽기 오류도 EOF를 반환할 수 있다. 종료 원인은 feof와 ferror로 판별한다.',
'단계 1: int ch를 char ch로 바꾸면 EOF와 실제 문자 값의 구별이 깨질 수 있다. 단계 2: rewind를 생략하면 현재 파일 위치에서 읽게 되어 기대 결과를 얻지 못한다.',
'부분 완성: 임시 파일에 HELLO를 기록하여 읽은 문자 수 5를 출력하라. 파일 끝 확인 방식은 유지한다.',
'''#include <stdio.h>
int main(void)
{
    FILE *fp = tmpfile();
    if (fp == NULL) return 1;
    if (fputs("HELLO", fp) == EOF) { fclose(fp); return 1; }
    rewind(fp);
    int ch;
    int count = 0;
    while ((ch = fgetc(fp)) != EOF) {
        ++count;
    }
    int bad = ferror(fp);
    int ended = feof(fp);
    if (fclose(fp) == EOF || bad) return 1;
    printf("읽은 문자=%d 파일끝=%d\\n", count, ended != 0);
    return 0;
}''','읽은 문자=5 파일끝=1\n',
'문자열만 HELLO로 바뀌었다. 정상적으로 기록한 다섯 바이트를 읽은 뒤 EOF를 확인하므로 5와 1이 출력된다.',
'1.45절의 FILE*와 1.41절의 malloc 반환값은 모두 NULL 검사를 요구하지만, 해제 함수는 각각 fclose와 free로 다르다.')

add('1.47','C','ungetc로 구분자를 한 글자 되돌려 보기',
'문자 단위 분석기에서는 한 글자를 미리 읽었다가 다음 처리 단계에 넘겨야 할 때가 있다. ungetc는 최소 한 글자를 다시 읽을 수 있게 하는 표준 C 기능이다. 여러 글자를 연속해서 되돌릴 수 있다고 가정하면 안 된다.',
'''#include <stdio.h>
int main(void)
{
    FILE *fp = tmpfile();
    if (fp == NULL) return 1;
    if (fputs("7,8", fp) == EOF) { fclose(fp); return 1; }
    rewind(fp);
    int first = fgetc(fp);
    int separator = fgetc(fp);
    if (first == EOF || separator != ',') { fclose(fp); return 1; }
    if (ungetc(separator, fp) == EOF) { fclose(fp); return 1; }
    int again = fgetc(fp);
    int last = fgetc(fp);
    if (again == EOF || last == EOF) { fclose(fp); return 1; }
    if (fclose(fp) == EOF) return 1;
    printf("처음=%c 다시=%c 끝=%c\\n", first, again, last);
    return 0;
}''','처음=7 다시=, 끝=8\n',
['첫 번째 fgetc는 문자 7을 읽고 스트림 위치를 다음 문자로 이동한다.','두 번째 fgetc는 구분자 쉼표를 읽는다.','ungetc(separator, fp)는 쉼표를 되돌려 놓아 다음 읽기에서 다시 쉼표를 받게 한다.','다시 읽은 문자를 int에 보관하여 EOF와 구별한다.','ungetc는 호출이 성공하면 EOF를 반환하지 않는다. 실패 경로를 검사한다.','실제 파서에서는 되돌린 문자를 다음 함수가 읽게 할 수 있다. 여기의 printf는 학습용 확인 출력이다.'],
'파일의 읽기 순서: 7 → , → ungetc(,) → , → 8. 따라서 first=7, again=쉼표, last=8이다.',
'구분자를 되돌리지 않으면 again과 last는 무엇을 읽을까? 왜 반환값을 int에 저장할까?',
'"ungetc는 파일 내용을 물리적으로 수정한다"는 틀리다. 스트림의 다음 읽기 동작을 바꾸는 것이며 원본 파일 데이터를 수정하는 함수가 아니다.',
'단계 1: ungetc 실패를 무시하면 다음 읽기 결과가 달라질 수 있다. 단계 2: 첫 글자와 구분자를 검증한 뒤에만 다음 단계로 넘어간다.',
'독립 변형: 자료를 9;4로 바꾸고 구분자를 세미콜론으로 검사한 뒤 같은 흐름을 완성하라.',
'''#include <stdio.h>
int main(void)
{
    FILE *fp = tmpfile();
    if (fp == NULL) return 1;
    if (fputs("9;4", fp) == EOF) { fclose(fp); return 1; }
    rewind(fp);
    int first = fgetc(fp);
    int separator = fgetc(fp);
    if (first == EOF || separator != ';') { fclose(fp); return 1; }
    if (ungetc(separator, fp) == EOF) { fclose(fp); return 1; }
    int again = fgetc(fp);
    int last = fgetc(fp);
    if (again == EOF || last == EOF) { fclose(fp); return 1; }
    if (fclose(fp) == EOF) return 1;
    printf("처음=%c 다시=%c 끝=%c\\n", first, again, last);
    return 0;
}''','처음=9 다시=; 끝=4\n',
'문자 값과 구분자 검사만 변경했다. ungetc의 성공 여부를 확인하고 한 문자만 되돌리는 계약은 그대로다.',
'1.46절의 fseek와 달리 ungetc는 단 한 글자를 미리 읽는 분석 과정에 특히 적합하다.')

add('1.48','A','snprintf로 고정 폭 번호를 안전하게 만들기',
'주문번호·파일 이름·로그 식별자를 일정한 폭으로 표현하면 사람이 비교하기 쉽다. 출력 버퍼의 크기를 함수에 알려 주면 길이 초과를 검사할 수 있다.',
'''#include <stdio.h>
int main(void)
{
    char buffer[32];
    int needed = snprintf(buffer, sizeof buffer, "ID-%04d", 23);
    if (needed < 0 || (size_t)needed >= sizeof buffer) {
        puts("서식 생성 실패");
        return 1;
    }
    printf("결과=%s 길이=%d\\n", buffer, needed);
    return 0;
}''','결과=ID-0023 길이=7\n',
['buffer[32]는 출력 문자와 마지막 널 문자를 저장할 배열이다.','sizeof buffer는 전체 32바이트를 반환한다. 함수 안의 포인터가 아니라 여기서 선언한 배열이기 때문이다.','%04d는 최소 네 자리, 부족한 자리는 0으로 채운 십진 정수 출력이다.','snprintf의 정상 반환값은 널 문자를 제외하고 필요했을 문자 수다.','반환값이 음수이면 인코딩 오류이고, 버퍼 크기 이상이면 잘렸을 수 있다.','size_t로 비교할 때는 먼저 needed<0을 검사한다. C의 || 단락 평가로 음수인 경우 오른쪽 형 변환을 평가하지 않는다.'],
'23 → %04d에서 0023 → 접두어 ID-를 붙여 ID-0023 → 글자 7개 + 널 문자 1개 → needed=7, 버퍼 32칸 안에 수용.',
'needed는 널 문자를 포함하는가? %04d의 0과 4는 각각 무슨 뜻인가?',
'"snprintf는 크기를 주었으므로 반환값을 검사할 필요가 없다"는 틀리다. 버퍼를 넘겨 쓰지는 않더라도 결과가 잘렸을 수 있다.',
'단계 1: sprintf로 바꾸면 목적지 크기를 전달할 수 없다. 단계 2: needed==sizeof buffer도 잘린 경우이므로 >=를 사용한다.',
'완성 예제 변형: 7을 다섯 자리 번호로 만들어 ID-00007을 출력하라. 필요한 길이도 계산해 보자.',
'''#include <stdio.h>
int main(void)
{
    char buffer[32];
    int needed = snprintf(buffer, sizeof buffer, "ID-%05d", 7);
    if (needed < 0 || (size_t)needed >= sizeof buffer) {
        puts("서식 생성 실패");
        return 1;
    }
    printf("결과=%s 길이=%d\\n", buffer, needed);
    return 0;
}''','결과=ID-00007 길이=8\n',
'다섯 자리 정수 00007과 접두어 ID- 세 글자를 더하면 필요한 문자는 여덟 개다. 마지막 널 문자를 위한 공간도 충분하다.',
'1.32절의 배열 경계와 1.4절의 sizeof 의미를 함께 떠올려, 왜 반환값과 목적지 크기를 비교하는지 설명하라.')

add('1.48','B','잘린 문자열을 감지하고 버퍼 내용을 구별하기',
'네트워크 메시지나 로그를 고정 크기 배열에 기록할 때 데이터가 잘렸는지 알아야 한다. 잘린 내용을 완전한 결과라고 믿으면 식별자가 달라지는 문제가 생긴다.',
'''#include <stdio.h>
int main(int argc, char *argv[])
{
    char buffer[6];
    const char *input = argc > 1 ? argv[1] : "ABCDEFGHI";
    int needed = snprintf(buffer, sizeof buffer, "%s", input);
    if (needed < 0) return 1;
    printf("필요=%d 저장=%s 잘림=%d\\n",
           needed, buffer, (size_t)needed >= sizeof buffer);
    return 0;
}''','필요=9 저장=ABCDE 잘림=1\n',
['buffer[6]에는 널 종료 문자를 포함해 최대 여섯 바이트가 들어간다.','ABCDEFGHI는 아홉 개의 영문자이므로 완전한 문자열에는 열 바이트가 필요하다.','snprintf는 목적지 크기를 6으로 받았으므로 최대 다섯 문자와 널 문자를 기록한다.','정상 반환값 needed=9는 저장된 다섯 문자가 아니라 저장하고 싶었던 아홉 문자 수다.','(size_t)needed >= sizeof buffer 비교는 결과가 잘렸는지 나타내는 0 또는 1을 출력한다.'],
'목적지 6칸 → A,B,C,D,E,\\0 기록 → needed=9 → 9>=6 참 → 잘림=1.',
'버퍼 크기가 6이면 왜 ABCDEF 여섯 글자를 모두 담을 수 없는가? 반환값은 5인가 9인가?',
'"반환값 9는 버퍼에 아홉 글자를 썼다는 뜻"은 틀리다. 정상 반환값은 널 문자를 제외한 전체 필요 문자 수다.',
'단계 1: 크기를 0으로 넘긴 snprintf의 사용법과 크기 6을 넘긴 경우를 혼동하지 않는다. 단계 2: 널 문자를 위한 한 칸을 항상 계산한다.',
'부분 완성: 버퍼를 5칸으로 줄이고 "MONDAY"를 넣어 잘린 결과를 예측하라.',
'''#include <stdio.h>
int main(int argc, char *argv[])
{
    char buffer[5];
    const char *input = argc > 1 ? argv[1] : "MONDAY";
    int needed = snprintf(buffer, sizeof buffer, "%s", input);
    if (needed < 0) return 1;
    printf("필요=%d 저장=%s 잘림=%d\\n",
           needed, buffer, (size_t)needed >= sizeof buffer);
    return 0;
}''','필요=6 저장=MOND 잘림=1\n',
'전체 문자 수는 6, 버퍼에 기록 가능한 문자 수는 4이므로 MOND와 널 문자를 저장한다. 반환값 6으로 잘림을 확인한다.',
'1.12절의 널 종료 문자열과 1.41절의 메모리 크기 계산을 복습하자.')

add('1.48','C','sscanf 필드 폭과 변환 개수로 입력 검증하기',
'파일 한 줄을 읽은 뒤 그 안의 짧은 코드와 수량을 해석할 수 있다. 숫자가 빠졌거나 문자열이 너무 길어도 안전하게 거부하려면 폭 제한과 반환값 검사가 필요하다.',
'''#include <stdio.h>
int main(void)
{
    const char *line = "APPLE 17";
    char name[8] = {0};
    int quantity = 0;
    int converted = sscanf(line, "%7s %d", name, &quantity);
    if (converted != 2) {
        puts("입력 형식 오류");
        return 1;
    }
    printf("품목=%s 수량=%d\\n", name, quantity);
    return 0;
}''','품목=APPLE 수량=17\n',
['sscanf는 파일 스트림이 아니라 이미 메모리에 있는 문자열에서 자료를 읽는다.','%7s는 최대 일곱 문자를 읽고 마지막 널 문자를 추가하므로 name[8]에 맞는다.','%d는 int 변수에 정수를 저장하므로 &quantity로 주소를 전달한다.','문자열 배열 name은 식에서 첫 원소의 포인터로 변환되므로 name 자체를 넘긴다.','변환 성공 개수가 2가 아니면 두 항목이 모두 들어왔다고 보장할 수 없다.','반환값 검사 뒤에만 name과 quantity를 결과로 사용한다.'],
'line의 첫 토큰 APPLE → name에 저장(널 종료) → 다음 정수 17 → quantity=17 → converted=2 → 출력.',
'"%7s"의 7은 목적지 배열 크기인가, 널 문자를 제외한 최대 입력 문자 수인가?',
'"%s에 배열 길이만 전달하면 scanf가 알아서 보호한다"는 틀리다. 필드 폭은 형식 문자열에서 명시해야 한다.',
'단계 1: %s로 바꾸면 길이 제한이 사라진다. 단계 2: converted 검사 없이 quantity를 신뢰하면 입력이 부족한 경우 논리 오류가 난다.',
'독립 실습: "BREAD 4"를 읽고 같은 형식으로 출력하라. 문자열 버퍼 크기와 반환값 검사는 유지한다.',
'''#include <stdio.h>
int main(void)
{
    const char *line = "BREAD 4";
    char name[8] = {0};
    int quantity = 0;
    int converted = sscanf(line, "%7s %d", name, &quantity);
    if (converted != 2) {
        puts("입력 형식 오류");
        return 1;
    }
    printf("품목=%s 수량=%d\\n", name, quantity);
    return 0;
}''','품목=BREAD 수량=4\n',
'BREAD는 다섯 글자이므로 일곱 문자 제한 안에 있다. 정수 4도 성공적으로 변환되어 두 항목이 출력된다.',
'1.44절의 strtol과 비교해 보자. sscanf는 간단한 구조의 입력을 읽지만 숫자 뒤 잔여 문자와 범위 문제를 더 엄격히 검사하려면 별도 검증이 필요하다.')

add('1.49','A','sizeof와 strlen이 서로 다른 길이를 뜻하는 이유',
'문자열을 복사하거나 출력할 때 배열 전체 크기와 실제 문자열 길이를 혼동하면 잘못된 길이를 전달하게 된다. 널 문자가 포함되는 위치를 눈으로 확인하는 예제다.',
'''#include <stdio.h>
#include <string.h>
int main(void)
{
    char word[12] = "CAT";
    size_t capacity = sizeof word;
    size_t length = strlen(word);
    printf("배열=%zu 문자열=%zu 널위치=%zu\\n",
           capacity, length, length);
    return 0;
}''','배열=12 문자열=3 널위치=3\n',
['char word[12]는 열두 개의 char 요소를 가진 배열을 만든다.','문자열 리터럴 CAT를 초기화하면 C,A,T,\\0가 들어가고 나머지 칸은 0으로 초기화된다.','sizeof word는 전체 배열의 크기를 바이트로 반환한다. C에서는 sizeof(char)가 1이다.','strlen(word)는 첫 널 문자 전의 세 바이트를 센다. 널 문자는 길이에 포함하지 않는다.','문자열 길이가 3이면 널 문자는 인덱스 3에 있다.','%zu는 size_t 형식의 값을 출력하기 위한 C99 이후의 표준 서식이다.'],
'인덱스 0=C, 1=A, 2=T, 3=\\0, 4~11=0 → sizeof=12, strlen=3, 널 인덱스=3.',
'배열 크기 12와 문자열 길이 3 중 어떤 값이 실제로 저장된 글자 수인가? 널 문자는 길이에 포함되는가?',
'"sizeof(word)는 항상 strlen(word)+1"이라는 주장은 틀리다. 배열에는 여유 공간이 있을 수 있다.',
'단계 1: 널 종료가 없는 배열에 strlen을 호출하면 배열 밖을 읽을 위험이 있다. 단계 2: 함수 매개변수 char *p에 sizeof p를 쓰면 포인터 크기를 얻는다.',
'완성 예제를 수정해 word[12]="DOGS"일 때 결과를 예측하라.',
'''#include <stdio.h>
#include <string.h>
int main(void)
{
    char word[12] = "DOGS";
    size_t capacity = sizeof word;
    size_t length = strlen(word);
    printf("배열=%zu 문자열=%zu 널위치=%zu\\n",
           capacity, length, length);
    return 0;
}''','배열=12 문자열=4 널위치=4\n',
'배열 크기는 그대로 12이고 실제 문자열은 네 바이트다. 널 종료 문자는 인덱스 4에 들어간다.',
'1.32절의 배열·포인터 차이를 떠올리자. 배열이 함수 인수로 전달될 때 왜 sizeof 결과가 달라질 수 있는가? 답: 매개변수 선언이 포인터로 조정되기 때문이다.')

add('1.49','B','겹치는 메모리 이동에는 memmove 사용하기',
'문자열 맨 앞에 접두 문자를 추가할 때 기존 내용을 한 칸 뒤로 밀어야 한다. 출발 영역과 목적지 영역이 겹치므로 memcpy 대신 memmove가 필요하다.',
'''#include <stdio.h>
#include <string.h>
int main(void)
{
    char text[16] = "ABCDE";
    size_t length = strlen(text);
    memmove(text + 1, text, length + 1);
    text[0] = 'A';
    printf("변경=%s 길이=%zu\\n", text, strlen(text));
    return 0;
}''','변경=AABCDE 길이=6\n',
['text[16]에는 기존 문자열과 새 접두 문자, 널 문자를 넣을 공간이 충분하다.','strlen(text)는 기존 문자열의 길이 5를 구한다.','memmove의 목적지는 text+1, 원본은 text다. 두 영역은 겹친다.','복사 바이트 수 length+1은 문자열 마지막 널 문자까지 함께 이동하기 위한 값이다.','text[0]에 접두 문자 A를 넣는다. 원래 문자열의 첫 문자도 뒤로 밀려 남아 있다.','memmove는 겹치는 영역에서도 올바른 결과를 보장하지만, 배열의 크기를 자동 검사하지는 않는다.'],
'처음 A B C D E \\0 → 인덱스 1부터 A B C D E \\0 복사 → 인덱스 0=A → AABCDE\\0 → 길이 6.',
'length가 5일 때 왜 이동 바이트 수는 6인가? memcpy로 바꾸어도 안전한가?',
'"memcpy와 memmove는 언제나 바꿔 쓸 수 있다"는 틀리다. 메모리 영역이 겹치면 memcpy는 정의되지 않은 동작을 일으킬 수 있다.',
'단계 1: length만 복사하면 널 문자가 이동하지 않아 문자열이 끝나지 않을 수 있다. 단계 2: 목적지에 length+2바이트 공간이 있는지 확인한다.',
'부분 완성: "XYZ" 앞에 X를 한 글자 붙여 "XXYZ"로 만들어 보라.',
'''#include <stdio.h>
#include <string.h>
int main(void)
{
    char text[16] = "XYZ";
    size_t length = strlen(text);
    memmove(text + 1, text, length + 1);
    text[0] = 'X';
    printf("변경=%s 길이=%zu\\n", text, strlen(text));
    return 0;
}''','변경=XXYZ 길이=4\n',
'원본 세 문자와 널 문자 네 바이트를 한 칸 뒤로 이동한다. 인덱스 0에 X를 저장하므로 XXYZ가 된다.',
'1.41절의 동적 메모리로 확장할 때는 새 접두 문자와 널 문자를 위한 할당 크기까지 계산해야 한다.')

add('1.49','C','strcspn과 strchr·strstr로 한 줄을 분석하기',
'사용자 입력의 마지막 줄바꿈을 제거하고 구분자를 찾는 일은 파일 처리에서 흔하다. 문자열 함수가 반환하는 포인터를 안전하게 확인하는 습관을 익힌다.',
'''#include <stdio.h>
#include <string.h>
int main(void)
{
    char line[64] = "blue,green,red\\n";
    line[strcspn(line, "\\n")] = '\\0';
    char *comma = strchr(line, ',');
    if (comma == NULL) return 1;
    size_t first_length = (size_t)(comma - line);
    int found = strstr(line, "green") != NULL;
    printf("첫 항목 길이=%zu 검색=%d\\n", first_length, found);
    return 0;
}''','첫 항목 길이=4 검색=1\n',
['strcspn(line,"\\n")은 줄바꿈 문자가 나타나기 전의 바이트 수를 반환한다.','그 위치에 널 문자를 써서 뒤쪽 줄바꿈을 문자열에서 제거한다. 줄바꿈이 없다면 기존 널 위치를 다시 널로 쓴다.','strchr(line, 쉼표)는 첫 쉼표 위치를 가리키는 포인터를 돌려준다. 없으면 NULL이다.','comma-line은 두 포인터가 같은 배열을 가리킬 때 가능한 포인터 차이 계산이다. 여기서는 첫 항목의 길이가 된다.','strstr는 부분 문자열 green의 시작 주소를 돌려준다. NULL이 아니면 찾았다는 뜻이다.','영문 ASCII 범위의 예제이므로 바이트 수와 영문자 수가 같다. 한국어 UTF-8 문자열에서는 이 둘이 다를 수 있다.'],
'줄바꿈 제거 후 blue,green,red → 첫 쉼표 인덱스 4 → first_length=4 → green 부분 문자열 발견 → found=1.',
'첫 쉼표가 없으면 comma-line을 계산해도 될까? 줄바꿈이 없는 문자열에서 strcspn은 무엇을 반환할까?',
'"strstr는 찾은 문자열을 새로 할당해 반환한다"는 틀리다. 원래 문자열 내부를 가리키는 포인터를 반환하며 free하면 안 된다.',
'단계 1: NULL 검사 없이 포인터 차이를 계산하면 위험하다. 단계 2: 줄바꿈이 없는 경우에도 배열 경계 내 널 위치를 쓸 수 있다.',
'독립 변형: "sun,moon,star\\n"에서 첫 항목 길이와 moon의 존재를 확인하라.',
'''#include <stdio.h>
#include <string.h>
int main(void)
{
    char line[64] = "sun,moon,star\\n";
    line[strcspn(line, "\\n")] = '\\0';
    char *comma = strchr(line, ',');
    if (comma == NULL) return 1;
    size_t first_length = (size_t)(comma - line);
    int found = strstr(line, "moon") != NULL;
    printf("첫 항목 길이=%zu 검색=%d\\n", first_length, found);
    return 0;
}''','첫 항목 길이=3 검색=1\n',
'첫 쉼표 앞에는 sun 세 바이트가 있고 moon이 연속된 부분 문자열로 존재한다. 따라서 3과 1을 출력한다.',
'1.35절의 문자 포인터와 1.44절의 문자열 비교를 복습하자. 문자열 함수 반환 포인터가 가리키는 데이터의 소유권은 누구에게 있는가? 답: 원래 배열 소유자에게 있다.')

FLAGS=['-std=c11','-Wall','-Wextra','-Wpedantic','-Werror','-O0']
SRC=ROOT/'sources'; SRC.mkdir(exist_ok=True)
RESULTS=[]
ARGS={'1.47-A':['C언어'],'1.47-A-variation':['C11']}
for e in E:
    for kind,code,expected in [('main',e['code'],e['expected']),('variation',e['variation'],e['variant_output'])]:
        fn=f"ch{e['sec'].replace('.','_')}_{e['key'].lower()}_{kind}.c"; sp=SRC/fn;sp.write_text(code,encoding='utf-8')
        args=ARGS.get(f"{e['sec']}-{e['key']}"+('-variation' if kind=='variation' else ''),[])
        for cc in ['gcc','clang']:
            exe=ROOT/(fn+'.'+cc)
            build=subprocess.run([cc,*FLAGS,str(sp),'-o',str(exe)],capture_output=True,text=True,timeout=30)
            if build.returncode: raise RuntimeError(f'Compile {fn} {cc}: {build.stderr}')
            run=subprocess.run([str(exe),*args],cwd=ROOT,capture_output=True,text=True,timeout=30)
            passed=(run.returncode==0 and run.stdout==expected and run.stderr=='')
            RESULTS.append({'source':fn,'compiler':cc,'flags':FLAGS,'args':args,'returncode':run.returncode,'stdout':run.stdout,'stderr':run.stderr,'expected_stdout':expected,'passed':passed})
            if not passed: raise RuntimeError(f'Run {fn} {cc}: {run.returncode}, {run.stdout!r}, {run.stderr!r} != {expected!r}')
            exe.unlink()

D=Document();sect=D.sections[0];sect.top_margin=Cm(2.2);sect.bottom_margin=Cm(2.0);sect.left_margin=Cm(2.45);sect.right_margin=Cm(2.45)
normal=D.styles['Normal'];normal.font.name='Malgun Gothic';normal.font.size=Pt(9.5);normal.font.color.rgb=RGBColor(32,44,48);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.15
for n,size,col in [('Title',17,(11,106,113)),('Heading 1',15,(9,103,113)),('Heading 2',11.5,(20,90,101))]:
    st=D.styles[n];st.font.name='Malgun Gothic';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(*col);st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
cs=D.styles.add_style('Editable C Code',WD_STYLE_TYPE.PARAGRAPH);cs.font.name='Consolas';cs.font.size=Pt(8.2);cs.font.color.rgb=RGBColor(24,47,52);cs.paragraph_format.space_after=Pt(0);cs.paragraph_format.line_spacing=1.04
ls=D.styles.add_style('Learning Label',WD_STYLE_TYPE.PARAGRAPH);ls.font.name='Malgun Gothic';ls.font.size=Pt(9.4);ls.font.bold=True;ls.font.color.rgb=RGBColor(0,110,116);ls.paragraph_format.keep_with_next=True;ls.paragraph_format.space_before=Pt(8);ls.paragraph_format.space_after=Pt(3)
def p(t='',style=None): return D.add_paragraph(t,style=style)
def label(t): p(t,'Learning Label')
def code(t):
    for i,line in enumerate(t.rstrip('\n').split('\n')):
        pp=p(line or ' ','Editable C Code'); pp.paragraph_format.keep_together=True
        if i<2:pp.paragraph_format.keep_with_next=True
        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F1F6F6');pp._p.get_or_add_pPr().append(shd)
def bullet(t): p('• '+t)
D.add_heading('한국어 C언어 종합교과서 — 이해 중심 정식 개정판',0)
p('제1장 1.47~1.49절 | 스트림 오류·종료 상태 · 안전한 형식 입출력 · 문자열과 바이트 이동')
p('앞서 검증된 제1장 1.46절 다음에서 이어지는 독립 신규 원고. 과거 거부 원고·제1~11장 구 초안의 문장을 재사용하지 않는다.')
p('읽는 순서: 왜 필요한가 → 개념 → 전체 실행 코드 → 실행 전 예측 → 문장별 설명 → 상태 추적 → 실제 출력 → 오개념 교정 → 디버깅 → 인접 변형 실습·정답.')
p('표준 범위: stdio.h, string.h, snprintf, sscanf, feof, ferror, ungetc, tmpfile, strlen, memmove, strcspn, strchr, strstr은 C11 표준 라이브러리 기능이다. 파일 위치·콘솔·셸의 구체적인 동작은 구현·운영체제와 구별한다. 모든 실습은 전역 객체 없이 작성한다.')
p('실습 안전: tmpfile은 임시 스트림이며 실제 파일을 덮어쓰지 않는다. 인수 전달 예제는 컴파일 후 명령행 인수를 설정하여 실행한다. 학습용 printf와 실제 데이터 처리의 역할을 구분한다.')
for sec,title,intro,concept in S:
    D.add_heading(f'{sec}절  {title}',1);p(intro);label('새 개념을 먼저 이해하기');p(concept)
    for e in [x for x in E if x['sec']==sec]:
        D.add_heading(f"{sec}-{e['key']}  {e['title']}",2)
        label('왜 필요한가 · 언제 쓰는가');p(e['why'])
        label('전체 C11 코드');code(e['code'])
        if sec=='1.47' and e['key']=='A':p('실행 인수: C언어')
        label('실행하기 전에 스스로 예측');p(e['predict'])
        label('문장별 해설 — 왜 이 문장을 쓰는가')
        for x in e['explain']:bullet(x)
        label('값과 실행 흐름 추적');p(e['trace'])
        label('실제 출력');code(e['expected'])
        label('그럴듯한 오개념 진단');p(e['misconception'])
        label('단계별 디버깅 — 원인과 수정');p(e['debug'])
        label('예제를 바꾸며 실력 다지기');p(e['exercise'])
        label('바로 이어지는 정답 — 독립 작성 후 비교');code(e['variation'])
        if sec=='1.47' and e['key']=='A':p('정답 실행 인수: C11')
        label('변형 정답의 실행 결과와 근거');code(e['variant_output']);p(e['solution'])
    label('간격 복습 · 이전 절과 연결하는 인출 질문');p(E[[x['sec'] for x in E].index(sec)]['review'])
D.add_heading('1.47~1.49절 종합 적용: 안전한 한 줄 처리기',1)
p('작은 하위 목표부터 구현한다. ① 입력 한 줄을 저장할 고정 크기 배열을 마련한다. ② fgets 반환값을 검사하고 파일 끝·오류를 구별한다. ③ strcspn으로 줄바꿈을 제거한다. ④ strchr로 첫 쉼표를 찾고 NULL인지 확인한다. ⑤ 첫 항목의 바이트 수를 계산한다. ⑥ snprintf로 결과를 만들고 반환값을 검사한다. ⑦ 정상 결과는 stdout, 진단은 stderr로 출력한다.')
D.add_heading('도움 없이 해결하는 통합 실습',2)
p('문제: 문자열 "alpha,beta\\n"을 배열에 저장한 뒤 첫 쉼표 앞의 바이트 수를 구하여 크기 32인 결과 배열에 "LEN=5" 형식으로 작성하라. strcspn, strchr, snprintf를 사용하고 모든 반환값과 경계를 검사하라. 출력은 결과 문자열 한 줄만 허용한다. 한글 UTF-8 문자열은 이 단계에서 제외하고 다음 문자 인코딩 단원으로 넘긴다.')
label('정답 설계와 검사 근거')
p('strcspn으로 널 종료 위치를 결정하고, strchr의 결과가 NULL이 아닌 경우에만 포인터 차이를 계산한다. 결과는 size_t이므로 snprintf 서식은 %zu를 사용한다. snprintf 반환값이 음수인지 먼저 확인한 뒤 (size_t)반환값과 sizeof 결과 배열을 비교한다. 이 절의 완성 예제 세 개를 연결하되 전역 객체 없이 작성한다. 경계 입력으로 쉼표 없음, 빈 첫 항목, 줄바꿈 없음, 길이가 긴 입력을 검증한다.')
D.add_heading('표준·오개념·다음 학습 위치',2)
p('C11 표준은 stderr 스트림과 tmpfile, snprintf, sscanf, string.h 함수를 규정한다. POSIX의 getline, Windows의 별도 파일 API는 ISO C11의 필수 기능이 아니므로 여기서 사용하지 않았다. C23에서 추가·변경된 라이브러리 함수는 후속 표준 비교 단원에서 구분한다. 함수 사용법을 이해했다고 해서 두 기준 원서의 개별 LAB·Exercise·부록까지 대조한 것은 아니다.')
p('다음 절은 1.50절부터 이어서 표준 라이브러리의 문자 분류·대소문자 변환, 문자열 변환과 검색, 정렬·검색 알고리즘의 호출 계약을 깊이 있게 학습한다. 모든 원서 학습목표 전수 대조는 미완료 상태다.')
footer=sect.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('C언어 종합교과서 · 이해 중심 정식 개정판 · 제1장')
DOCX=ROOT/TITLE;D.save(DOCX)
ZIP=ROOT/'official_ch01_1_47-1_49_sources.zip'
with zipfile.ZipFile(ZIP,'w',zipfile.ZIP_DEFLATED) as z:
    for f in sorted(SRC.glob('*.c')):z.write(f,'sources/'+f.name)
VERIFY=ROOT/'official_ch01_1_47-1_49_c11_verification.json'
VERIFY.write_text(json.dumps({'sections':['1.47','1.48','1.49'],'source_count':len(E)*2,'compile_run_count':len(RESULTS),'passed':all(z['passed'] for z in RESULTS),'results':RESULTS},ensure_ascii=False,indent=2),encoding='utf-8')
text='\n'.join(x.text for x in Document(DOCX).paragraphs)
checks=[]
for e in E:
    for kind in ('code','variation'):
        lines=e[kind].rstrip().splitlines();ok='\n'.join(lines) in text
        checks.append({'section':e['sec'],'example':e['key'],'kind':kind,'exact_code_block_present':ok})
if not all(x['exact_code_block_present'] for x in checks):raise RuntimeError('DOCX code mismatch')
(ROOT/'official_ch01_1_47-1_49_code_match.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
progress={'completed_range_reported':'1.1~1.49','new_sections':['1.47','1.48','1.49'],'next_section':'1.50','progress_percent':None,'source_objective_full_audit':False,'stage':'스트림 오류·형식 입출력·문자열 메모리 처리 집필 및 검증','note':'이전 1.1~1.46절은 선행 회차 보고를 참조하며, 이번에는 새 1.47~1.49절만 코드/문서 검증'}
(ROOT/'official_rewrite_progress_1_1-1_49.json').write_text(json.dumps(progress,ensure_ascii=False,indent=2),encoding='utf-8')
coverage={'full_audit_complete':False,'sources':[{'name':'c언어 압축 (1)(1).pdf','status':'원본 스캔 PDF 7쪽 구성표 확인, 15장 파일 스트림 범위 연관. 장·절·LAB 전수 대조 미완료'},{'name':'C_Programming.pdf','status':'원본 스캔 PDF 3쪽 표지 확인, 개별 학습목표 전수 대조 미완료'},{'name':'C Programming Language - 2nd Edition (OCR).pdf','status':'OCR 목차의 K&R 7.2, 7.4, 7.6, 7.7, 7.8.1, 7.8.3 관련 절 확인. 본문·연습문제 전수 대조 미완료'}],'new_sections':[{'section':'1.47','objectives':'stderr와 종료 상태, EOF와 feof/ferror, ungetc와 tmpfile','original_pdf_item_by_item':'partial: K&R OCR 7.6/7.7/7.8.3'},{'section':'1.48','objectives':'snprintf 버퍼 한계, 서식 폭·반환값, sscanf 변환 개수와 입력 폭','original_pdf_item_by_item':'partial: K&R OCR 7.2/7.4'},{'section':'1.49','objectives':'strlen/sizeof, memmove 겹침, strcspn/strchr/strstr 검색','original_pdf_item_by_item':'partial: K&R OCR 7.8.1'}],'verified_compiler_runs':len(RESULTS),'source_audit_status':'partial'}
(ROOT/'official_rewrite_objective_tracking_1_1-1_49.json').write_text(json.dumps(coverage,ensure_ascii=False,indent=2),encoding='utf-8')
manifest={'run':'ch01_1.47-1.49','created_at':datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=9))).isoformat(),'files':[],'github_delivery':'pending','library_delivery':'pending','source_full_audit_complete':False}
for f in [DOCX,ZIP,VERIFY,ROOT/'official_ch01_1_47-1_49_code_match.json',ROOT/'official_rewrite_progress_1_1-1_49.json',ROOT/'official_rewrite_objective_tracking_1_1-1_49.json',ROOT/'build_ch01_1_47_to_1_49.py']:
    manifest['files'].append({'name':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(ROOT/'official_rewrite_manifest_1_47-1_49.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('SUCCESS',len(E),'examples',len(RESULTS),'compiler-runs',DOCX.stat().st_size)