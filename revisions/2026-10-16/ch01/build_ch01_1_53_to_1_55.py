#!/usr/bin/env python3
# Reproducible C11 textbook chapter authoring and verification.
import os, json, zipfile, hashlib, subprocess, re, sys, textwrap
from pathlib import Path
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parent
TITLE='C언어_종합교과서_이해중심_정식개정판_제1장_1.53-1.55.docx'
STEM='official_ch01_1_53-1_55'
FLAGS=['-std=c11','-Wall','-Wextra','-Wpedantic','-Werror','-O0']
CODES=ROOT/'sources'; CODES.mkdir(exist_ok=True)

S=[
{'id':'1.53','title':'수학 함수와 실수 계산: <math.h>','intro':'직각삼각형의 거리, 음수의 절댓값, 반올림 방식, 거듭제곱을 매번 직접 계산하면 중복 코드와 오류가 늘어난다. <math.h>는 이러한 수치 계산을 표준 인터페이스로 제공한다. 실수는 유한한 비트로 근삿값을 표현하므로 수학적 실수와 컴퓨터의 부동소수점 값을 구분해야 한다.','concept':'double은 보통 실수 계산에 사용하는 부동소수점 자료형이다. sqrt(x)는 제곱근, hypot(x,y)는 두 값의 제곱합의 제곱근을 계산한다. floor(x)는 x 이하의 가장 큰 정수값을 double로, ceil(x)는 x 이상의 가장 작은 정수값을 double로, trunc(x)는 0 방향으로 소수 부분을 버린 값을 double로 돌려준다. pow(x,y)는 x의 y제곱, fabs(x)는 절댓값이다. 이 함수들은 표준 C11의 <math.h>에 있으며 GCC/Clang 환경에서는 링크할 때 -lm이 필요할 수 있다.','items':[
{'key':'A','name':'제곱근과 거리 계산','why':'지도 좌표 사이 거리나 벡터 크기를 구할 때 제곱근이 필요하다. 직접 반복문으로 근삿값을 만드는 대신 표준 함수를 쓰면 의도를 명확히 표현할 수 있다.','code':'''#include <math.h>\n#include <stdio.h>\n\nint main(void)\n{\n    double root = sqrt(81.0);\n    double distance = hypot(3.0, 4.0);\n    printf("root=%.1f distance=%.1f\\n", root, distance);\n    return 0;\n}''','out':'root=9.0 distance=5.0\n','predict':'81의 제곱근은 얼마인가? hypot(3,4)는 단순히 3+4인가? 출력의 %.1f는 몇 자리까지 표시하는가?','lines':['#include <math.h>는 sqrt와 hypot의 선언을 제공한다. 선언 없이 호출하면 정확한 인수·반환 형식을 검사할 수 없다.','#include <stdio.h>는 결과를 확인할 printf 선언을 제공한다.','sqrt(81.0)는 9.0을, hypot(3.0,4.0)는 sqrt(3²+4²)인 5.0을 계산한다.','%.1f는 double 값을 소수 첫째 자리까지 표시하는 출력 형식이다. printf는 학습 확인용이고 실제 계산은 두 수학 함수가 담당한다.'],'trace':'root ← 9.0 → distance ← 5.0 → printf가 두 값을 순서대로 표시 → main 종료.','pitfall':'sqrt(-1.0)은 실수 범위의 제곱근이 없다. 반환값과 오류 처리에는 수학 라이브러리의 도메인 규칙을 고려해야 한다.','misconception':'“hypot(a,b)는 언제나 sqrt(a*a+b*b)와 수치적으로 완전히 같다”는 생각은 틀리다. hypot은 중간 제곱에서 발생할 수 있는 오버플로·언더플로를 피하도록 설계된 연산이다.','variant':'직각삼각형 변 길이를 5.0, 12.0으로 바꾸고 49.0의 제곱근을 출력하라.','vcode':'''#include <math.h>\n#include <stdio.h>\n\nint main(void)\n{\n    double root = sqrt(49.0);\n    double distance = hypot(5.0, 12.0);\n    printf("root=%.1f distance=%.1f\\n", root, distance);\n    return 0;\n}''','vout':'root=7.0 distance=13.0\n','vwhy':'5²+12²=169이고 그 제곱근은 13이다. 학습용 printf 외의 계산 문장은 그대로 재사용한다.'},
{'key':'B','name':'음수에서 달라지는 버림 방식','why':'금액, 좌표, 격자 인덱스를 변환할 때 “버림”이 아래 방향인지 0 방향인지 구분하지 않으면 음수 입력에서 오답이 발생한다.','code':'''#include <math.h>\n#include <stdio.h>\n\nint main(void)\n{\n    double x = -2.7;\n    printf("floor=%.0f ceil=%.0f trunc=%.0f\\n",\n           floor(x), ceil(x), trunc(x));\n    return 0;\n}''','out':'floor=-3 ceil=-2 trunc=-2\n','predict':'-2.7을 “내림”하면 -2인가 -3인가? 세 함수 중 어떤 것이 0 방향으로 이동하는가?','lines':['x에 -2.7을 저장한다. 음수이므로 수직선에서 왼쪽과 0 방향이 서로 다르다.','floor(x)는 x보다 크지 않은 가장 큰 정수값인 -3.0이다.','ceil(x)는 x보다 작지 않은 가장 작은 정수값인 -2.0이다.','trunc(x)는 소수 부분을 0 방향으로 제거하므로 -2.0이다.','%.0f는 소수 자릿수를 출력하지 않는다. 함수의 반환 자료형 자체가 정수형으로 바뀌는 것은 아니다.'],'trace':'x=-2.7 → floor=-3.0, ceil=-2.0, trunc=-2.0 → 표시 형식으로 소수점 없이 출력.','pitfall':'(int)x와 floor(x)는 음수에서 결과가 다를 수 있다. 정수형 변환은 표현 가능 범위 조건도 확인해야 한다.','misconception':'“floor는 소수점만 없애는 함수다”는 오개념이다. -2.7의 floor는 -3.0이고 trunc는 -2.0이다.','variant':'x를 -3.2로 바꿔 세 함수를 비교하라.','vcode':'''#include <math.h>\n#include <stdio.h>\n\nint main(void)\n{\n    double x = -3.2;\n    printf("floor=%.0f ceil=%.0f trunc=%.0f\\n",\n           floor(x), ceil(x), trunc(x));\n    return 0;\n}''','vout':'floor=-4 ceil=-3 trunc=-3\n','vwhy':'수직선에서 -3.2 이하 최대 정수는 -4, 이상 최소 정수는 -3, 0 방향으로 버리면 -3이다.'},
{'key':'C','name':'거듭제곱과 절댓값','why':'성장률의 반복 적용, 오차의 크기 비교에 거듭제곱과 절댓값을 사용한다. 연산 의도를 코드에서 명시하는 편이 유지보수에 유리하다.','code':'''#include <math.h>\n#include <stdio.h>\n\nint main(void)\n{\n    double power = pow(2.0, 5.0);\n    double magnitude = fabs(-3.5);\n    printf("power=%.1f magnitude=%.1f\\n", power, magnitude);\n    return 0;\n}''','out':'power=32.0 magnitude=3.5\n','predict':'pow(2,5)는 10인가 32인가? fabs(-3.5)는 원래 부호를 남기는가?','lines':['pow(2.0,5.0)은 2를 다섯 번 곱한 값 32.0을 구한다.','fabs(-3.5)는 double 절댓값 3.5를 구한다. int의 abs와 혼동하지 않는다.','printf는 계산된 값을 사람이 확인하기 위한 출력이다. 실제 계산은 pow와 fabs 호출에서 끝난다.'],'trace':'power ← 32.0 → magnitude ← 3.5 → 두 값 출력.','pitfall':'pow를 정수 거듭제곱의 정확한 대체품이라고 가정하면 안 된다. 반환은 부동소수점이며 큰 값은 표현 범위를 넘을 수 있다.','misconception':'“fabs는 정수만 받는다”는 오해다. fabs는 double에 쓰는 함수이고 정수 절댓값 abs는 <stdlib.h>에 있다.','variant':'3의 세제곱과 -4.25의 절댓값을 출력하라.','vcode':'''#include <math.h>\n#include <stdio.h>\n\nint main(void)\n{\n    double power = pow(3.0, 3.0);\n    double magnitude = fabs(-4.25);\n    printf("power=%.1f magnitude=%.2f\\n", power, magnitude);\n    return 0;\n}''','vout':'power=27.0 magnitude=4.25\n','vwhy':'3×3×3=27이고 -4.25의 절댓값은 4.25다. %.2f로 소수 둘째 자리까지 확인한다.'}]},
{'id':'1.54','title':'날짜·시각·시간 차이: <time.h>','intro':'일정 표시와 처리 시간 측정은 “시간”이라는 한 단어로 묶이지만 서로 다른 작업이다. 달력의 연월일·시분초 표현, 두 시각의 경과 시간, CPU 사용 시간은 각각 다른 인터페이스가 필요하다.','concept':'<time.h>의 struct tm은 연도·월·일·시·분·초를 멤버로 저장한다. tm_year는 1900을 뺀 연도이고 tm_mon은 0부터 시작한다. strftime은 이 구조체의 필드를 문자열로 형식화한다. mktime은 지역 시간 기준의 struct tm을 time_t로 바꾸며 유효하지 않은 값은 정규화할 수 있다. difftime은 두 time_t 사이의 초 차이를 double로 계산한다. time()은 현재 달력 시간을 얻고 clock()은 프로그램의 CPU 시간을 측정하는 용도로 구분한다. time_t의 내부 인코딩이나 초 단위 숫자 자체를 가정해서는 안 된다.','items':[
{'key':'A','name':'struct tm의 연도·월 인코딩','why':'날짜를 저장하는 구조체가 사람이 읽는 연도·월과 다른 숫자를 사용하기 때문에, 변환 규칙을 모르면 2026년이 126년으로 표시되는 버그가 생긴다.','code':'''#include <stdio.h>\n#include <time.h>\n\nint main(void)\n{\n    struct tm date = {0};\n    char buffer[32];\n    date.tm_year = 2026 - 1900;\n    date.tm_mon = 10 - 1;\n    date.tm_mday = 16;\n    if (strftime(buffer, sizeof buffer, "%Y-%m-%d", &date) == 0)\n        return 1;\n    puts(buffer);\n    return 0;\n}''','out':'2026-10-16\n','predict':'tm_year에 2026을 그대로 저장하면 어떤 문제가 생기는가? tm_mon에 10을 저장하면 10월인가 11월인가?','lines':['struct tm date = {0}은 모든 멤버를 0으로 초기화한다. 이 예제에서 필요한 연월일만 뒤이어 지정한다.','tm_year는 1900을 뺀 값 126, tm_mon은 0부터 세므로 9, tm_mday는 날짜 16이다.','buffer는 형식화된 텍스트를 받을 배열이다. sizeof buffer는 배열 전체 바이트 수를 계산한다.','strftime이 0을 반환하면 결과를 안전하게 출력할 수 없다고 보고 중단한다.','puts는 버퍼의 문자열을 표시하는 학습 확인 출력이다.'],'trace':'date.tm_year=126 → tm_mon=9 → tm_mday=16 → strftime이 “2026-10-16” 작성 → puts 출력.','pitfall':'strftime의 반환값을 무시하면 버퍼가 충분하지 않은 상황에서 유효하지 않은 문자열을 출력할 수 있다.','misconception':'“struct tm의 월은 1부터 12다”는 생각은 틀리다. tm_mon=0이 1월이고 9가 10월이다.','variant':'2024년 3월 1일을 표시하라.','vcode':'''#include <stdio.h>\n#include <time.h>\n\nint main(void)\n{\n    struct tm date = {0};\n    char buffer[32];\n    date.tm_year = 2024 - 1900;\n    date.tm_mon = 3 - 1;\n    date.tm_mday = 1;\n    if (strftime(buffer, sizeof buffer, "%Y-%m-%d", &date) == 0)\n        return 1;\n    puts(buffer);\n    return 0;\n}''','vout':'2024-03-01\n','vwhy':'연도는 124, 월은 2, 일은 1을 저장한다. %m과 %d는 두 자리 숫자로 표시한다.'},
{'key':'B','name':'시·분·초를 문자열로 표현','why':'로그와 보고서에는 내부 시간 구조체보다 14:07:05 같은 읽기 쉬운 표기가 필요하다.','code':'''#include <stdio.h>\n#include <time.h>\n\nint main(void)\n{\n    struct tm t = {0};\n    char text[16];\n    t.tm_hour = 14;\n    t.tm_min = 7;\n    t.tm_sec = 5;\n    if (strftime(text, sizeof text, "%H:%M:%S", &t) == 0)\n        return 1;\n    printf("clock=%s\\n", text);\n    return 0;\n}''','out':'clock=14:07:05\n','predict':'분과 초가 7, 5이면 14:7:5인가 14:07:05인가?','lines':['tm_hour, tm_min, tm_sec는 각각 시·분·초를 담는다. 여기서는 현재 시각을 읽지 않고 고정 데이터를 사용한다.','%H:%M:%S는 24시간제 시각·분·초를 두 자리로 나타낸다.','strftime이 0이면 출력하지 않고 종료한다.','printf는 형식화된 문자열이 기대대로 생성됐는지 확인하는 출력이다.'],'trace':'시=14, 분=7, 초=5 → strftime의 결과 14:07:05 → clock= 접두어와 함께 출력.','pitfall':'시간대와 일광절약시간제 정보가 없는 고정 struct tm 값만으로 절대 시점을 판정할 수는 없다.','misconception':'“strftime은 현재 시간을 자동으로 읽는다”는 오해다. 호출자가 전달한 struct tm의 내용을 형식화한다.','variant':'09:03:12를 출력하라.','vcode':'''#include <stdio.h>\n#include <time.h>\n\nint main(void)\n{\n    struct tm t = {0};\n    char text[16];\n    t.tm_hour = 9;\n    t.tm_min = 3;\n    t.tm_sec = 12;\n    if (strftime(text, sizeof text, "%H:%M:%S", &t) == 0)\n        return 1;\n    printf("clock=%s\\n", text);\n    return 0;\n}''','vout':'clock=09:03:12\n','vwhy':'%H, %M, %S가 각 숫자를 두 자리로 형식화한다.'},
{'key':'C','name':'mktime과 difftime으로 시간 간격 계산','why':'“두 시각 사이에 몇 초가 지났는가?”를 직접 시·분·초 필드의 뺄셈으로 계산하면 날짜 변경과 시간대 처리를 놓치기 쉽다.','code':'''#include <stdio.h>\n#include <time.h>\n\nint main(void)\n{\n    struct tm a = {0};\n    struct tm b = {0};\n    a.tm_year = b.tm_year = 2026 - 1900;\n    a.tm_mon = b.tm_mon = 10 - 1;\n    a.tm_mday = b.tm_mday = 16;\n    a.tm_hour = b.tm_hour = 12;\n    b.tm_min = 1;\n    b.tm_sec = 30;\n    a.tm_isdst = b.tm_isdst = -1;\n    time_t start = mktime(&a);\n    time_t end = mktime(&b);\n    if (start == (time_t)-1 || end == (time_t)-1)\n        return 1;\n    printf("elapsed=%.0f seconds\\n", difftime(end, start));\n    return 0;\n}''','out':'elapsed=90 seconds\n','predict':'12:00:00에서 12:01:30까지 몇 초인가? mktime은 지역 시간 기준인가 UTC 기준인가?','lines':['두 struct tm을 0으로 초기화한 뒤 같은 날짜와 정오를 설정한다.','b의 분과 초만 1분 30초로 지정한다. a의 분과 초는 초기화된 0이다.','tm_isdst=-1은 일광절약시간제 적용 여부를 라이브러리가 판단하도록 한다.','mktime은 지역 시간 표현을 time_t로 바꾼다. -1 반환은 변환 실패 가능성을 나타내므로 확인한다.','difftime(end,start)는 두 time_t의 경과 초를 double로 돌려준다. time_t 숫자를 직접 빼지 않는 이유는 내부 표현이 표준에서 고정되지 않았기 때문이다.'],'trace':'a=12:00:00, b=12:01:30 → 각각 mktime 변환 → difftime = 90초 → 출력.','pitfall':'일광절약시간제 전환 구간에서는 겉보기 시각 차이와 실제 경과 초가 달라질 수 있다. 이 예제는 같은 날짜의 정오를 선택했다.','misconception':'“time_t는 반드시 유닉스 epoch 이후의 정수 초다”는 가정은 ISO C만으로 보장되지 않는다.','variant':'종료 시각을 12:02:30으로 바꾸어 150초를 출력하라.','vcode':'''#include <stdio.h>\n#include <time.h>\n\nint main(void)\n{\n    struct tm a = {0};\n    struct tm b = {0};\n    a.tm_year = b.tm_year = 2026 - 1900;\n    a.tm_mon = b.tm_mon = 10 - 1;\n    a.tm_mday = b.tm_mday = 16;\n    a.tm_hour = b.tm_hour = 12;\n    b.tm_min = 2;\n    b.tm_sec = 30;\n    a.tm_isdst = b.tm_isdst = -1;\n    time_t start = mktime(&a);\n    time_t end = mktime(&b);\n    if (start == (time_t)-1 || end == (time_t)-1)\n        return 1;\n    printf("elapsed=%.0f seconds\\n", difftime(end, start));\n    return 0;\n}''','vout':'elapsed=150 seconds\n','vwhy':'2분 30초는 150초다. 두 시각 모두 같은 지역 시간대의 같은 날짜 정오이므로 일반적인 환경에서 이 차이가 성립한다.'}]},
{'id':'1.55','title':'오류를 알아채고 실패를 처리하기: <errno.h>·<assert.h>','intro':'프로그램은 성공했을 때뿐 아니라 잘못된 입력과 실패 상황에서도 예측 가능한 동작을 해야 한다. 오류를 무시한 채 계산을 계속하면 뒤늦게 더 큰 손실이 생길 수 있다.','concept':'errno는 라이브러리 함수가 오류 정보를 전달할 때 사용하는 수정 가능한 정수 lvalue 매크로이며, <errno.h>의 ERANGE는 범위를 벗어난 변환 등의 오류를 알린다. errno는 함수가 성공했다고 자동으로 0으로 돌아가지 않으므로 해당 호출 전에 0으로 초기화한 뒤 규약에 따라 검사한다. <assert.h>의 assert(조건)는 개발 중 반드시 참이어야 하는 내부 전제 조건을 점검한다. NDEBUG가 정의되면 assert 검사가 비활성화될 수 있으므로 사용자 입력의 필수 검증을 assert에 맡기면 안 된다.','items':[
{'key':'A','name':'숫자 변환의 끝 문자 검사','why':'입력 문자열 “120x”를 120으로 받아들이면 잘못된 사용자 입력이 정상 데이터처럼 저장될 수 있다.','code':'''#include <errno.h>\n#include <stdio.h>\n#include <stdlib.h>\n\nint main(void)\n{\n    const char input[] = "120x";\n    char *end = NULL;\n    errno = 0;\n    long value = strtol(input, &end, 10);\n    if (end == input || *end != '\\0' || errno == ERANGE)\n        puts("invalid");\n    else\n        printf("value=%ld\\n", value);\n    return 0;\n}''','out':'invalid\n','predict':'strtol이 120까지 읽고 x를 남긴다면 프로그램은 성공으로 처리해야 하는가?','lines':['input은 숫자처럼 시작하지만 뒤에 x가 남아 있는 문자열이다.','end는 strtol이 변환을 중단한 지점을 가리키도록 전달하는 포인터다.','errno=0은 이 호출의 범위 오류를 판정하기 위한 준비다.','end==input은 변환할 숫자가 하나도 없었음을, *end != 종료 널은 남은 문자가 있음을 뜻한다.','errno==ERANGE는 변환 결과가 long 범위를 벗어난 경우 등을 판정한다.','puts는 입력 검증 결과를 확인하는 출력이고 핵심 로직은 세 오류 조건의 결합이다.'],'trace':'input=120x → strtol은 120을 읽고 end를 x에 둔다 → *end != 종료 널이 참 → invalid 출력.','pitfall':'strtol의 반환값만 검사하면 0이 정상 입력 “0”인지 변환 실패인지 구분하지 못할 수 있다.','misconception':'“strtol이 120을 돌려줬으니 입력 전체가 유효하다”는 오개념이다. end가 문자열 끝을 가리키는지도 검사해야 한다.','variant':'입력을 “120”으로 바꾸고 성공 분기를 실행하라.','vcode':'''#include <errno.h>\n#include <stdio.h>\n#include <stdlib.h>\n\nint main(void)\n{\n    const char input[] = "120";\n    char *end = NULL;\n    errno = 0;\n    long value = strtol(input, &end, 10);\n    if (end == input || *end != '\\0' || errno == ERANGE)\n        puts("invalid");\n    else\n        printf("value=%ld\\n", value);\n    return 0;\n}''','vout':'value=120\n','vwhy':'변환 뒤 end는 문자열 종료 널을 가리키고 errno는 ERANGE가 아니므로 정상 출력한다.'},
{'key':'B','name':'범위를 벗어난 숫자 변환','why':'아주 긴 숫자를 long에 저장하려 할 때 일부 숫자가 잘려 정상값처럼 보이는 문제를 막아야 한다.','code':'''#include <errno.h>\n#include <stdio.h>\n#include <stdlib.h>\n\nint main(void)\n{\n    const char input[] = "999999999999999999999999999999999999999999999";\n    char *end = NULL;\n    errno = 0;\n    (void)strtol(input, &end, 10);\n    if (errno == ERANGE)\n        puts("out of range");\n    else if (end == input || *end != '\\0')\n        puts("invalid text");\n    else\n        puts("within range");\n    return 0;\n}''','out':'out of range\n','predict':'45자리 숫자가 일반적인 32비트 또는 64비트 long에 들어가는가? 변환 결과를 출력하기 전에 무엇을 검사해야 하는가?','lines':['입력은 일반적인 long 표현 범위를 훨씬 초과하는 양의 정수 문자열이다.','(void)strtol(...)은 반환값을 사용하지 않고 오류 상태만 조사하겠다는 의도를 표시한다.','errno는 호출 전에 0으로 초기화한다. ERANGE이면 변환 결과를 정상값으로 사용하지 않는다.','문자 오류와 범위 오류를 분리하여 진단하면 사용자에게 더 정확한 피드백을 줄 수 있다.'],'trace':'긴 숫자 → strtol이 long 범위를 벗어남 → errno=ERANGE → out of range 출력.','pitfall':'errno는 모든 라이브러리 호출이 성공할 때 0으로 초기화해 주는 자동 상태 변수가 아니다. 호출 규약을 먼저 확인한다.','misconception':'“strtol이 LONG_MAX를 돌려주면 항상 오류다”는 오해다. LONG_MAX 자체가 유효한 입력일 수도 있으므로 errno 검사와 함께 판단해야 한다.','variant':'정상 입력 “214”로 바꿔 오류 없이 변환되었는지 확인하라.','vcode':'''#include <errno.h>\n#include <stdio.h>\n#include <stdlib.h>\n\nint main(void)\n{\n    const char input[] = "214";\n    char *end = NULL;\n    errno = 0;\n    (void)strtol(input, &end, 10);\n    if (errno == ERANGE)\n        puts("out of range");\n    else if (end == input || *end != '\\0')\n        puts("invalid text");\n    else\n        puts("within range");\n    return 0;\n}''','vout':'within range\n','vwhy':'214는 C11이 보장하는 최소 long 범위 안에도 있으므로 ERANGE가 발생하지 않고 문자열 전체가 소비된다.'},
{'key':'C','name':'assert로 내부 전제 조건 검증','why':'함수 작성자가 내부 배열 포인터가 유효하고 길이가 양수라는 전제를 놓치면 버그가 조용히 확산될 수 있다.','code':'''#include <assert.h>\n#include <stddef.h>\n#include <stdio.h>\n\nint sum_positive(const int *data, size_t count)\n{\n    assert(data != NULL);\n    assert(count > 0);\n    int total = 0;\n    for (size_t i = 0; i < count; ++i)\n        total += data[i];\n    return total;\n}\n\nint main(void)\n{\n    const int numbers[] = {1, 2, 3};\n    int result = sum_positive(numbers, sizeof numbers / sizeof numbers[0]);\n    printf("sum=%d\\n", result);\n    return 0;\n}''','out':'sum=6\n','predict':'count가 0이라면 이 함수의 assert가 무엇을 검사하는가? NDEBUG가 정의되면 어떻게 되는가?','lines':['함수 매개변수 data는 읽기 전용 정수 배열을 가리키고 count는 원소 개수다.','assert(data != NULL)과 assert(count > 0)은 개발 중 내부 호출 계약을 검사한다.','반복문은 count개 원소를 누적해 total을 계산한다.','main의 sizeof numbers / sizeof numbers[0]은 배열 전체 바이트를 원소 하나의 바이트로 나누어 원소 수 3을 구한다.','printf는 학습용 확인 출력이고 핵심 기능은 sum_positive의 누적 합 계산이다.'],'trace':'numbers={1,2,3}, count=3 → total: 0→1→3→6 → sum=6 출력.','pitfall':'assert로 사용자 입력을 검증하지 않는다. NDEBUG가 정의되면 assert 검사가 빠질 수 있기 때문이다.','misconception':'“assert는 예외 처리이므로 배포 버전에서도 항상 실행된다”는 생각은 틀리다. NDEBUG에 의해 비활성화될 수 있다.','variant':'배열을 {2,4,6}으로 바꿔 sum=12를 출력하라.','vcode':'''#include <assert.h>\n#include <stddef.h>\n#include <stdio.h>\n\nint sum_positive(const int *data, size_t count)\n{\n    assert(data != NULL);\n    assert(count > 0);\n    int total = 0;\n    for (size_t i = 0; i < count; ++i)\n        total += data[i];\n    return total;\n}\n\nint main(void)\n{\n    const int numbers[] = {2, 4, 6};\n    int result = sum_positive(numbers, sizeof numbers / sizeof numbers[0]);\n    printf("sum=%d\\n", result);\n    return 0;\n}''','vout':'sum=12\n','vwhy':'total은 0→2→6→12로 바뀐다. 배열 원소 수는 여전히 3이다.'}]}
]

EXTRA={
'1.53':{'guided':'빈칸 채우기: 음수 x의 0 방향 버림을 구하려면 ______(x), 아래 방향 내림을 구하려면 ______(x)를 사용한다.','guided_answer':'trunc(x), floor(x). 음수 -2.7에서 -2와 -3으로 결과가 다르다.','independent':'독립 구현: 실수형 세 값 중 0과의 거리가 가장 큰 값을 찾아 크기만 출력하는 함수를 설계하라. 절댓값 비교를 왜 먼저 해야 하는지 설명하라.','independent_answer':'fabs(a), fabs(b), fabs(c)로 크기를 비교하고 최댓값을 선택한다. 음수의 부호를 그대로 비교하면 실제 거리 순서와 달라진다.','debug':'단계별 디버깅: #include <math.h>가 있는데도 sqrt의 링크 단계에서 undefined reference가 난다. 먼저 선언 문제인지 링크 문제인지 구분하고 GCC/Clang 명령 끝에 -lm을 추가해 보라.','review':'인출 복습: 1.49절의 문자열 길이 검사와 연결해, 문자열을 숫자로 바꾼 후 범위 검사를 생략하면 어떤 오류가 생기는지 말해 보라.'},
'1.54':{'guided':'코드 순서 구성: ① tm_year/tm_mon 설정 ② strftime 결과 검사 ③ struct tm 초기화 ④ 출력. 올바른 순서를 쓰라.','guided_answer':'③→①→②→④. 초기화하지 않은 구조체를 먼저 읽거나, 형식화 실패 후 출력하면 안 된다.','independent':'독립 구현: 고정된 두 날짜를 각각 struct tm으로 만들고 mktime/difftime으로 시간 차이를 계산하라. 실패 처리와 지역 시간대 주석을 포함하라.','independent_answer':'각 구조체를 0 초기화하고 연도에서 1900, 월에서 1을 빼서 저장한다. tm_isdst=-1 후 mktime 두 번의 반환을 검사하고 difftime(end,start)을 출력한다. 지역 시간대와 DST 전환 가능성을 문서화한다.','debug':'단계별 디버깅: 2026년 10월을 tm_year=2026, tm_mon=10으로 넣은 프로그램이 엉뚱한 날짜를 출력한다. 컴파일 오류가 아니라 논리 오류임을 분류하고 126, 9로 수정하라.','review':'인출 복습: 1.38절 구조체 멤버 접근(. 연산자)과 1.46절 파일 출력에 사용한 형식화 출력의 역할을 설명하라.'},
'1.55':{'guided':'빈칸 채우기: strtol 직전에 ______ = 0; 변환 후 ______ == ERANGE인지 검사한다.','guided_answer':'errno, errno. 단, 변환된 전체 문자열 여부도 end 포인터로 별도 확인한다.','independent':'독립 구현: 문자열 3개 {"17", "-3", "4x"}를 차례로 strtol로 검사하고 유효한 값만 합산하는 프로그램을 작성하라. 입력 오류를 건너뛰는 이유를 설명하라.','independent_answer':'각 문자열마다 errno=0, end 검사, ERANGE 검사를 수행한다. 유효한 17과 -3만 더하면 14다. 4x를 4로 받아들이면 입력 검증 요구를 위반한다.','debug':'단계별 디버깅: assert(n > 0)를 사용자 입력 검사로 사용한 프로그램을 -DNDEBUG로 빌드하면 검사 없이 실행된다. 오류의 종류는 계약 오해다. 사용자 입력은 if로 확인하고 assert는 내부 불변식에 한정하라.','review':'인출 복습: 1.44절의 strtol, 1.43절의 malloc 실패 검사, 1.46절의 파일 열기 실패 검사는 어떤 공통 원칙을 공유하는가?'}
}

# Validate sources before making the document.
checks=[]
for section in S:
 for item in section['items']:
  for suffix,code,out in [('main',item['code'],item['out']),('variation',item['vcode'],item['vout'])]:
   name=f"ch{section['id'].replace('.','_')}_{item['key'].lower()}_{suffix}.c"
   src=CODES/name; src.write_text(code+'\n',encoding='utf-8')
   for compiler in ('gcc','clang'):
    exe=ROOT/f'.test_{name}_{compiler}'
    cmd=[compiler,*FLAGS,str(src),'-lm','-o',str(exe)]
    comp=subprocess.run(cmd,text=True,capture_output=True)
    if comp.returncode:
     raise RuntimeError(f'{name} {compiler} compile: {comp.stderr}')
    run=subprocess.run([str(exe)],text=True,capture_output=True,timeout=5,env={**os.environ,'TZ':'Asia/Seoul'})
    if run.returncode!=0 or run.stdout!=out or run.stderr:
     raise RuntimeError(f'{name} {compiler} mismatch: expected {out!r} actual {run.stdout!r}, stderr={run.stderr!r}, rc={run.returncode}')
    checks.append({'source':name,'compiler':compiler,'compile_command':' '.join(cmd),'exit_code':run.returncode,'stdout':run.stdout,'expected_stdout':out,'passed':True})
    exe.unlink()

D=Document();sec=D.sections[0];sec.top_margin=Cm(2.15);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
styles=D.styles
base=styles['Normal'];base.font.name='Noto Sans CJK KR';base.font.size=Pt(9.2);base.font.color.rgb=RGBColor(35,49,58);base.paragraph_format.space_after=Pt(5);base.paragraph_format.line_spacing=1.24
for nm,size,color in [('Title',19,RGBColor(15,95,104)),('Heading 1',14,RGBColor(0,112,121)),('Heading 2',11.4,RGBColor(14,97,107))]:
 s=styles[nm];s.font.name='Noto Sans CJK KR';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=color;s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
for nm in ['Learning Label','Editable C Code']:
 if nm not in styles: styles.add_style(nm,WD_STYLE_TYPE.PARAGRAPH)
label=styles['Learning Label'];label.font.name='Noto Sans CJK KR';label.font.size=Pt(9.2);label.font.bold=True;label.font.color.rgb=RGBColor(0,109,116);label.paragraph_format.space_before=Pt(6);label.paragraph_format.space_after=Pt(3);label.paragraph_format.keep_with_next=True
code_style=styles['Editable C Code'];code_style.font.name='DejaVu Sans Mono';code_style.font.size=Pt(7.7);code_style.font.color.rgb=RGBColor(20,46,55);code_style.paragraph_format.space_after=Pt(0);code_style.paragraph_format.line_spacing=1.0

def P(s,style=None):
 p=D.add_paragraph(style=style);p.add_run(s);return p

def L(t):P(t,'Learning Label')
def C(t):
 for line in t.rstrip('\n').split('\n'):
  p=P(line,'Editable C Code');p.paragraph_format.keep_together=True

def H(t,level=2):D.add_heading(t,level=level)

def add_example(section,item):
 H(f"{section['id']}-{item['key']}  {item['name']}")
 L('왜 필요한가 · 언제 쓰는가');P(item['why'])
 L('전체 C11 코드 — 그대로 컴파일 가능한 독립 프로그램');C(item['code'])
 L('실행하기 전에 스스로 예측');P(item['predict'])
 L('문장별 해설 — 왜 이 문장을 쓰는가')
 for v in item['lines']:P('• '+v)
 L('값과 실행 흐름 추적');P(item['trace'])
 L('실제 출력 — 예측과 비교');C(item['out'].rstrip())
 L('흔한 실수와 그 원인');P(item['pitfall'])
 L('오개념 진단 · 반례');P(item['misconception'])
 L('예제를 바꾸며 실력 다지기 — 바로 여기서 변형');P(item['variant'])
 L('인접 정답: 완성 코드');C(item['vcode'])
 L('변형 실행 결과');C(item['vout'].rstrip())
 L('정답 근거');P(item['vwhy'])

D.add_heading('한국어 C언어 종합교과서 — 이해 중심 정식 개정판',0)
P('제1장 1.53~1.55절 | 수학 함수 · 시간 처리 · 오류 진단')
P('새 정식 개정판의 1.52절 다음 원고. 과거 거부 초안 및 구 제1~11장 초안을 본문으로 재사용하지 않았다.')
P('학습 순서: 필요성 → 새 개념의 정의 → 완전한 C11 코드 → 실행 전 예측 → 문장별 이유 → 값 추적 → 실제 출력 → 오류·오개념 → 바로 인접한 변형 실습·정답.')
P('실습은 전역 객체를 사용하지 않는다. GCC·Clang의 수학 라이브러리 링크에는 -lm을 사용했다. time_t 내부 표현·지역 시간·일광절약시간제는 환경에 따라 다를 수 있음을 별도로 표시한다.')
P('원서 Express·K&R의 전체 LAB·Exercise·Programming·Mini Project·Advanced Topic·부록 학습목표 전수 대조는 아직 완료되지 않았다. 아래 절의 검증 완료는 원서 전수 대조 완료를 뜻하지 않는다.')
for section in S:
 H(f"{section['id']}절  {section['title']}",1)
 P(section['intro'])
 L('새 개념을 먼저 이해하기');P(section['concept'])
 for item in section['items']:add_example(section,item)
 ex=EXTRA[section['id']]
 H(f"{section['id']}절 누적 실습 · 점진적 독립성")
 L('부분 완성 / 코드 순서 구성 — 도움을 남긴 단계');P(ex['guided']);L('바로 확인하는 정답과 이유');P(ex['guided_answer'])
 L('도움 없는 독립 코드 작성 — 다음 단계');P(ex['independent']);L('풀이 방향 및 검증 기준');P(ex['independent_answer'])
 L('단계별 디버깅: 문법·링크·논리·계약 오류 구별');P(ex['debug'])
 L('이전 절을 떠올리는 간격 복습');P(ex['review'])

H('학습 경계와 이 절 이후의 심화 연결',1)
P('수학: IEEE 754 부동소수점의 반올림·NaN·무한대·예외 플래그는 <fenv.h>, <float.h> 및 플랫폼 구현 차이와 연결된다. 여기서는 함수 사용의 기초만 검증했으며, 정밀도 분석 전체를 완료했다고 표시하지 않는다.')
P('시간: clock()은 CPU 시간이며 실제 경과 시간(wall-clock time)과 다르다. time()·localtime()·gmtime()의 현재 시간 예제는 실행 환경과 시간대에 따라 출력이 달라지므로 고정 출력 비교가 필요한 이 회차에서는 별도의 심화 실습으로 남긴다.')
P('오류: assert는 개발 중 내부 계약을 확인하는 도구이고, 사용자 입력 검증·파일 오류 복구·동적 메모리 실패 처리의 대체물이 아니다. C 표준 errno와 운영체제 오류 코드를 혼동하지 않는다.')
P('다음 집필 예정: 1.56절 이후에서 표준 라이브러리의 남은 영역과 원서 미수록 학습목표를 추적해 누락을 보충한다. 최종 절 수는 미확정이다.')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('C언어 종합교과서 | 이해 중심 정식 개정판 · 1.53~1.55')
D.save(ROOT/TITLE)

# Check that all complete code blocks appear in DOCX text and source files exactly.
from docx import Document as ReadDoc
r=ReadDoc(ROOT/TITLE)
alltext='\n'.join(p.text for p in r.paragraphs)
match=[]
for section in S:
 for item in section['items']:
  for kind,code in [('main',item['code']),('variation',item['vcode'])]:
   ok=code in alltext
   match.append({'section':section['id'],'example':item['key'],'kind':kind,'match':ok})
   if not ok: raise RuntimeError(f'DOCX code mismatch: {section["id"]}-{item["key"]}-{kind}')

ZIP=ROOT/(STEM+'_sources.zip')
with zipfile.ZipFile(ZIP,'w',compression=zipfile.ZIP_DEFLATED) as z:
 for p in sorted(CODES.glob('*.c')):z.write(p,'sources/'+p.name)
 z.writestr('README.txt','C11 독립 예제 9개 + 인접 변형 정답 9개. GCC/Clang: -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 source.c -lm -o program\n')

now=datetime.now(timezone(timedelta(hours=9))).isoformat(timespec='seconds')
def dump(name,obj):
 (ROOT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

VERIFY=STEM+'_c11_verification.json'
dump(VERIFY,{'timestamp':now,'compiler_flags':FLAGS,'link_libraries':['m'],'programs':18,'compiler_runs':len(checks),'passed':len(checks),'failed':0,'checks':checks,'note':'mktime/difftime local timezone test run with TZ=Asia/Seoul; DST and timezone may affect other environments.'})
dump(STEM+'_code_match.json',{'source_files':18,'docx_code_blocks_matched':len(match),'all_match':all(x['match'] for x in match),'details':match})
dump('official_rewrite_progress_1_1-1_55.json',{'reported_completed_range':'1.1~1.55','new_sections':['1.53','1.54','1.55'],'next_section':'1.56','progress_percent':None,'source_objective_full_audit':False,'stage':'math/time/errno/assert new manuscript and C11 test verified','github_delivery':'pending','library_delivery':'pending','timestamp':now})
dump('official_rewrite_objective_tracking_1_1-1_55.json',{'full_audit_complete':False,'source_files':[{'name':'c언어 압축 (1)(1).pdf','current_access':'not found in accessible Library search/root list this run','full_item_audit':False},{'name':'C_Programming.pdf','current_access':'not found in accessible Library search/root list this run','full_item_audit':False},{'name':'C Programming Language - 2nd Edition (OCR).pdf','current_access':'not found in accessible Library search/root list this run','full_item_audit':False},{'name':'C언어_통합교과서_전체수록범위_확정안.md','current_access':'not found in accessible Library search/root list this run','full_item_audit':False}],'new_sections':[{'section':'1.53','objectives':'<math.h> sqrt/hypot/floor/ceil/trunc/pow/fabs; domain and precision warnings','original_source_crosscheck':'unverified this run'},{'section':'1.54','objectives':'<time.h> struct tm/strftime/mktime/difftime; time_t and local timezone caveats','original_source_crosscheck':'unverified this run'},{'section':'1.55','objectives':'<errno.h> errno/ERANGE and strtol full input validation; <assert.h> internal contract','original_source_crosscheck':'unverified this run'}],'c11_verified_runs':len(checks),'previous_sections':'1.1~1.52 reported by prior run; not re-audited here'})
dump(STEM+'_source_inspection.json',{'timestamp':now,'Library_folder':'/C언어_통합교과서_제작/','latest_previous_manifest':'official_rewrite_manifest_1_50-1_52.json read','latest_previous_objective_tracking':'official_rewrite_objective_tracking_1_1-1_52.json read','root_source_pdf_search':'no matching file found; source item-level comparison not claimed','source_full_audit_complete':False})
# Manifest does not hash itself; all other deliverables are hashed.
files=[p for p in ROOT.iterdir() if p.is_file() and p.name not in ('official_rewrite_manifest_1_53-1_55.json',)]
dump('official_rewrite_manifest_1_53-1_55.json',{'run':'ch01_1.53-1.55','timestamp':now,'files':[{'name':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(files)],'source_full_audit_complete':False,'github_delivery':'pending','library_delivery':'pending','next_section':'1.56'})
print(json.dumps({'docx':str(ROOT/TITLE),'source_zip':str(ZIP),'verification':str(ROOT/VERIFY),'programs':18,'runs':len(checks),'docx_code_matches':len(match)},ensure_ascii=False))
