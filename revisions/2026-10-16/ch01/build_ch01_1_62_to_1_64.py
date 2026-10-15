#!/usr/bin/env python3
"""ISO C11 한국어 이해 중심 교과서 1.62-1.64 신규 원고·검증·배포자료 생성."""
from __future__ import annotations
from pathlib import Path
from datetime import datetime, timezone, timedelta
from collections import Counter
import hashlib, json, os, re, subprocess, zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
SOURCES=ROOT/'sources'
SOURCES.mkdir(parents=True,exist_ok=True)
DOCX_NAME='C언어_종합교과서_이해중심_정식개정판_제1장_1.62-1.64.docx'
STEM='official_ch01_1_62-1_64'
FLAGS=['-std=c11','-Wall','-Wextra','-Wpedantic','-Werror','-O0']
NOW=datetime.now(timezone(timedelta(hours=9))).isoformat(timespec='seconds')

SECTIONS=[
 dict(number='1.62',title='로캘과 문자열 비교: <locale.h>를 정확히 사용하기',
  need='숫자 12.5를 읽는 프로그램이 어떤 나라에서는 12,5를 만나고, 이름을 정렬할 때에도 지역별 문자 순서가 다를 수 있다. 문자 데이터의 해석·비교 규칙을 무조건 내 컴퓨터의 설정과 같다고 생각하면 입력 검증과 정렬 결과가 흔들린다. 표준 C는 로캘(locale)이라는 설정을 통해 숫자 표기와 문자 분류·정렬 규칙을 다룬다. 이 절에서는 출력이 재현되도록 표준으로 반드시 존재하는 "C" 로캘을 명시한다.',
  concepts=[
   '로캘은 문자 분류, 문자열 정렬, 숫자 표기 등 일부 표준 라이브러리 동작을 제어하는 설정이다. `setlocale(LC_ALL, "C")`는 모든 범주를 표준 "C" 로캘로 설정한다. `setlocale`은 성공하면 문자열 포인터, 실패하면 NULL을 반환하므로 결과를 확인한다. "C" 로캘은 모든 호스티드 C 구현에 존재한다.',
   '`LC_NUMERIC`은 숫자 변환·소수점 기호 관련 범주이고 `LC_COLLATE`는 `strcoll`과 `strxfrm`의 정렬 범주다. `LC_CTYPE`은 문자 분류·대소문자 변환 및 멀티바이트 변환에 영향을 준다. 다른 범주를 섞어 바꾸면 프로그램 전체의 라이브러리 상태가 달라질 수 있으므로 설정 시점을 분명히 해야 한다.',
   '`localeconv()`는 `struct lconv` 구조체를 가리키는 포인터를 반환한다. `decimal_point`와 `thousands_sep`는 숫자 표기 관련 문자열이다. 이 포인터가 가리키는 데이터는 라이브러리가 관리하며 임의로 수정하면 안 된다. `strcoll`은 현재 로캘의 정렬 규칙에 따른 비교이고 `strxfrm`은 비교 가능한 변환 키를 만든다.',
   '`strtod`는 문자열의 숫자 부분을 `double`로 변환하고 `endptr`로 어디까지 읽었는지 알려준다. `strtod`의 소수점 해석은 현재 `LC_NUMERIC` 범주의 영향을 받는다. 변환 성공 여부와 남은 문자를 검사하지 않고 값만 쓰면 잘못된 입력을 놓친다.'
  ],
  cases=[
   dict(title='표준 C 로캘을 설정하고 소수점 규칙 확인하기',why='파일 형식이나 자동 채점처럼 동일한 소수점 규칙이 필요한 환경에서 로캘에 따른 흔들림을 피한다.',
    code=r'''#include <locale.h>
#include <stdio.h>
#include <string.h>

int main(void)
{
    if (setlocale(LC_ALL, "C") == NULL) {
        return 1;
    }
    struct lconv *format = localeconv();
    int decimal_ok = strcmp(format->decimal_point, ".") == 0;
    int no_group = format->thousands_sep[0] == '\0';
    printf("decimal=%d grouping_empty=%d\n", decimal_ok, no_group);
    return 0;
}''',expected='decimal=1 grouping_empty=1\n',
    predict='"C" 로캘의 소수점 문자열은 무엇인가? `thousands_sep[0]`은 어떤 값인가? 두 조건식은 참을 나타내는 1을 출력할까?',
    steps=['`<locale.h>`는 setlocale·localeconv·LC_ALL을, `<string.h>`는 strcmp를, `<stdio.h>`는 printf를 선언한다.',
           '`setlocale`의 반환값을 NULL과 비교한다. 설정이 실패하면 뒤의 비교를 신뢰할 수 없으므로 main에서 1을 반환한다.',
           '`localeconv`는 현재 로캘 정보를 담은 라이브러리 관리 구조체의 주소를 돌려준다. 구조체 멤버 접근에 `->`를 쓰는 이유는 format이 구조체 포인터이기 때문이다.',
           '`strcmp`는 두 문자열의 내용이 같은지 검사한다. 소수점이 "."인지 확인하고, 천 단위 구분 문자열이 비어 있는지도 검사한다.',
           '`printf`는 학습 확인용 출력이다. 로캘 설정과 정보 조회가 실제 동작의 핵심이다. 사용자 정의 전역 객체는 없다.'],
    trace='처음에는 실행 환경의 로캘을 모른다 → setlocale 성공 후 모든 범주가 "C"가 된다 → decimal_point="." → strcmp 결과 0 → decimal_ok=1 → thousands_sep[0]은 종료 문자 → no_group=1.',
    mistake='`localeconv()->decimal_point`를 수정해 소수점 기호를 바꾸려 하면 안 된다. 라이브러리 소유 데이터를 임의로 변경하는 대신 지원되는 로캘을 선택해야 한다.',
    misconception='“`setlocale`은 출력할 때만 적용되는 지역 변수다”는 틀리다. 로캘은 라이브러리의 전역적인 설정이며 다른 관련 함수에도 영향을 준다. 다중 스레드에서 변경하면 특히 주의해야 한다.',
    exercise='`LC_NUMERIC`만 "C"로 설정하고 decimal_point가 점인지 검사하여 numeric_c=1을 출력하라.',
    variant=r'''#include <locale.h>
#include <stdio.h>
#include <string.h>

int main(void)
{
    if (setlocale(LC_NUMERIC, "C") == NULL) {
        return 1;
    }
    struct lconv *format = localeconv();
    printf("numeric_c=%d\n", strcmp(format->decimal_point, ".") == 0);
    return 0;
}''',variant_out='numeric_c=1\n',reason='숫자 범주만 바꿔도 소수점 표기는 "C" 규칙을 따른다. LC_COLLATE 같은 다른 범주를 의도 없이 바꾸지 않는다.'),
   dict(title='strcoll과 strcmp, strxfrm의 역할 구별하기',why='사용자 이름이나 항목을 정렬할 때 문자 코드 순서와 로캘별 정렬 규칙을 구별해야 한다.',
    code=r'''#include <locale.h>
#include <stdio.h>
#include <string.h>

int main(void)
{
    if (setlocale(LC_COLLATE, "C") == NULL) {
        return 1;
    }
    const char *left = "apple";
    const char *right = "banana";
    int a = strcoll(left, right);
    int b = strcmp(left, right);
    int same_sign = (a < 0) == (b < 0) &&
                    (a == 0) == (b == 0) &&
                    (a > 0) == (b > 0);
    printf("same_sign=%d\n", same_sign);
    return 0;
}''',expected='same_sign=1\n',
    predict='"C" 로캘에서 strcoll과 strcmp의 결과는 정렬 방향이 같을까? 두 함수가 반환하는 정수의 정확한 절댓값도 같아야 할까?',
    steps=['`LC_COLLATE`는 문자열 정렬 규칙 범주다. 표준 "C" 로캘에서는 strcoll의 비교 관계가 strcmp와 같다.',
           '`left`와 `right`는 읽기 전용 문자열을 가리키는 지역 포인터다. 문자열 리터럴을 수정하지 않는다.',
           '`strcoll`은 로캘 규칙으로, `strcmp`는 문자열의 바이트 값을 순서대로 비교한다. 두 함수 모두 음수·0·양수의 부호에 의미가 있다.',
           '`same_sign`은 정확한 반환 정수를 비교하지 않고 세 가지 부호 관계가 같은지 검사한다. 이것이 표준 함수의 계약에 맞는 비교 방법이다.',
           '출력은 정렬 결과를 눈으로 확인하기 위한 부분이고, 비교의 본질은 두 함수 호출과 부호 검사다.'],
    trace='"C" 정렬 규칙을 선택 → strcoll(left,right)와 strcmp(left,right)의 부호가 일치 → 세 비교 결과가 각각 같음 → same_sign=1.',
    mistake='`if (strcmp(a,b)==-1)`처럼 반환값이 정확히 -1이라고 가정하지 않는다. 음수인지 검사한다.',
    misconception='“strcoll은 언제나 한국어 사전 순서를 보장한다”는 틀리다. 현재 로캘과 구현이 제공하는 정렬 규칙에 따라 달라지며 "C" 로캘은 한국어 자연어 정렬을 제공하지 않는다.',
    exercise='두 단어를 "pear", "peach"로 바꾸고 strcoll과 strcmp의 부호가 같은지 출력하라.',
    variant=r'''#include <locale.h>
#include <stdio.h>
#include <string.h>

int main(void)
{
    if (setlocale(LC_COLLATE, "C") == NULL) {
        return 1;
    }
    int a = strcoll("pear", "peach");
    int b = strcmp("pear", "peach");
    printf("same_sign=%d\n", (a < 0) == (b < 0) &&
           (a == 0) == (b == 0) && (a > 0) == (b > 0));
    return 0;
}''',variant_out='same_sign=1\n',reason='어떤 두 문자열이든 "C" 로캘의 strcoll과 strcmp는 같은 정렬 관계를 준다. 반환 정수의 크기는 비교하지 않는다.'),
   dict(title='strtod의 소수점 해석과 입력 종료 위치 검증하기',why='설정 파일의 숫자와 단위를 구별할 때 문자열 전체를 숫자로 잘못 간주하지 않도록 한다.',
    code=r'''#include <locale.h>
#include <stdio.h>
#include <stdlib.h>

int main(void)
{
    if (setlocale(LC_NUMERIC, "C") == NULL) {
        return 1;
    }
    char text[] = "12.5kg";
    char *end = NULL;
    double value = strtod(text, &end);
    printf("value=%.1f suffix=%s\n", value, end);
    return 0;
}''',expected='value=12.5 suffix=kg\n',
    predict='strtod는 "12.5kg" 전체를 읽을까? end는 문자열의 어느 문자부터 가리킬까? 값은 얼마가 될까?',
    steps=['`setlocale(LC_NUMERIC,"C")`는 숫자 표기를 점 소수점으로 고정한다. `strtod`는 `<stdlib.h>`에 선언된다.',
           '`text`는 수정 가능한 지역 문자 배열이다. `end`는 아직 아무 위치도 가리키지 않도록 NULL로 초기화한다.',
           '`strtod(text,&end)`는 숫자로 해석할 수 있는 앞부분을 읽어 12.5를 돌려준다. 동시에 end가 남은 문자열의 첫 문자 k를 가리키게 한다.',
           '`printf`의 %s는 end에서 시작하는 널 종료 문자열 "kg"를 출력한다. 이 코드는 단위를 허용하는 예제이며, 숫자만 허용하는 입력이라면 *end가 종료 문자인지 추가 검사해야 한다.',
           '이 출력은 입력 파싱 결과를 관찰하기 위한 것이며, 실제 입력 검증에서는 errno와 범위·끝 위치를 더 점검한다.'],
    trace='text[0..3]의 "12.5"가 숫자로 변환 → value=12.5 → end는 text[4]의 k를 가리킴 → suffix="kg".',
    mistake='숫자 변환 함수의 반환값만 보고 입력 전체가 숫자라고 판단하면 "12.5kg"처럼 뒤에 붙은 문자를 놓친다.',
    misconception='“`strtod`는 어떤 로캘에서든 점(.)만 소수점으로 읽는다”는 틀리다. LC_NUMERIC에 따라 해석이 달라질 수 있다.',
    exercise='입력 문자열을 "7.25m"로 바꾸고 value=7.25 suffix=m을 출력하라.',
    variant=r'''#include <locale.h>
#include <stdio.h>
#include <stdlib.h>

int main(void)
{
    if (setlocale(LC_NUMERIC, "C") == NULL) {
        return 1;
    }
    char text[] = "7.25m";
    char *end = NULL;
    double value = strtod(text, &end);
    printf("value=%.2f suffix=%s\n", value, end);
    return 0;
}''',variant_out='value=7.25 suffix=m\n',reason='변환은 숫자 부분에서 멈춘다. %.2f는 소수점 아래 두 자리로 확인 출력한다.'),
   dict(title='strxfrm으로 비교 키를 만들고 정렬 관계 확인하기',why='같은 로캘에서 많은 문자열을 반복 비교할 때 미리 만든 변환 키로 정렬 관계를 검사하는 원리를 배운다.',
    code=r'''#include <locale.h>
#include <stdio.h>
#include <string.h>

int main(void)
{
    if (setlocale(LC_COLLATE, "C") == NULL) {
        return 1;
    }
    char first[32];
    char second[32];
    size_t n = strxfrm(first, "pear", sizeof first);
    size_t m = strxfrm(second, "peach", sizeof second);
    if (n >= sizeof first || m >= sizeof second) {
        return 1;
    }
    int a = strcoll("pear", "peach");
    int b = strcmp(first, second);
    printf("keys_ok=1 same_order=%d\n",
           (a < 0) == (b < 0) && (a == 0) == (b == 0) &&
           (a > 0) == (b > 0));
    return 0;
}''',expected='keys_ok=1 same_order=1\n',
    predict='strxfrm의 반환값이 버퍼 크기보다 작으면 변환 키가 널 종료될까? 두 변환 키를 strcmp로 비교하면 strcoll과 같은 정렬 관계일까?',
    steps=['`strxfrm`은 현재 LC_COLLATE에 맞춰 정렬 비교용 문자열을 만든다. 변환 결과가 반드시 원본과 같은 바이트열일 필요는 없다.',
           '`sizeof first`와 `sizeof second`는 각 배열의 바이트 용량이다. strxfrm 반환값이 용량 이상이면 출력이 널 종료된다고 가정하지 않고 중단한다.',
           '`strcoll`로 원본 문자열의 정렬 관계를 얻고 `strcmp`로 변환 키의 정렬 관계를 얻는다.',
           '부호 관계를 비교해 두 정렬 결과가 같음을 확인한다. keys_ok=1은 위에서 두 용량 검사를 통과했음을 나타내는 학습 출력이다.'],
    trace='C 정렬 규칙 선택 → 두 변환 키 생성 → 각 반환 길이 < 32 → 원본 strcoll과 키 strcmp의 부호 관계 일치 → same_order=1.',
    mistake='strxfrm의 반환값은 종료 문자를 제외한 필요한 길이이므로 `n == sizeof buffer`인 경우에도 버퍼가 충분하지 않다.',
    misconception='“strxfrm 결과는 원본 문자열과 언제나 똑같다”는 틀리다. 로캘에 따라 비교용 키가 달라질 수 있다.',
    exercise='문자열을 "red"와 "blue"로 바꾸어 키 생성과 정렬 관계 일치를 검증하라.',
    variant=r'''#include <locale.h>
#include <stdio.h>
#include <string.h>

int main(void)
{
    if (setlocale(LC_COLLATE, "C") == NULL) {
        return 1;
    }
    char first[32];
    char second[32];
    size_t n = strxfrm(first, "red", sizeof first);
    size_t m = strxfrm(second, "blue", sizeof second);
    if (n >= sizeof first || m >= sizeof second) {
        return 1;
    }
    int a = strcoll("red", "blue");
    int b = strcmp(first, second);
    printf("keys_ok=1 same_order=%d\n",
           (a < 0) == (b < 0) && (a == 0) == (b == 0) &&
           (a > 0) == (b > 0));
    return 0;
}''',variant_out='keys_ok=1 same_order=1\n',reason='strxfrm으로 얻은 키를 strcmp로 비교하면 현재 로캘의 strcoll과 같은 정렬 관계를 갖는다.')
  ],
  guided='부분 완성: `setlocale(________, "C")`에서 숫자 변환에만 적용하려면 어떤 로캘 범주를 넣어야 할까?',
  guided_answer='`LC_NUMERIC`. 전체 범주를 변경하는 LC_ALL과 구별한다.',
  debug='1단계(컴파일): setlocale을 호출하면서 `<locale.h>`를 누락한 코드를 고친다. 2단계(논리): strcmp의 결과를 -1과 직접 비교하는 조건을 `< 0`으로 고친다. 3단계(입력 검증): strtod 뒤에서 end==text(숫자를 전혀 못 읽음), *end(남은 문자), errno(범위 오류)를 순서대로 점검한다.',
  review='1.44절의 strtol 검증을 떠올려 보자. end 포인터가 필요한 이유와 문자열 끝까지 검사해야 하는 경우를 설명하라. 1.47절의 sscanf와도 입력 검증 차이를 비교해 보자.',
  independent_prompt='힌트 없이 작성: 숫자만 허용하는 문자열 "18.75"를 strtod로 변환하고, 남은 문자가 없을 때만 `ok=1 value=18.75`를 출력하라. "C" 숫자 로캘을 설정하라.',
  independent_code=r'''#include <locale.h>
#include <stdio.h>
#include <stdlib.h>

int main(void)
{
    if (setlocale(LC_NUMERIC, "C") == NULL) {
        return 1;
    }
    char input[] = "18.75";
    char *end = NULL;
    double value = strtod(input, &end);
    int ok = end != input && *end == '\0';
    printf("ok=%d value=%.2f\n", ok, value);
    return 0;
}''',independent_out='ok=1 value=18.75\n',independent_reason='end!=input은 적어도 하나의 숫자를 읽었다는 뜻이고 *end가 종료 문자이면 뒤에 쓰레기 문자가 없다. 실제 사용자 입력은 범위·유한성까지 추가 검사해야 한다.'),
 dict(number='1.63',title='와이드 문자와 멀티바이트 변환: <wchar.h>',
  need='한국어 한 글자가 UTF-8에서 여러 바이트를 차지하는 환경에서 `strlen`을 글자 수라고 생각하면 문자 개수와 바이트 개수를 혼동하게 된다. 표준 C에는 와이드 문자와 멀티바이트 문자열을 다루는 별도의 자료형·함수가 있다. 다만 와이드 문자가 곧 UTF-16 또는 유니코드 코드 포인트라는 뜻은 아니다. 이 절은 구현 의존성을 피하기 위해 "C" 로캘의 기본 실행 문자만 사용한다.',
  concepts=[
   '`wchar_t`는 와이드 문자 코드 단위를 담는 정수형이며 `L\'A\'`는 와이드 문자 상수, `L"ABC"`는 널 종료 와이드 문자열이다. `wchar_t`의 크기·부호·인코딩은 구현에 따라 다르므로 UTF-16/UTF-32와 무조건 동일시하지 않는다.',
   '`wcslen`은 와이드 문자열의 종료 문자 전까지의 wchar_t 원소 수를 세고 `wcscmp`는 와이드 문자열을 비교한다. `wmemmove`는 와이드 원소가 겹치는 구간도 이동할 수 있다. 단, 버퍼 크기와 널 종료는 프로그래머가 관리한다.',
   '`mbstowcs`는 멀티바이트 문자열을 와이드 문자열로, `wcstombs`는 그 반대로 변환한다. 변환은 현재 로캘의 문자 인코딩 규칙을 따른다. 반환값 `(size_t)-1`은 변환 오류를 나타내며, 충분한 공간이 없으면 종료 문자가 저장되지 않을 수 있다.',
   '멀티바이트 인코딩은 여러 바이트가 하나의 문자를 나타내는 규칙이다. UTF-8이 흔하지만 표준 C의 "C" 로캘이 UTF-8을 뜻하지는 않는다. `mbstate_t`와 `mbrtowc` 같은 상태 기반 변환은 뒤 절에서 별도로 확장한다.'
  ],
  cases=[
   dict(title='wcslen으로 와이드 원소 수 세고 wcscmp로 비교하기',why='와이드 문자열을 다루는 API에서 바이트 문자열 함수 strlen/strcmp를 잘못 적용하지 않도록 한다.',
    code=r'''#include <stdio.h>
#include <wchar.h>

int main(void)
{
    const wchar_t *word = L"ABC";
    size_t length = wcslen(word);
    int equal = wcscmp(word, L"ABC") == 0;
    printf("units=%zu equal=%d\n", length, equal);
    return 0;
}''',expected='units=3 equal=1\n',
    predict='L"ABC"에는 종료 와이드 문자까지 총 몇 원소가 저장될까? wcslen은 종료 원소를 길이에 포함할까?',
    steps=['`<wchar.h>`는 wchar_t 및 와이드 문자열 함수의 선언을 제공한다. 여기서는 기본 실행 문자 A,B,C만 사용한다.',
           '`L"ABC"`는 네 개의 wchar_t 원소(A,B,C,종료값 0)를 갖는 와이드 문자열 리터럴이다. word는 수정하지 않는 문자열을 가리킨다.',
           '`wcslen(word)`는 종료 원소를 제외한 3을 돌려준다. 이 수는 UTF-8 바이트 수와 같은 개념이 아니다.',
           '`wcscmp`의 반환값이 0이면 두 와이드 문자열이 같다. 정확한 양수/음수의 크기는 가정하지 않는다.',
           'printf는 계산 결과를 눈으로 확인하기 위한 출력이고, 핵심 처리는 와이드 문자열 함수 호출이다.'],
    trace='word[0]=L\'A\', word[1]=L\'B\', word[2]=L\'C\', word[3]=0 → wcslen=3 → wcscmp(...)=0 → equal=1.',
    mistake='와이드 문자열 포인터를 `(char *)`로 바꿔 strlen에 넘겨 글자 수를 구하지 않는다. 객체 표현 바이트에 따라 전혀 다른 결과가 나올 수 있다.',
    misconception='“wchar_t는 모든 운영체제에서 4바이트 UTF-32다”는 틀리다. Windows와 여러 Unix 계열에서 크기와 표현이 다를 수 있다.',
    exercise='문자열을 L"HELLO"로 바꾸고 units=5 equal=1을 출력하라.',
    variant=r'''#include <stdio.h>
#include <wchar.h>

int main(void)
{
    const wchar_t *word = L"HELLO";
    size_t length = wcslen(word);
    int equal = wcscmp(word, L"HELLO") == 0;
    printf("units=%zu equal=%d\n", length, equal);
    return 0;
}''',variant_out='units=5 equal=1\n',reason='L"HELLO"는 다섯 와이드 문자와 종료 원소 하나로 구성된다. 길이에는 종료 원소가 포함되지 않는다.'),
   dict(title='mbstowcs와 wcstombs로 기본 문자 왕복 변환하기',why='파일이나 네트워크에서 읽은 멀티바이트 문자열을 와이드 API에 전달한 뒤 다시 저장할 때 변환 경계를 이해한다.',
    code=r'''#include <locale.h>
#include <stdio.h>
#include <stdlib.h>
#include <wchar.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    char input[] = "CAT";
    wchar_t wide[16];
    char output[16];
    size_t count = mbstowcs(wide, input, 16);
    if (count == (size_t)-1 || count >= 16) {
        return 1;
    }
    size_t bytes = wcstombs(output, wide, 16);
    if (bytes == (size_t)-1 || bytes >= 16) {
        return 1;
    }
    printf("wide_units=%zu bytes=%zu text=%s\n", count, bytes, output);
    return 0;
}''',expected='wide_units=3 bytes=3 text=CAT\n',
    predict='기본 문자 "CAT"을 왕복 변환하면 와이드 원소 수와 출력 바이트 수는 얼마일까? 버퍼가 작으면 종료 문자가 자동으로 붙을까?',
    steps=['`LC_CTYPE`은 문자 인코딩과 분류에 관련된 로캘 범주다. "C"를 선택해 기본 실행 문자의 변환을 재현 가능하게 한다.',
           '`wide[16]`은 wchar_t 원소 16개를 저장하고 `output[16]`은 char 16개를 저장한다. 두 배열은 서로 다른 표현을 담는다.',
           '`mbstowcs`는 입력의 종료 문자 전까지 세 개의 문자를 변환하고 3을 반환한다. 이 예제는 공간이 충분해 종료 문자도 저장된다.',
           '`wcstombs`는 와이드 문자열을 다시 멀티바이트 문자열로 변환한다. `(size_t)-1`이면 오류, 반환값이 버퍼 크기 이상이면 종료 저장을 보장할 수 없으므로 사용을 중단한다.',
           '출력은 왕복 변환을 관찰하는 단계다. 실제 데이터 처리에서는 필요한 길이를 먼저 계산하거나 넉넉한 용량을 검증한다.'],
    trace='input="CAT\\0" → mbstowcs: wide={L\'C\',L\'A\',L\'T\',0} count=3 → wcstombs: output="CAT\\0" bytes=3.',
    mistake='버퍼 길이를 채운 변환 뒤에 종료 문자가 반드시 있다고 가정하고 %s로 출력하면 범위를 벗어날 수 있다. 길이와 종료 여부를 확인한다.',
    misconception='“mbstowcs는 입력이 UTF-8이면 언제나 제대로 변환한다”는 틀리다. 현재 로캘이 그 인코딩을 지원해야 한다.',
    exercise='입력을 "DOGS"로 바꾸어 wide_units=4 bytes=4 text=DOGS를 출력하라.',
    variant=r'''#include <locale.h>
#include <stdio.h>
#include <stdlib.h>
#include <wchar.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    char input[] = "DOGS";
    wchar_t wide[16];
    char output[16];
    size_t count = mbstowcs(wide, input, 16);
    if (count == (size_t)-1 || count >= 16) {
        return 1;
    }
    size_t bytes = wcstombs(output, wide, 16);
    if (bytes == (size_t)-1 || bytes >= 16) {
        return 1;
    }
    printf("wide_units=%zu bytes=%zu text=%s\n", count, bytes, output);
    return 0;
}''',variant_out='wide_units=4 bytes=4 text=DOGS\n',reason='기본 실행 문자 네 개는 "C" 로캘에서 네 개의 와이드 문자로 변환되고, 역변환한 바이트 문자열도 네 문자다.'),
   dict(title='겹치는 와이드 배열을 wmemmove로 안전하게 이동하기',why='문자열 앞에 표시를 끼워 넣거나 구간을 이동할 때 원본·대상 영역이 겹칠 수 있다.',
    code=r'''#include <stdio.h>
#include <wchar.h>

int main(void)
{
    wchar_t word[8] = L"ABCD";
    wmemmove(word + 1, word, 4);
    word[5] = L'\0';
    printf("result_ok=%d units=%zu\n",
           wcscmp(word, L"AABCD") == 0, wcslen(word));
    return 0;
}''',expected='result_ok=1 units=5\n',
    predict='원래 A B C D 0인 배열에서 앞의 네 원소를 한 칸 오른쪽으로 옮기면 어떻게 될까? 원본과 대상이 겹쳐도 될까?',
    steps=['`word[8]`은 지역 와이드 배열이며 초기 문자열은 L"ABCD"다. 실제 원소는 A,B,C,D,0과 나머지 초기화된 0이다.',
           '`wmemmove(word+1,word,4)`는 네 개의 wchar_t 원소를 원소 1~4 위치로 이동한다. 원본 0~3과 대상 1~4가 겹쳐도 wmemmove는 올바르게 처리한다.',
           '`word[5]=L\'\\0\'`는 새 길이 다섯 원소 뒤에 종료 와이드 문자를 확실히 놓는다. 종료 원소를 이동 대상으로 포함하지 않았기 때문에 직접 설정한다.',
           '`wcscmp`와 `wcslen`은 결과를 검증한다. 핵심 로직은 겹침을 허용하는 이동과 종료 설정이다.'],
    trace='시작: [A,B,C,D,0,0,...] → 4원소 이동: [A,A,B,C,D,0,...] → word[5]=0 → L"AABCD" → 길이 5.',
    mistake='겹치는 구간을 wmemcpy로 복사하는 것은 정의되지 않은 동작이므로 wmemmove를 사용한다. 옮길 원소 수와 종료 문자 공간을 함께 계산한다.',
    misconception='“wmemmove의 세 번째 인수는 바이트 수다”는 틀리다. `wmemmove`는 wchar_t 원소 개수를 받는다. `memmove`의 세 번째 인수는 바이트 수다.',
    exercise='L"XYZ"의 첫 세 원소를 한 칸 오른쪽으로 이동해 L"XXYZ"가 되도록 하라.',
    variant=r'''#include <stdio.h>
#include <wchar.h>

int main(void)
{
    wchar_t word[8] = L"XYZ";
    wmemmove(word + 1, word, 3);
    word[4] = L'\0';
    printf("result_ok=%d units=%zu\n",
           wcscmp(word, L"XXYZ") == 0, wcslen(word));
    return 0;
}''',variant_out='result_ok=1 units=4\n',reason='원래 세 문자를 오른쪽으로 옮겨 X,X,Y,Z가 된다. 인수 3은 바이트가 아니라 wchar_t 원소 세 개다.'),
   dict(title='wcstol로 와이드 문자열의 숫자 부분과 접미사 구분하기',why='와이드 문자 API에서 받은 숫자와 단위가 섞인 입력을 분리하고 변환 실패를 검사한다.',
    code=r'''#include <errno.h>
#include <stdio.h>
#include <stdlib.h>
#include <wchar.h>

int main(void)
{
    wchar_t input[] = L"42kg";
    wchar_t *end = NULL;
    errno = 0;
    long value = wcstol(input, &end, 10);
    if (end == input || errno == ERANGE) {
        return 1;
    }
    printf("value=%ld suffix_ok=%d\n",
           value, wcscmp(end, L"kg") == 0);
    return 0;
}''',expected='value=42 suffix_ok=1\n',
    predict='wcstol이 L"42kg"를 읽으면 end는 어느 와이드 문자를 가리킬까? 문자열 전체를 숫자로 취급하면 안 되는 이유는 무엇일까?',
    steps=['`wcstol`은 와이드 문자열의 정수 부분을 long으로 변환한다. 10은 10진수 해석을 지정한다.',
           '`errno=0`은 이번 변환의 오류를 확인하기 위한 준비다. errno가 0이 아닌 모든 경우를 실패라고 단정하지 않고 범위 오류 ERANGE를 검사한다.',
           '`end==input`은 변환할 숫자가 없다는 뜻이다. ERANGE는 결과가 long의 범위를 벗어났음을 나타낸다.',
           '`wcscmp(end,L"kg")`는 남은 접미사가 예상한 단위인지 확인한다. 숫자만 허용한다면 *end==L\'\\0\'을 요구해야 한다.'],
    trace='입력 L"42kg" → 42를 long으로 변환 → end는 L\'k\' 위치 → wcscmp(L"kg",L"kg")=0 → suffix_ok=1.',
    mistake='wcstol의 반환값 0만 보고 실패라고 판단하지 않는다. 입력 "0"은 정상이고, end와 errno를 함께 검사한다.',
    misconception='“와이드 문자열의 숫자 변환은 atoi를 사용해도 된다”는 틀리다. atoi는 char*용이며 wcstol은 wchar_t*용이다.',
    exercise='L"125ms"를 변환하고 value=125 suffix_ok=1을 출력하라.',
    variant=r'''#include <errno.h>
#include <stdio.h>
#include <stdlib.h>
#include <wchar.h>

int main(void)
{
    wchar_t input[] = L"125ms";
    wchar_t *end = NULL;
    errno = 0;
    long value = wcstol(input, &end, 10);
    if (end == input || errno == ERANGE) {
        return 1;
    }
    printf("value=%ld suffix_ok=%d\n",
           value, wcscmp(end, L"ms") == 0);
    return 0;
}''',variant_out='value=125 suffix_ok=1\n',reason='세 자리 숫자를 읽고 end는 접미사 m을 가리킨다. 남은 와이드 문자열을 wcscmp로 비교한다.')
  ],
  guided='부분 완성: 겹치는 와이드 문자 구간을 옮기려면 `wmem____(destination, source, count)`의 빈칸은 무엇인가?',
  guided_answer='`move`. `wmemmove`는 겹치는 영역을 지원하고 count는 wchar_t 원소 수다.',
  debug='1단계(형식): wchar_t*에 strlen을 적용하려는 코드를 wcslen으로 고친다. 2단계(논리): mbstowcs의 반환값을 확인하지 않고 버퍼를 출력하는 코드를 오류·용량 검사로 보강한다. 3단계(경계): 겹치는 구간의 wmemcpy를 wmemmove로 교체하고 널 종료 위치를 검증한다.',
  review='1.49절의 memmove와 1.62절의 LC_CTYPE를 떠올리자. 왜 같은 이동 문제라도 원소 수 단위가 달라지는지, 로캘이 문자 변환에 어떤 영향을 주는지 설명하라.',
  independent_prompt='독립 작성: 와이드 문자열 L"LEVEL"의 길이 5와 L"LEVEL"과의 일치 여부 1을 출력하라. 필요한 헤더와 올바른 서식 지정자를 직접 선택하라.',
  independent_code=r'''#include <stdio.h>
#include <wchar.h>

int main(void)
{
    wchar_t text[] = L"LEVEL";
    size_t count = wcslen(text);
    int equal = wcscmp(text, L"LEVEL") == 0;
    printf("length=%zu equal=%d\n", count, equal);
    return 0;
}''',independent_out='length=5 equal=1\n',independent_reason='널 종료 원소는 길이에 포함하지 않는다. size_t에는 %zu를 사용하고 wcscmp 결과가 0인지 검사한다.'),
 dict(number='1.64',title='와이드 문자 분류와 C11 UTF 코드 단위: <wctype.h>·<uchar.h>',
  need='문자 분류를 위해 `isalpha`에 와이드 문자를 무심코 전달하면 좁은 문자 함수의 인수 범위를 벗어날 수 있다. 또한 UTF-16·UTF-32 코드 단위와 구현 정의 wchar_t를 혼동하면 파일 형식이나 네트워크 프로토콜에서 오류가 발생한다. 표준 C11은 와이드 문자 분류와 char16_t·char32_t 변환 인터페이스를 구별해 제공한다.',
  concepts=[
   '`wint_t`는 와이드 문자 값 또는 WEOF를 나타낼 수 있는 정수형이다. `<wctype.h>`의 `iswalpha`, `iswdigit`, `towupper` 등은 와이드 문자 분류·변환을 수행한다. `isw*` 함수는 참이면 0이 아닌 값을 반환하므로 결과를 1로 정규화해 출력한다.',
   '`wctype("digit")`는 이름 붙은 문자 분류 속성의 기술자를 반환하고 `iswctype`이 이를 사용한다. `wctrans("toupper")`는 변환 기술자를 반환하고 `towctrans`가 적용한다. 기술자를 얻지 못했을 때 0을 검사해야 한다.',
   '`<uchar.h>`는 C11의 `char16_t`·`char32_t` 및 `mbrtoc16`, `c16rtomb`, `mbrtoc32`, `c32rtomb`을 제공한다. char16_t/char32_t는 UTF-16/UTF-32 코드 단위를 표현하기 위한 정수형이다. 변환은 로캘의 멀티바이트 인코딩과 mbstate_t 변환 상태를 사용한다.',
   '`mbstate_t`는 이어지는 변환에서 상태를 보존한다. 0으로 초기화된 상태는 초기 변환 상태를 나타낸다. 변환 함수는 `(size_t)-1`(인코딩 오류), `(size_t)-2`(입력 부족), 경우에 따라 `(size_t)-3`(추가 코드 단위) 등을 반환할 수 있다. ASCII 기본 문자만 쓰는 예제에서 이런 경계 조건을 먼저 분리해 배운다.'
  ],
  cases=[
   dict(title='iswalpha·iswdigit·towupper의 참/거짓 구별하기',why='와이드 문자 입력을 분류하고 대문자로 정규화하는 규칙을 이해한다.',
    code=r'''#include <locale.h>
#include <stdio.h>
#include <wctype.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    int letter = iswalpha(L'k') != 0;
    int number = iswdigit(L'7') != 0;
    wint_t upper = towupper(L'k');
    printf("letter=%d number=%d upper_ok=%d\n",
           letter, number, upper == L'K');
    return 0;
}''',expected='letter=1 number=1 upper_ok=1\n',
    predict='k는 알파벳으로, 7은 숫자로 분류될까? towupper(L\'k\')가 반환하는 것은 char 문자열일까, 와이드 문자 값일까?',
    steps=['`<wctype.h>`는 와이드 문자 분류와 변환 함수를 선언한다. `LC_CTYPE`을 "C"로 고정해 결과가 재현되도록 한다.',
           '`iswalpha(L\'k\')`와 `iswdigit(L\'7\')`는 참이면 0이 아닌 값을 반환한다. `!=0`을 사용해 출력용 0 또는 1로 바꾼다.',
           '`towupper(L\'k\')`는 대응되는 대문자 와이드 문자 값 L\'K\'를 반환한다. 결과를 wint_t 변수에 담는다.',
           '마지막 비교식은 결과가 기대한 와이드 문자와 같은지 확인한다. 출력 자체는 학습 확인용이며 분류·변환 호출이 실제 핵심이다.'],
    trace='C 로캘 선택 → k는 알파벳(참) → 7은 숫자(참) → k를 대문자 K로 변환 → 세 비교가 모두 참.',
    mistake='와이드 문자에 좁은 문자 함수 toupper를 그대로 적용하지 않는다. toupper의 인수는 EOF 또는 unsigned char로 표현 가능한 값이어야 한다.',
    misconception='“iswalpha는 참이면 반드시 정수 1을 반환한다”는 틀리다. 표준은 0이 아닌 값만 보장한다. 0/1 출력이 필요하면 `!=0`을 사용한다.',
    exercise='문자 L\'m\'을 대문자로 변환하고 L\'M\'과 같은지 확인하라. 숫자는 L\'8\'로 바꾸라.',
    variant=r'''#include <locale.h>
#include <stdio.h>
#include <wctype.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    int letter = iswalpha(L'm') != 0;
    int number = iswdigit(L'8') != 0;
    wint_t upper = towupper(L'm');
    printf("letter=%d number=%d upper_ok=%d\n",
           letter, number, upper == L'M');
    return 0;
}''',variant_out='letter=1 number=1 upper_ok=1\n',reason='C 로캘에서 m은 알파벳이고 8은 숫자다. towupper는 m을 M으로 변환한다.'),
   dict(title='wctype·wctrans 기술자로 문자 규칙 선택하기',why='문자 분류 종류와 변환 종류를 문자열 이름으로 선택해야 하는 일반화된 처리 코드에 활용한다.',
    code=r'''#include <locale.h>
#include <stdio.h>
#include <wctype.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    wctype_t category = wctype("digit");
    wctrans_t mapping = wctrans("toupper");
    if (category == 0 || mapping == 0) {
        return 1;
    }
    int digit = iswctype(L'4', category) != 0;
    int upper = towctrans(L'q', mapping) == L'Q';
    printf("digit=%d upper=%d\n", digit, upper);
    return 0;
}''',expected='digit=1 upper=1\n',
    predict='"digit" 기술자를 받은 iswctype는 L\'4\'를 어떻게 판정할까? "toupper" 변환 기술자는 L\'q\'를 어떻게 바꿀까?',
    steps=['`wctype("digit")`는 숫자 분류를 가리키는 기술자이고 `wctrans("toupper")`는 대문자 변환을 가리키는 기술자다.',
           '두 기술자 반환값이 0이면 지원되는 분류/변환을 얻지 못한 것이므로 사용하지 않고 종료한다.',
           '`iswctype`은 지정한 기술자에 따라 L\'4\'가 숫자인지 검사한다. `towctrans`는 지정한 변환을 L\'q\'에 적용한다.',
           '두 결과를 0 또는 1로 정규화해 출력한다. 기술자를 얻는 과정은 함수 종류를 선택하는 준비 단계다.'],
    trace='category="digit", mapping="toupper" 선택 → 4는 숫자 → q는 Q로 변환 → digit=1, upper=1.',
    mistake='wctype가 0을 반환했는데 그대로 iswctype에 넘기면 안 된다. 현재 로캘이 제공하지 않는 속성 이름일 수 있다.',
    misconception='“wctype_t는 문자열 포인터이므로 printf %s로 출력할 수 있다”는 틀리다. 표준이 정한 불투명한 기술자 자료형이며 내부 표현을 가정하면 안 된다.',
    exercise='L\'9\'와 L\'z\'를 사용해 digit=1 upper=1을 확인하라.',
    variant=r'''#include <locale.h>
#include <stdio.h>
#include <wctype.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    wctype_t category = wctype("digit");
    wctrans_t mapping = wctrans("toupper");
    if (category == 0 || mapping == 0) {
        return 1;
    }
    int digit = iswctype(L'9', category) != 0;
    int upper = towctrans(L'z', mapping) == L'Z';
    printf("digit=%d upper=%d\n", digit, upper);
    return 0;
}''',variant_out='digit=1 upper=1\n',reason='9는 digit 분류에 속하고 z를 대문자 변환하면 Z가 된다. 기술자 유효성 검사도 유지한다.'),
   dict(title='mbrtoc32와 c32rtomb로 코드 단위 왕복 변환하기',why='멀티바이트 입력을 char32_t 코드 단위로 변환한 뒤 다시 바이트로 인코딩하는 절차를 분리한다.',
    code=r'''#include <limits.h>
#include <locale.h>
#include <stdio.h>
#include <string.h>
#include <uchar.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    mbstate_t read_state = {0};
    mbstate_t write_state = {0};
    char32_t unit = 0;
    size_t consumed = mbrtoc32(&unit, "R", 1, &read_state);
    if (consumed != 1) {
        return 1;
    }
    char output[MB_LEN_MAX + 1];
    size_t written = c32rtomb(output, unit, &write_state);
    if (written == (size_t)-1 || written > MB_LEN_MAX) {
        return 1;
    }
    output[written] = '\0';
    printf("read=%zu written=%zu same=%d\n",
           consumed, written, strcmp(output, "R") == 0);
    return 0;
}''',expected='read=1 written=1 same=1\n',
    predict='기본 문자 R을 "C" 로캘에서 읽으면 몇 바이트를 소비할까? 변환 상태는 왜 별도로 두 개 만들었을까?',
    steps=['`<uchar.h>`는 char32_t, mbstate_t 및 mbrtoc32/c32rtomb을 제공한다. `<limits.h>`의 MB_LEN_MAX는 멀티바이트 문자의 최대 바이트 수를 나타내는 정수 상수다.',
           '`read_state`와 `write_state`는 각각 0으로 초기화된 독립 변환 상태다. 서로 다른 방향의 변환 상태를 섞지 않는다.',
           '`mbrtoc32(&unit,"R",1,&read_state)`는 기본 문자 R을 char32_t 코드 단위로 변환한다. 성공 시 입력 바이트 하나를 소비해 1을 반환한다.',
           '`c32rtomb`는 코드 단위를 현재 로캘의 멀티바이트 표현으로 변환한다. 오류를 검사한 뒤 output[written]에 종료 문자를 직접 넣는다.',
           'strcmp로 왕복 결과를 검증한다. 출력용 printf는 관찰 장치이고 핵심은 상태를 유지하며 양방향 변환하는 과정이다.'],
    trace='초기 read_state/write_state={0} → R 한 바이트 소비 → unit은 R에 해당하는 코드 단위 → 다시 한 바이트 R 기록 → 종료 문자 추가 → same=1.',
    mistake='변환 함수의 반환값이 `(size_t)-1` 또는 `(size_t)-2`일 수 있다는 사실을 무시하고 그대로 배열 인덱스로 사용하면 안 된다. 이 예제는 소비 바이트가 정확히 1인지 먼저 검사한다.',
    misconception='“char32_t 코드 단위는 언제나 wchar_t와 같은 자료형이다”는 틀리다. char32_t와 wchar_t는 서로 다른 목적의 자료형이고 크기·인코딩이 같다고 보장하지 않는다.',
    exercise='문자 R 대신 S를 변환해 같은 형식으로 read=1 written=1 same=1을 출력하라.',
    variant=r'''#include <limits.h>
#include <locale.h>
#include <stdio.h>
#include <string.h>
#include <uchar.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    mbstate_t read_state = {0};
    mbstate_t write_state = {0};
    char32_t unit = 0;
    size_t consumed = mbrtoc32(&unit, "S", 1, &read_state);
    if (consumed != 1) {
        return 1;
    }
    char output[MB_LEN_MAX + 1];
    size_t written = c32rtomb(output, unit, &write_state);
    if (written == (size_t)-1 || written > MB_LEN_MAX) {
        return 1;
    }
    output[written] = '\0';
    printf("read=%zu written=%zu same=%d\n",
           consumed, written, strcmp(output, "S") == 0);
    return 0;
}''',variant_out='read=1 written=1 same=1\n',reason='기본 실행 문자 S도 "C" 로캘에서 한 바이트로 표현된다. 변환 상태와 종료 처리 규칙은 그대로 적용된다.'),
   dict(title='mbrtoc16과 c16rtomb의 상태 기반 왕복 변환하기',why='UTF-16 코드 단위를 사용하는 자료형과 UTF-32 코드 단위를 사용하는 자료형을 구별한다.',
    code=r'''#include <limits.h>
#include <locale.h>
#include <stdio.h>
#include <string.h>
#include <uchar.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    mbstate_t in_state = {0};
    mbstate_t out_state = {0};
    char16_t unit = 0;
    size_t consumed = mbrtoc16(&unit, "A", 1, &in_state);
    if (consumed != 1) {
        return 1;
    }
    char output[MB_LEN_MAX + 1];
    size_t written = c16rtomb(output, unit, &out_state);
    if (written == (size_t)-1 || written > MB_LEN_MAX) {
        return 1;
    }
    output[written] = '\0';
    printf("read=%zu written=%zu same=%d\n",
           consumed, written, strcmp(output, "A") == 0);
    return 0;
}''',expected='read=1 written=1 same=1\n',
    predict='C 로캘의 기본 문자 A를 UTF-16 코드 단위로 변환했다가 되돌리면 몇 바이트가 소비·기록될까?',
    steps=['`char16_t`는 C11 UTF-16 코드 단위 자료형이다. `wchar_t`와 혼동하지 않는다.',
           '`mbrtoc16`은 멀티바이트 입력을 char16_t 코드 단위로 바꾼다. 기본 문자 A에서는 1바이트를 소비한다.',
           '`c16rtomb`는 char16_t 코드 단위를 멀티바이트 출력으로 변환한다. 입력·출력의 mbstate_t를 따로 유지한다.',
           '변환 오류와 출력 버퍼 범위를 검사한 뒤 종료 문자를 추가한다. 복잡한 문자에서는 한 글자가 여러 UTF-16 코드 단위로 나뉠 수 있어 상태를 이어 처리해야 한다.'],
    trace='초기 상태 두 개 → A 한 바이트 읽음 → char16_t 코드 단위 → A 한 바이트 기록 → 종료 문자 추가 → 원본과 동일.',
    mistake='mbrtoc16이 반환하는 (size_t)-3을 무조건 치명적 오류로 처리하면 여러 코드 단위가 필요한 문자에서 이어지는 결과를 놓칠 수 있다. 이 예제는 기본 문자 한 개만 다룬다.',
    misconception='“char16_t 한 원소는 항상 유니코드 글자 하나다”는 틀리다. UTF-16에서는 하나의 문자에 서로게이트 쌍 두 코드 단위가 필요할 수 있다.',
    exercise='입력을 "B"로 바꾸고 같은 방식으로 read=1 written=1 same=1을 확인하라.',
    variant=r'''#include <limits.h>
#include <locale.h>
#include <stdio.h>
#include <string.h>
#include <uchar.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    mbstate_t in_state = {0};
    mbstate_t out_state = {0};
    char16_t unit = 0;
    size_t consumed = mbrtoc16(&unit, "B", 1, &in_state);
    if (consumed != 1) {
        return 1;
    }
    char output[MB_LEN_MAX + 1];
    size_t written = c16rtomb(output, unit, &out_state);
    if (written == (size_t)-1 || written > MB_LEN_MAX) {
        return 1;
    }
    output[written] = '\0';
    printf("read=%zu written=%zu same=%d\n",
           consumed, written, strcmp(output, "B") == 0);
    return 0;
}''',variant_out='read=1 written=1 same=1\n',reason='B도 C 로캘의 기본 문자이므로 변환과 역변환에서 각각 한 바이트를 처리한다.')
  ],
  guided='부분 완성: C11의 32비트 코드 단위 자료형은 `char____`이고, 멀티바이트 문자열을 코드 단위로 읽는 함수는 `mbrtoc____`이다.',
  guided_answer='두 빈칸 모두 `32`. char32_t는 wchar_t와 별개이고 변환에는 현재 로캘과 mbstate_t 상태가 사용된다.',
  debug='1단계(헤더): iswalpha를 사용하면서 `<ctype.h>`만 포함한 코드를 `<wctype.h>`로 고친다. 2단계(오개념): iswalpha 반환값이 반드시 1이라고 가정하는 비교를 `!=0`으로 고친다. 3단계(경계): mbrtoc32 반환값을 무조건 배열 인덱스로 사용하지 말고 오류/입력 부족을 검사한다.',
  review='1.62절의 LC_CTYPE와 1.63절의 mbstowcs를 떠올려 보자. 로캘 설정이 인코딩 변환에 영향을 주는 이유, wchar_t와 char32_t를 구별해야 하는 이유를 설명하라.',
  independent_prompt='독립 작성: C 로캘에서 L\'B\'는 알파벳인지, L\'5\'는 숫자인지 검사하고 `alpha=1 digit=1`을 출력하라. 참값은 반드시 0/1로 정규화하라.',
  independent_code=r'''#include <locale.h>
#include <stdio.h>
#include <wctype.h>

int main(void)
{
    if (setlocale(LC_CTYPE, "C") == NULL) {
        return 1;
    }
    int alpha = iswalpha(L'B') != 0;
    int digit = iswdigit(L'5') != 0;
    printf("alpha=%d digit=%d\n", alpha, digit);
    return 0;
}''',independent_out='alpha=1 digit=1\n',independent_reason='iswalpha/iswdigit는 0이 아닌 참값을 돌려주므로 `!=0`으로 정규화한다. 분류는 로캘에 의존하므로 먼저 C 로캘을 선택한다.')
]

# Save source files and compile/run with both toolchains before producing a final manuscript.
PROGRAMS=[]
for sec in SECTIONS:
    for idx,c in enumerate(sec['cases'],1):
        for kind,key,out_key in [('example','code','expected'),('variation','variant','variant_out')]:
            name=f"ch{sec['number'].replace('.','_')}_{idx}_{kind}.c"
            source=c[key]+'\n'
            (SOURCES/name).write_text(source,encoding='utf-8')
            PROGRAMS.append({'name':name,'section':sec['number'],'kind':kind,'code':source,'expected':c[out_key]})
    name=f"ch{sec['number'].replace('.','_')}_independent_solution.c"
    source=sec['independent_code']+'\n'
    (SOURCES/name).write_text(source,encoding='utf-8')
    PROGRAMS.append({'name':name,'section':sec['number'],'kind':'independent_solution','code':source,'expected':sec['independent_out']})

checks=[]
for program in PROGRAMS:
    for cc in ('gcc','clang'):
        executable=SOURCES/(program['name']+'.'+cc+'.bin')
        cmd=[cc,*FLAGS,str(SOURCES/program['name']),'-o',str(executable)]
        cp=subprocess.run(cmd,capture_output=True,text=True,timeout=30)
        if cp.returncode != 0:
            raise RuntimeError(f"{program['name']} {cc} compile failed:\n{cp.stderr}")
        rp=subprocess.run([str(executable)],capture_output=True,text=True,timeout=10)
        expected=program['expected']
        ok=(rp.returncode==0 and rp.stdout==expected and rp.stderr=='')
        checks.append({'file':program['name'],'compiler':cc,'flags':FLAGS,'link_options':[],
                       'compile_returncode':cp.returncode,'run_returncode':rp.returncode,
                       'expected_stdout':expected,'actual_stdout':rp.stdout,'stderr':rp.stderr,'passed':ok})
        executable.unlink()
        if not ok:
            raise RuntimeError(f"{program['name']} {cc} mismatch: expected={expected!r}, actual={rp.stdout!r}, stderr={rp.stderr!r}")

# DOCX layout
D=Document(); section=D.sections[0]
section.page_width=Cm(21.0); section.page_height=Cm(29.7)
section.top_margin=Cm(2.05); section.bottom_margin=Cm(1.95)
section.left_margin=Cm(2.3); section.right_margin=Cm(2.3)
section.header_distance=Cm(0.9); section.footer_distance=Cm(0.9)
styles=D.styles
normal=styles['Normal']; normal.font.name='Noto Sans CJK KR'; normal.font.size=Pt(9.4)
normal.font.color.rgb=RGBColor(34,47,53)
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.18
for name,sz,color,before,after in [('Title',18,RGBColor(16,90,94),0,13),('Heading 1',14,RGBColor(15,103,108),16,9),('Heading 2',11.2,RGBColor(28,88,100),11,6)]:
    s=styles[name]; s.font.name='Noto Sans CJK KR'; s.font.size=Pt(sz); s.font.bold=True; s.font.color.rgb=color
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True
for name,sz,color,bold in [('Study Label',9.2,RGBColor(13,113,113),True),('C Code',8.3,RGBColor(32,42,47),False),('Expected Output',8.8,RGBColor(22,74,91),False)]:
    if name not in styles:styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s=styles[name]; s.font.name='DejaVu Sans Mono' if name in ('C Code','Expected Output') else 'Noto Sans CJK KR'
    s.font.size=Pt(sz); s.font.color.rgb=color; s.font.bold=bold
    s.paragraph_format.space_after=Pt(3)
    s.paragraph_format.line_spacing=1.05 if name in ('C Code','Expected Output') else 1.15
    if name=='Study Label':s.paragraph_format.keep_with_next=True

# Header/footer
header=section.header.paragraphs[0];header.text='C언어 종합교과서  |  이해 중심 정식 개정판  |  제1장'
header.style='Normal';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor(98,116,119)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('독립 집필 · ISO C11 예제 검증   |   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

codes_in_doc=[]
def p(text,style=None):
    q=D.add_paragraph(style=style or 'Normal');q.add_run(text)
    q.paragraph_format.widow_control=True
    return q
def label(text):p(text,'Study Label')
def code(text):
    codes_in_doc.append(text+'\n')
    lines=text.splitlines()
    for i,line in enumerate(lines):
        q=D.add_paragraph(style='C Code')
        q.paragraph_format.left_indent=Cm(0.45)
        q.paragraph_format.keep_together=True
        q.add_run(line if line else ' ')
def heading(text,level):
    q=D.add_heading(text,level=level)
    q.paragraph_format.keep_with_next=True
    return q

D.add_heading('한국어 C언어 종합교과서 — 이해 중심 정식 개정판',0)
p('제1장 1.62~1.64절 | 로캘·와이드 문자·C11 UTF 코드 단위')
p('1.61절 다음에 이어지는 새 본문. 기존 거부 원고 및 과거 제1~11장 초안은 새 정식 본문으로 재사용하지 않았다.')
p('학습 순서: 필요성 → 새 개념 → 전체 코드 → 실행 전 예측 → 문장별 해설 → 값·실행 추적 → 실제 출력 → 오류·오개념 → 바로 이어지는 변형 실습과 정답 → 단계적 독립 작성·누적 복습.')
p('ISO C11 기준으로 작성했으며 GCC·Clang 엄격 경고 옵션과 실제 실행을 검증한다. 전역 객체는 사용하지 않는다. 기본 실행 문자와 C 로캘을 사용해 구현별 로캘·문자 인코딩 차이를 통제했다.')
p('기준 PDF 두 권 및 범위 MD의 실제 파일 접근과 일부 페이지를 확인했다. 원서의 모든 예제 유형·실습 목표에 대한 항목별 전수 대조는 아직 진행 중이며, 미확인 항목을 완료로 간주하지 않는다.')

for sec in SECTIONS:
    heading(sec['number']+'절  '+sec['title'],1)
    label('필요성 · 없으면 생기는 문제')
    p(sec['need'])
    label('처음 등장하는 개념과 기술적 경계')
    for i,concept in enumerate(sec['concepts'],1):p(f'{i}. {concept}')
    for i,c in enumerate(sec['cases'],1):
        heading(f"{sec['number']}-{i}  {c['title']}",2)
        label('왜 필요한가 · 언제 사용하는가')
        p(c['why'])
        label('전체 C11 코드 — 독립 실행 가능')
        code(c['code'])
        label('실행하기 전에 먼저 예측')
        p(c['predict'])
        label('문장별 설명 · 이 줄을 왜 쓰는가')
        for j,s in enumerate(c['steps'],1):p(f'{j}) {s}')
        label('값과 실행 흐름 추적')
        p(c['trace'])
        label('실제 실행 결과 (GCC·Clang 검증)')
        p(c['expected'].rstrip('\n'),'Expected Output')
        label('흔한 실수와 단계별 오류 수정')
        p(c['mistake'])
        label('오개념 진단 · 그럴듯하지만 틀린 생각')
        p(c['misconception'])
        label('예제를 바꾸며 실력 다지기 — 지금 바로 풀기')
        p(c['exercise'])
        label('변형 문제 정답: 전체 코드 (바로 옆에 배치)')
        code(c['variant'])
        label('변형 정답의 실제 출력과 이유')
        p(c['variant_out'].rstrip('\n'),'Expected Output')
        p(c['reason'])
    heading(sec['number']+'절  단계적으로 도움을 줄이는 연습과 누적 복습',2)
    label('1단계 · 부분 완성/빈칸 채우기')
    p(sec['guided']);p('정답과 근거: '+sec['guided_answer'])
    label('2단계 · 오류의 종류를 구별하며 고치기')
    p(sec['debug'])
    label('3단계 · 힌트 없이 독립적으로 코드 작성')
    p(sec['independent_prompt'])
    label('독립 작성 정답 코드와 해설')
    code(sec['independent_code'])
    p(sec['independent_out'].rstrip('\n'),'Expected Output')
    p(sec['independent_reason'])
    label('시간을 두고 다시 떠올리는 누적 인출 복습')
    p(sec['review'])

heading('표준과 플랫폼 구별 · 원서 전수 대조 상태',1)
p('setlocale·localeconv·strcoll·strxfrm·strtod, wchar_t와 와이드 문자열 함수, wctype 분류 및 C11 uchar.h의 변환 인터페이스는 표준 C 범위다. 실제 로캘 이름과 지원 문자 인코딩은 구현마다 달라질 수 있으므로 예제는 C 로캘만 사용한다.')
p('wchar_t의 크기와 표현은 구현에 따라 다르고, UTF-8이 C 로캘의 기본 인코딩이라는 보장은 없다. 표준 C와 POSIX locale_t/newlocale 및 Windows 전용 로캘 API를 구별한다. GCC/Clang 컴파일 옵션은 ISO C 문법이 아니라 검증 도구의 설정이다.')
p('K&R 원본 스캔 288쪽, OCR 보조본 238쪽, Express 스캔 750쪽과 전체 범위 MD에 접근했다. 이번 회차에는 Express 목차 일부와 K&R 표지/판권 및 OCR 목차를 확인했다. 각 장의 LAB·Exercise·Programming·Mini Project·Advanced Topic별 전수 대조는 미완료이며 추적표에서 보류한다.')
p('다음 절에서는 원서 누락 목표를 우선 확인하면서 미수록 표준 라이브러리·언어 기능을 계속 보강한다. 총 PART 수나 완료 시점은 제한하지 않는다.')

out=ROOT/DOCX_NAME
D.save(out)
# Verify every full code block in manuscript corresponds exactly to a compiled source.
expected_codes=[x['code'] for x in PROGRAMS]
if Counter(codes_in_doc)!=Counter(expected_codes):
    raise RuntimeError(f'DOCX code mismatch: {len(codes_in_doc)} vs {len(expected_codes)}')
assert len(codes_in_doc)==len(PROGRAMS)
assert all(not re.search(r'(?m)^\s*(?:int|double|float|char|size_t|ptrdiff_t)\s+\w+\s*=',x['code'].split('int main')[0]) for x in PROGRAMS)

with zipfile.ZipFile(ROOT/(STEM+'_sources.zip'),'w',zipfile.ZIP_DEFLATED) as z:
    for x in PROGRAMS:z.write(SOURCES/x['name'],arcname=x['name'])

(ROOT/(STEM+'_c11_verification.json')).write_text(json.dumps({
 'run':STEM,'timestamp':NOW,'standard':'ISO C11','flags':FLAGS,'compilers':['gcc','clang'],
 'link_options':[],'source_count':len(PROGRAMS),'compile_run_count':len(checks),
 'all_passed':all(x['passed'] for x in checks),'cases':checks},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(ROOT/(STEM+'_code_match.json')).write_text(json.dumps({
 'docx_code_blocks':len(codes_in_doc),'compiled_source_count':len(PROGRAMS),'matched':True,
 'per_file_sha256':{x['name']:hashlib.sha256(x['code'].encode()).hexdigest() for x in PROGRAMS}},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

