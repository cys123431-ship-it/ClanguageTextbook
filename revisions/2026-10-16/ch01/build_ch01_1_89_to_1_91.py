#!/usr/bin/env python3
"""독창적 C11 원고·소스·검증·진행 추적을 재현 가능하게 생성."""
from pathlib import Path
from datetime import datetime, timezone, timedelta
import subprocess, json, hashlib, zipfile, re, os, shutil, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from content_run29 import SECTIONS
ROOT=Path(__file__).resolve().parent
SRC=ROOT/'sources'
SRC.mkdir(exist_ok=True)
NOW=datetime.now(timezone(timedelta(hours=9))).isoformat(timespec='seconds')
FLAGS=['-std=c11','-Wall','-Wextra','-Wpedantic','-Werror','-O0']
STEM='official_ch01_1_89-1_91'
DOCX=ROOT/'C언어_종합교과서_이해중심_정식개정판_제1장_1.89-1.91.docx'

def digest(b): return hashlib.sha256(b).hexdigest()

ENTRIES=[]
for s in SECTIONS:
 for i,e in enumerate(s['examples'],1):
  for typ,code,expected in [('example',e['code'],e['expected']),('variation',e['variation_code'],e['variation_expected'])]:
   ENTRIES.append({'section':s['id'],'example':e['key'],'kind':typ,'path':f"ch01_{s['id'].replace('.','_')}_{i:02d}_{e['key']}_{typ}.c",'code':code,'expected':expected})
 ind=s['independent']
 ENTRIES.append({'section':s['id'],'example':'independent','kind':'independent','path':f"ch01_{s['id'].replace('.','_')}_independent.c",'code':ind['code'],'expected':ind['expected']})

VERIFIED=[]
for e in ENTRIES:
 src=SRC/e['path'];src.write_text(e['code'],encoding='utf-8')
 v={'section':e['section'],'example':e['example'],'kind':e['kind'],'file':e['path'],'sha256':digest(e['code'].encode()),'expected_stdout':e['expected'],'results':{}}
 for cc in ['gcc','clang']:
  exe=SRC/(src.stem+'_'+cc)
  cmd=[cc,*FLAGS,str(src),'-o',str(exe)]
  comp=subprocess.run(cmd,capture_output=True,text=True,timeout=30)
  if comp.returncode or comp.stdout or comp.stderr:raise RuntimeError(f'compile failed: {e["path"]} {cc} {comp.returncode} {comp.stderr}')
  run=subprocess.run([str(exe)],capture_output=True,timeout=10)
  if run.returncode!=0 or run.stdout!=e['expected'].encode('utf-8') or run.stderr:
   raise RuntimeError(f'output mismatch {e["path"]} {cc}: {run.returncode} {run.stdout!r} {run.stderr!r} vs {e["expected"]!r}')
  v['results'][cc]={'compiled':True,'executed':True,'exit_code':0,'stdout_matched':True,'stderr_empty':True,'command':' '.join(cmd)}
  exe.unlink()
 VERIFIED.append(v)

# test diagnostic tooling on safe examples (nonstandard tool, distinct from ISO C11 verification)
SUPPLEMENTARY=[]
for cc in ['gcc','clang']:
 src=SRC/'ch01_1_91_02_safe_buffer_example.c'
 exe=SRC/f'.sanitizer_{cc}'
 cmd=[cc,*FLAGS,'-g','-fsanitize=address,undefined','-fno-omit-frame-pointer',str(src),'-o',str(exe)]
 comp=subprocess.run(cmd,capture_output=True,text=True,timeout=30)
 if comp.returncode:
  SUPPLEMENTARY.append({'compiler':cc,'available':False,'reason':comp.stderr[:800]})
 else:
  # Running ASan binaries in restricted sandboxes can be unstable; record observed result, do not claim if failed.
  run=subprocess.run([str(exe)],capture_output=True,timeout=15)
  SUPPLEMENTARY.append({'compiler':cc,'available':True,'exit_code':run.returncode,'stdout_matched':run.stdout==b'\xeb\xb3\xb5\xec\x82\xac \xec\x8b\xa4\xed\x8c\xa8\n','stderr':run.stderr.decode('utf-8','replace')[:800],'passed':run.returncode==0 and run.stdout=='복사 실패\n'.encode() and not run.stderr})
 if exe.exists():exe.unlink()

# Editable manuscript
D=Document();sec=D.sections[0]
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.74);sec.left_margin=Inches(.86);sec.right_margin=Inches(.84)
sec.header_distance=Inches(.32);sec.footer_distance=Inches(.35)
normal=D.styles['Normal'];normal.font.name='Malgun Gothic';normal.font.size=Pt(9.6);normal.font.color.rgb=RGBColor(32,49,57)
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.17
for name,size,before,after in [('Title',20,0,10),('Heading 1',15,13,7),('Heading 2',12.1,10,5),('Heading 3',10.6,8,4)]:
 st=D.styles[name];st.font.name='Malgun Gothic';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,113,117)
 st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
code_style=D.styles.add_style('Editable C Code',WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name='Consolas';code_style.font.size=Pt(7.55);code_style.font.color.rgb=RGBColor(16,45,54)
code_style.paragraph_format.space_after=Pt(1);code_style.paragraph_format.line_spacing=1.06
note_style=D.styles.add_style('Note',WD_STYLE_TYPE.PARAGRAPH)
note_style.font.name='Malgun Gothic';note_style.font.size=Pt(8.5);note_style.font.color.rgb=RGBColor(80,88,95)
header=sec.header.paragraphs[0];header.text='C언어 종합교과서  |  이해 중심 정식 개정판  |  제1장';header.style='Note';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('제1장 · 1.89–1.91   |   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def para(text,style=None):
 p=D.add_paragraph(style=style);p.add_run(text);return p

def label(head,text):
 p=D.add_paragraph();p.paragraph_format.keep_together=True
 r=p.add_run(head+'  ');r.bold=True;r.font.color.rgb=RGBColor(0,109,115);p.add_run(text)
 return p

def code(code):
 p=D.add_paragraph(style='Editable C Code');p.paragraph_format.left_indent=Inches(.12)
 # Avoid page break inside small code blocks; long code may move to next page.
 p.paragraph_format.keep_together=True
 p.add_run(code.rstrip('\n').replace('    ','\u00a0\u00a0\u00a0\u00a0'))
 return p

def output(value):
 p=D.add_paragraph(style='Note');p.paragraph_format.keep_with_next=True;p.add_run('실제 출력 — GCC·Clang에서 일치 확인').bold=True
 p=D.add_paragraph(style='Editable C Code');p.add_run(value.rstrip('\n'))

D.add_heading('C언어 종합교과서',0)
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('이해 중심 정식 개정판 · 제1장 1.89~1.91절');r.bold=True;r.font.size=Pt(13.5);r.font.color.rgb=RGBColor(0,113,117)
para('학습 흐름: 진단 메시지의 의미 → 실행 경로·변수 상태 관찰 → 경계값과 회귀 테스트. 각 절은 완성 예제, 인접 변형 과제, 부분 완성 연습, 독립 과제 순서로 도움을 줄인다.')
para('이번 원고의 모든 완성·변형·독립 정답은 C11 프로그램이다. GCC·Clang 엄격 옵션으로 각각 컴파일하고 표준 출력 바이트·표준 오류·종료 상태를 비교한다. GDB·LLDB·Sanitizer는 ISO C11 외부의 도구로 별도 표시한다.')
para('이전 제1~11장 거부 초안은 사용하지 않았다. 이전 1.1~1.88절의 집필 완료는 Library 보고를 계승하며 이번 회차에 전체를 재검증하지 않았다. Express·K&R 원서의 모든 문제 유형별 전수 대조는 미완료다.')
for s in SECTIONS:
 D.add_heading(s['id']+'  '+s['title'],1)
 for t in s['intro']:para(t)
 D.add_heading('처음 등장하는 개념을 정확히 구별하기',2)
 for k,t in s['concepts']:label(k,t)
 D.add_heading('복잡한 일을 작은 목표로 나누기',2)
 for t in s['plan']:para(t)
 for i,e in enumerate(s['examples'],1):
  D.add_heading(f"{s['id']}.{i}  {e['title']}",2)
  label('필요성·실제 사용 사례',e['why'])
  D.add_heading('완성 예제 — 편집 가능한 전체 C 코드',3);code(e['code'])
  label('해설 전에 먼저 예측하기',e['predict'])
  D.add_heading('문장별 해설: 이 문장을 왜 쓰는가',3)
  for j,x in enumerate(e['walk'],1):label('해설 '+str(j),x)
  D.add_heading('실행 중 값과 제어 흐름의 실제 변화',3)
  for x in e['trace']:label('추적',x)
  output(e['expected'])
  label('오개념 진단·교정',e['misconception'])
  label('단계별 디버깅',e['debug'])
  D.add_heading('예제를 바꾸며 실력 다지기 — 바로 옆에서 해결',3)
  label('변형 문제',e['task'])
  if i==2:label('도움 줄이기','먼저 변경할 자료와 실행 흐름을 스스로 적고, 아래 정답과 비교한다. 실행 결과만 외우지 말고 조건문이 참·거짓이 되는 이유를 확인하자.')
  D.add_heading('인접 정답 — 변경 후 전체 코드',3);code(e['variation_code']);output(e['variation_expected'])
  label('정답의 근거',e['reason'])
 D.add_heading('부분 완성·빈칸 문제',2);label('문제',s['scaffold'][0]);label('정답과 이유',s['scaffold'][1])
 D.add_heading('도움 없이 독립적으로 작성하기',2)
 para(s['independent']['task']+' 구현하기 전에 입력·반환 규약·경계 사례를 적고, 아래 정답을 보지 않은 상태로 코드를 작성하자.')
 D.add_heading('독립 작성 정답 — 전체 코드',3);code(s['independent']['code']);output(s['independent']['expected']);label('정답 근거',s['independent']['reason'])
 D.add_heading('간격 복습·인출 연습',2);label('이전 절을 다시 떠올리기',s['review'][0]);label('정답과 근거',s['review'][1])
 if s['id']=='1.89':
  D.add_heading('컴파일러 진단과 실제 오류를 구별하는 연습',2)
  para('GCC·Clang에서는 -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 옵션으로 컴파일한다. -Werror는 활성화된 경고를 오류로 취급하지만 모든 논리 오류를 발견하지는 못한다. -Wpedantic도 컴파일러가 모든 미정의 동작을 증명한다는 뜻은 아니다.')
  label('의도적으로 잘못된 문장','int number = "12"; 는 정수 객체에 문자열 리터럴이 가리키는 주소를 대입하려는 잘못된 코드다. 수정하려면 정수 리터럴 12를 사용하거나 문자열을 strtol로 변환해야 한다. 이 반례는 검증된 완성 프로그램에 포함하지 않는다.')
  label('논리 오류 반례','if (value = 0) 는 비교가 아니라 대입이다. GCC·Clang이 경고할 수 있지만, 경고가 없다고 해서 의도한 비교가 되는 것은 아니다. 비교 목적이라면 value == 0을 사용한다.')
 if s['id']=='1.90':
  D.add_heading('실제 디버거에서 따라 하는 순서 — 플랫폼 도구',2)
  code('gcc -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 -g main.c -o app\ngdb ./app\n(gdb) break positive_sum\n(gdb) run\n(gdb) next\n(gdb) print sum\n(gdb) backtrace\n(gdb) continue')
  para('위 명령은 1.90.1의 코드를 main.c로 저장한 GCC/GDB 예시다. gdb는 C 표준의 일부가 아니다. break는 중단점, run은 시작, next는 다음 소스 줄, print는 값 관찰, backtrace는 호출 경로, continue는 재개를 뜻한다. LLDB는 breakpoint set, run, next, frame variable, bt 등 대응 명령을 제공한다. Visual Studio는 F9 중단점, F5 시작, F10 단계 실행, F11 함수 진입, 호출 스택 창을 사용한다.')
  label('실행 환경 주의','디버그 정보·최적화·프로그램 입출력에 따라 한 줄에 대응하는 실제 기계어 명령과 변수 표시가 다를 수 있다. 디버거가 보여 주는 주소나 메모리 배치는 ISO C가 고정하지 않는다.')
 if s['id']=='1.91':
  D.add_heading('Sanitizer·정적 분석을 사용할 때의 범위',2)
  code('gcc -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 -g -fsanitize=address,undefined -fno-omit-frame-pointer main.c -o app\nclang -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 -g -fsanitize=address,undefined -fno-omit-frame-pointer main.c -o app')
  para('이 옵션들은 GCC·Clang의 확장이며 도구 체인에 런타임 라이브러리가 필요하다. AddressSanitizer는 메모리 접근 오류의 일부를, UndefinedBehaviorSanitizer는 여러 미정의 동작을 탐지한다. 모든 오류를 검출하는 증명기는 아니다. Windows/MSVC의 AddressSanitizer는 /fsanitize=address와 지원 환경을 별도로 확인해야 한다.')
  label('위험 코드의 읽기 전용 반례','char tiny[3]; strcpy(tiny, "ABCDE"); 는 배열에 필요한 공간보다 많은 바이트를 쓰므로 정의되지 않은 동작이다. 이 코드는 실행하지 않는다. 수정은 복사 전 용량 검사 또는 충분한 크기의 배열 사용이다.')
  label('추가 점검','컴파일러의 정적 분석은 GCC -fanalyzer, Clang --analyze, MSVC /analyze 등으로 수행할 수 있다. Valgrind는 별도 설치가 필요한 플랫폼 도구다. 이 회차에서는 해당 도구가 모두 존재하거나 모든 분석을 통과했다고 주장하지 않는다.')
D.add_heading('이번 원고의 범위·검증 한계',1)
para('C11 완성·변형·독립 정답 21개를 각각 GCC·Clang으로 컴파일·실행해 실제 출력과 종료 상태를 비교한다. 경고를 일부러 발생시키는 문장과 메모리 오류 반례는 교육용 설명만 제공하며 정상 프로그램의 검증 건수에 포함하지 않는다.')
para('기준 원서 두 권의 전체 LAB·Exercise·Programming·Mini Project·Advanced Topic 및 부록 학습목표를 한 항목씩 대조하는 작업은 별도 추적표에서 미완료로 유지한다. 원서의 원문 문장·문제·그림을 복제하지 않고 학습목표를 독립적으로 구현했다.')
D.save(DOCX)

# DOCX <-> source exact match (paragraph strings may contain NBSP replacing spaces)
D2=Document(DOCX)
blocks=[p.text.replace('\u00a0',' ') for p in D2.paragraphs if p.style.name=='Editable C Code']
matched=[]
for e in ENTRIES:
 normalized=e['code'].rstrip('\n')
 if normalized not in blocks:
  raise RuntimeError('DOCX source mismatch: '+e['path'])
 matched.append(e['path'])

ZIP=ROOT/(STEM+'_sources.zip')
with zipfile.ZipFile(ZIP,'w',zipfile.ZIP_DEFLATED) as z:
 for e in ENTRIES:z.write(SRC/e['path'],e['path'])
 z.writestr('README.txt','21개 독립 C11 프로그램. gcc/clang -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 source.c -o app\nGDB, Sanitizer는 별도 도구이며 ISO C 표준이 아닙니다.\n')

verify={'timestamp':NOW,'sections':[s['id'] for s in SECTIONS],'compilers':['gcc','clang'],'flags':FLAGS,'program_count':len(ENTRIES),'compile_run_count':len(ENTRIES)*2,'all_passed':True,'programs':VERIFIED,'supplementary_sanitizer_probe':SUPPLEMENTARY,'note':'Sanitizer probes are additional tool checks, not part of C11 compile/run count.'}
(ROOT/(STEM+'_c11_verification.json')).write_text(json.dumps(verify,ensure_ascii=False,indent=2)+'\n')
(ROOT/(STEM+'_code_match.json')).write_text(json.dumps({'timestamp':NOW,'source_program_count':len(matched),'all_matched':True,'filenames':matched},ensure_ascii=False,indent=2)+'\n')
(ROOT/(STEM+'_pedagogy_qa.json')).write_text(json.dumps({'timestamp':NOW,'sections':[s['id'] for s in SECTIONS],'gradual_independence':True,'spaced_retrieval':True,'prediction_before_trace':True,'debugging_stages':True,'subgoal_decomposition':True,'misconception_correction':True,'variation_adjacent_to_example':True,'independent_answers':3,'first_appearance_notes':True,'prior_sections_not_reaudited':True},ensure_ascii=False,indent=2)+'\n')
(ROOT/(STEM+'_source_inspection.json')).write_text(json.dumps({'timestamp':NOW,'scope_md':'C언어_통합교과서_전체수록범위_확정안.md','scope_unit':'30 디버깅·검증·품질 관리','express_pdf':'Not found in current Library root/search; previously reported 750 pages; this run cannot claim fresh PDF inspection','kr_original_pdf':'Not found in current Library root/search; previously reported 288 pages; this run cannot claim fresh PDF inspection','kr_ocr_pdf':'Not found in current Library root/search; previously reported 238 pages; this run cannot claim fresh PDF inspection','original_item_audit_complete':False,'copyright':'Original explanatory text and independently created examples, no original exercises reproduced'},ensure_ascii=False,indent=2)+'\n')
(ROOT/'official_rewrite_progress_1_1-1_91.json').write_text(json.dumps({'timestamp':NOW,'reported_previous_completed':'1.1-1.88 (Library reports, not independently rechecked)','this_run_verified':['1.89','1.90','1.91'],'last_completed_section':'1.91','next_section':'1.92','overall_percentage':None,'original_objective_audit_complete':False,'github_delivery':'pending verification','note':'Earlier rejected chapters 1-11 not counted'},ensure_ascii=False,indent=2)+'\n')
(ROOT/'official_rewrite_objective_tracking_1_1-1_91.json').write_text(json.dumps({'timestamp':NOW,'original_sources_item_by_item_audit':'incomplete','prior_reported_1_1_to_1_88':'not independently reaudited','new_sections':[{'section':s['id'],'topic':s['title'],'editorial_unit':'30','new_original_examples':len(s['examples']),'variations':len(s['examples']),'independent_solutions':1,'c11_compiler_verification':'passed','original_express_item_ids':'unverified','original_kr_item_ids':'unverified'} for s in SECTIONS],'unmapped_original_items':'Unknown until both originals fully inspected; do not mark complete'},ensure_ascii=False,indent=2)+'\n')
# manifest excludes itself to avoid recursive hash problem
artifacts=[DOCX,ZIP,ROOT/'content_run29.py',ROOT/'build_ch01_1_89_to_1_91.py']+sorted(ROOT.glob('*.json'))
manifest={'timestamp':NOW,'range':'1.89-1.91','verified_new_sections':['1.89','1.90','1.91'],'new_programs':21,'compiler_runs':42,'original_item_audit_complete':False,'artifacts':[{'name':p.name,'sha256':digest(p.read_bytes()),'bytes':p.stat().st_size} for p in artifacts]}
(ROOT/'official_rewrite_manifest_1_89-1_91.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'docx':str(DOCX),'sources_zip':str(ZIP),'programs':len(ENTRIES),'compiler_runs':len(ENTRIES)*2,'sanitizer_probes':SUPPLEMENTARY,'timestamp':NOW},ensure_ascii=False,indent=2))