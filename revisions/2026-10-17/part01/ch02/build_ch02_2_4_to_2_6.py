from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
ROOT=Path('/mnt/data/ctextbook/run_20261017_0312'); SRC=ROOT/'sources'; SRC.mkdir(exist_ok=True,parents=True)
BASE=Path('/mnt/data/ctextbook/source/C언어_종합교과서_이해중심_정식개정판_PART01_제02장_2.1-2.3.docx')
OUT=ROOT/'C언어_종합교과서_PART01_제02장_2.1-2.6_누적1차초안.docx'
d=Document(BASE)
# 기존 본문 2.1~2.3 전체 재활용. 마지막 결산 부분은 2.4~2.6 추가 후 새 결산으로 교체.
for p in list(d.paragraphs):
    if p.style.name=='Heading 1' and p.text in ['제2장 2.1~2.3 학습 완료 점검','원서 학습목표 대응 및 현재 검증 한계']:
        start=next(i for i,q in enumerate(d.paragraphs) if q._element is p._element)
        for q in list(d.paragraphs)[start:]:
            q._element.getparent().remove(q._element)
        break

codes={
'2_4_01_status.c':r'''#include <stdio.h>

int main(void)
{
    int seats = 12;
    int reserved = 5;
    int available = seats - reserved;

    printf("available=%d\n", available);
    return 0;
}
''',
'2_4_02_price.h':r'''#ifndef PRICE_H
#define PRICE_H

int final_price(int unit_price, int quantity);

#endif
''',
'2_4_02_price.c':r'''#include "2_4_02_price.h"

int final_price(int unit_price, int quantity)
{
    return unit_price * quantity;
}
''',
'2_4_02_main.c':r'''#include <stdio.h>
#include "2_4_02_price.h"

int main(void)
{
    int total = final_price(1200, 3);
    printf("total=%d\n", total);
    return 0;
}
''',
'2_5_01_label_v1.c':r'''#include <stdio.h>

int main(void)
{
    printf("revision=1\n");
    return 0;
}
''',
'2_5_02_label_v2.c':r'''#include <stdio.h>

int main(void)
{
    printf("revision=2\n");
    return 0;
}
''',
'2_6_01_portable.c':r'''#include <stdio.h>

int main(void)
{
    int first = 8;
    int second = 3;
    int total = first + second;
    printf("total=%d\n", total);
    return 0;
}
''',
'2_6_02_warning_fixed.c':r'''#include <stdio.h>

int main(void)
{
    int expected = 10;
    int actual = 7;
    printf("difference=%d\n", expected - actual);
    return 0;
}
''',
}
for k,v in codes.items(): (SRC/k).write_text(v,encoding='utf-8')

# Helper functions

def h(text,level=2): return d.add_heading(text,level)
def p(text,style=None): return d.add_paragraph(text,style=style)
def note(text): return p(text,'Note')
def code(name,shown=True):
    note('전체 C 코드 · '+name)
    for line in codes[name].rstrip('\n').split('\n'):p(line,'Code')
def console(text):
    p('실제 출력' if '\n' not in text else '실제 출력 / 명령 기록')
    for line in text.rstrip('\n').split('\n'):p(line,'Code')
def steps(*items):
    for n,s in enumerate(items,1):p(f'{n}단계. {s}')

def paragraph_bundle(items):
    for item in items:p(item)

h('2.4 Visual Studio Community에서 빈 프로젝트를 실제 C 프로그램으로 만들기',1)
h('왜 배우는가 — .c 파일을 저장했다고 프로젝트가 자동으로 알아보지는 않는다',2)
p('앞의 2.1절에서는 소스 코드를 컴파일해야 실행 파일이 생긴다고 배웠고, 2.2절에서는 main.c·기능을 담당하는 .c·선언을 적은 .h가 협력하는 모습을 보았다. 이제 그 파일을 Windows 11의 Visual Studio Community에서 실제로 빌드하는 방법을 익힌다. 컴퓨터에 파일이 존재한다는 사실과, 현재 열어 둔 프로젝트가 그 파일을 빌드 대상으로 알고 있다는 사실은 다르기 때문이다.')
p('책상 위에 계산 종이 세 장을 올려놓았다고 해서 제출 서류철 안에 세 장이 자동으로 들어가지는 않는다. 실제 종이를 서류철에 끼워 넣어야 제출 목록에 포함된다. Visual Studio의 프로젝트도 비슷하다. 디스크의 폴더는 실제 파일이 있는 장소이고, 프로젝트는 어떤 파일을 묶어서 컴파일할지 관리하는 목록이다. 다만 프로젝트는 단순한 서류철보다 복잡해서 컴파일 옵션과 연결 대상도 기억한다. 이 비유는 파일을 자동으로 복사한다는 뜻이 아니다.')
p('실무에서 이 구분을 모르면 USB로 main.c, function.c, header.h를 복사했는데 솔루션 탐색기에 main.c만 보이는 상황에서 “파일이 사라졌다”고 오해하기 쉽다. 파일이 진짜 사라졌는지, 프로젝트에만 등록되지 않았는지를 먼저 구분해야 한다.')
h('Windows 11에서 새 프로젝트를 만드는 순서',2)
steps('Visual Studio Community를 실행한다. 필요한 워크로드는 일반적으로 “C++를 사용한 데스크톱 개발”이다. Visual Studio의 C 지원은 별도 C 전용 프로젝트 템플릿이 아니라 C/C++ 도구 집합을 통해 제공된다.',
'시작 화면에서 “새 프로젝트 만들기”를 선택하고 “빈 프로젝트(Empty Project)”를 찾는다. 버전과 언어 필터에 따라 표시 위치가 다를 수 있으므로 이름으로 검색한다.',
'프로젝트 이름을 예를 들어 LibraryPractice로 정하고 저장 위치를 확인한다. “솔루션”은 프로젝트를 묶는 상위 작업 단위이며, “프로젝트”는 컴파일할 파일과 설정을 관리한다.',
'솔루션 탐색기에서 프로젝트 이름을 오른쪽 클릭해 “추가 → 새 항목”으로 이동한다. 소스 파일을 만들고 확장자를 .c로 저장한다. 기본값이 .cpp로 제안되면 C 실습 파일은 .c로 명시한다.',
'파일 내용을 입력하고 Ctrl+S로 저장한다. 이어 “빌드 → 솔루션 빌드” 또는 Ctrl+Shift+B로 빌드한다. 컴파일에 실패하면 “오류 목록”과 “출력” 창에서 첫 원인을 읽는다.',
'실행은 일반적으로 Ctrl+F5(디버깅하지 않고 시작) 또는 F5(디버깅 시작)를 사용한다. 단축키가 다르면 메뉴를 이용한다. 실행 결과를 확인할 때는 실행한 프로젝트가 맞는지도 점검한다.')
code('2_4_01_status.c')
h('이 코드의 모든 줄을 쉬운 말로 읽기',3)
p('`#include <stdio.h>`는 화면에 글자를 쓰는 printf의 선언을 컴파일러가 알 수 있도록 표준 헤더를 포함한다. `int main(void)`는 프로그램의 시작 함수이고, 중괄호는 그 함수에 속한 명령을 묶는다. 여기까지는 2.1절에서 이미 사용한 구조다.')
p('`int seats = 12;`는 정수를 담는 그릇을 준비해 seats라는 이름을 붙이고 12를 처음 넣는다. `int reserved = 5;`는 예약된 자리 5개를 따로 기억한다. `int available = seats - reserved;`는 전체 12개에서 예약 5개를 빼 7을 계산한 뒤 available에 저장한다. 이때 `=`는 수학의 등호가 아니라 선언하면서 처음 값을 정하는 초기화 문법의 일부다.')
p('`printf("available=%d\\n", available);`는 문장 속 `%d` 자리에 available의 정수 값 7을 끼워 넣어 출력하고, `\\n`으로 줄을 바꾼다. `return 0;`은 main의 정상 종료 상태를 반환한다. 화면에 숫자 0을 추가로 인쇄하는 명령은 아니다.')
console('available=7')
h('값과 파일의 흐름을 따라가기',3)
p('입력 단계: 소스 파일에 seats=12와 reserved=5를 적는다. 저장 단계: 실제 파일 2_4_01_status.c가 디스크에 기록된다. 빌드 단계: 프로젝트가 이 .c 파일을 대상으로 컴파일과 링크를 수행한다. 실행 단계: 만들어진 실행 파일이 시작되면 available=12-5=7이 되고 한 줄이 출력된다. “저장”과 “빌드” 사이를 건너뛰면 이전 실행 파일이 남아 있을 수 있다.')
h('바로 옆 변형 실습 — 예측하고 수정하기',2)
p('문제: 전체 좌석 12개 중 예약된 좌석을 9개로 바꾸면 결과가 어떻게 될까? `reserved`를 9로 수정하기 전에 종이에 먼저 계산해 보자. 정답은 available=3이다. `seats`나 printf를 고칠 필요 없이 예약된 개수만 바꾸면 된다. 변경 후 반드시 다시 빌드해야 새 값이 반영된다.')
note('오개념 경고  “솔루션 탐색기에서 파일이 안 보이니 디스크에서도 삭제되었다”는 결론은 성급하다. Windows 파일 탐색기에서 실제 경로를 확인하고, Visual Studio 프로젝트의 파일 포함 여부를 따로 확인한다.')
h('이미 가지고 있는 세 파일을 빈 프로젝트에 연결하기',2)
p('앞 절에서 만든 main.c·기능 담당 .c·선언 담당 .h를 다른 컴퓨터나 USB에서 가져왔다고 하자. 새 파일을 다시 타이핑할 필요는 없다. Windows 파일 탐색기에서 세 파일이 실제로 있는 폴더를 확인한 뒤 프로젝트 이름을 오른쪽 클릭하고 “추가 → 기존 항목”으로 각 파일을 선택하면 된다. 프로젝트가 파일을 참조하는 방식은 설정과 작업 방법에 따라 다를 수 있으므로, 원본이 어디에 있는지와 프로젝트 폴더에 복사했는지를 별도로 확인한다.')
p('예제로 가격 계산 기능을 나누어 보자. 2.2절에서 이미 배운 것처럼 헤더는 함수가 어떻게 불리는지를 알려 주고, 소스 파일은 실제 계산을 정의하며, main 파일은 그 기능을 사용한다. 다음 세 파일은 **같은 폴더**에 둔다. 실제 프로젝트에는 두 .c 파일을 빌드 대상으로 넣고, .h 파일은 편집·관리용으로 함께 포함한다.')
for name in ['2_4_02_price.h','2_4_02_price.c','2_4_02_main.c']:code(name)
h('세 파일을 말로 연결해 보기',3)
p('헤더의 `int final_price(int unit_price, int quantity);`는 “정수 두 개를 받아 정수 하나를 돌려주는 final_price라는 함수가 있다”는 약속이다. `#ifndef`·`#define`·`#endif`는 같은 헤더 내용이 중복 포함될 때 선언이 여러 번 처리되는 문제를 예방하는 보호 장치다. 세부 전처리 규칙은 제14장에서 다시 배운다.')
p('price.c의 `return unit_price * quantity;`는 단가에 개수를 곱한 결과를 호출한 곳에 돌려준다. main.c의 `final_price(1200, 3)`는 1200과 3을 보내 계산을 요청하고, 돌려받은 3600을 total에 저장한다. 마지막 printf는 그 값을 보여 준다. `#include "..."`의 큰따옴표 표기는 이 프로젝트에서 작성한 헤더를 찾도록 요청한다. 헤더 이름과 실제 파일 이름은 글자 하나까지 일치해야 한다.')
console('total=3600')
p('빌드가 “final_price를 찾을 수 없다”는 링크 오류로 실패하면 main.c의 함수 호출만 고칠 일이 아니다. 먼저 price.c가 프로젝트에 포함되어 있는지 확인한다. 헤더 파일만 추가해도 함수의 실제 몸체가 들어 있는 price.c가 빠지면 링크할 정의가 없다. 이 오류를 직접 한 번 경험한 뒤, price.c를 다시 포함하고 빌드가 성공하는지 확인하자.')
note('변형 실습  단가 1200을 1500으로 바꾸고 수량은 3으로 유지한다. 예상 출력은 total=4500이다. 정답: main.c의 호출을 final_price(1500, 3)으로 바꾸면 된다. 헤더나 price.c의 계산식을 불필요하게 수정하지 않는다.')
h('파일이 안 보일 때 점검 순서',2)
steps('Windows 파일 탐색기에서 실제 .c와 .h 파일의 존재와 확장자를 확인한다. 확장자가 숨겨져 있으면 보기 설정에서 표시한다.',
'Visual Studio의 솔루션 탐색기에서 해당 프로젝트를 선택했는지 확인한다. 폴더만 열어 둔 상태와 솔루션을 연 상태는 다르다.',
'프로젝트에 없으면 “추가 → 기존 항목”으로 등록한다. 필요하다면 솔루션 탐색기의 “모든 파일 표시”를 이용해 파일 시스템과 프로젝트 표시의 차이를 살핀다.',
'프로젝트에 main 함수를 가진 .c가 둘 이상 들어가면 진입점이 중복되어 링크 오류가 날 수 있다. 예제별로 프로젝트를 나누거나 해당 빌드에서 제외할 파일을 정한다.',
'정상 빌드한 뒤 실행 결과가 이전 값이라면 실행 대상 프로젝트와 실행 파일 경로를 확인한다.')

h('2.5 저장 폴더, 작업 폴더, 실행 파일을 구별하기',1)
h('왜 배우는가 — 수정한 코드를 실행했는데 옛날 숫자가 나온다',2)
p('2.4절에서 프로젝트에 파일을 등록했다. 그런데 컴파일을 다시 했다고 생각했는데도 화면에 옛날 숫자가 나오는 경우가 있다. 가장 흔한 원인 중 하나는 수정한 소스 파일과 실제 빌드 대상이 다르거나, 새 실행 파일 대신 예전에 만들어 둔 실행 파일을 시작했기 때문이다. 이 절에서는 “내가 편집한 파일”, “컴파일러가 읽은 파일”, “실제로 실행한 파일”을 서로 구별한다.')
p('시험 답안지를 두 장 복사했다고 생각해 보자. 원본에 2라고 고쳐 적어도, 이미 인쇄한 사본에는 여전히 1이 적혀 있다. 실행 파일도 비슷하다. 소스 코드를 고쳤다고 이미 만들어 둔 실행 파일 내부의 명령이 자동으로 바뀌지는 않는다. 단, 소스와 실행 파일의 관계는 단순 종이 복사보다 복잡하고 컴파일·링크 과정을 거친다는 차이가 있다.')
h('눈으로 확인하는 두 버전 실험',2)
p('아래 두 파일은 같은 프로그램을 서로 다른 시점에 저장했다고 가정한 실험용 코드다. 실제로는 하나의 파일을 수정하면서 매번 다른 이름의 실행 파일을 만들 수도 있다. 두 소스 파일을 같은 프로젝트에 동시에 넣으면 main 함수가 두 개가 되므로 **각각 별도로 컴파일**한다.')
code('2_5_01_label_v1.c');console('revision=1')
code('2_5_02_label_v2.c');console('revision=2')
p('`printf("revision=1\\n");`의 큰따옴표 안 글자는 그대로 출력된다. `revision=1`은 프로그램이 만든 버전 표지일 뿐 컴파일러가 자동으로 붙인 빌드 번호가 아니다. 두 번째 파일에서는 문자열의 1을 2로 바꾸었으므로 출력도 2로 바뀐다.')
h('같은 컴퓨터에서 일부러 다른 실행 파일 만들기',2)
p('GCC를 사용하는 터미널이라면 다음 두 명령을 각각 실행한다. 각 명령 끝의 `-o` 다음 이름은 만들어질 실행 파일의 이름이다. Windows에서 GCC를 쓰면 확장자 .exe가 붙을 수 있으며 실행 방식은 사용 중인 셸에 따라 달라진다.')
for line in ['gcc -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 2_5_01_label_v1.c -o old_program','gcc -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 2_5_02_label_v2.c -o new_program','./old_program   # Linux/macOS에서 실행하는 예','./new_program   # Linux/macOS에서 실행하는 예']:p(line,'Code')
p('첫 번째 실행 파일을 실행하면 revision=1, 두 번째를 실행하면 revision=2가 나온다. 이것은 “새 소스가 맞느냐”와 “어느 실행 파일을 실행했느냐”가 별개의 질문이라는 것을 보여 준다. Windows PowerShell에서는 일반적으로 `./` 대신 `.\\old_program.exe`처럼 실행할 수 있다. 실제 생성된 파일 이름과 폴더를 먼저 확인한다.')
h('세 가지 경로를 구별하기',2)
p('① 소스 파일 경로: 편집기가 저장한 2_5_02_label_v2.c의 실제 위치다. ② 출력 파일 경로: 컴파일·링크 결과인 실행 파일이 생성된 위치다. ③ 현재 작업 폴더(working directory): 실행한 프로그램이 상대 경로를 해석할 때 기준으로 삼는 폴더다. 세 경로는 같을 수도 있고 서로 다를 수도 있다.')
p('예를 들어 Visual Studio 프로젝트 폴더에 main.c가 있고, 빌드 결과가 x64\\Debug 아래에 만들어졌다고 하자. 이것은 흔히 볼 수 있는 배치의 예시일 뿐, 실제 경로는 플랫폼·구성·프로젝트 속성에 따라 달라진다. 프로젝트를 Debug에서 Release로 바꾸면 다른 출력 폴더를 사용하도록 설정되어 있을 수도 있다. “파일이 없어진 것 같다”는 느낌만으로 추정하지 말고 출력 창의 실제 경로를 확인한다.')
p('작업 폴더는 이후 파일 입출력 단원에서 특히 중요해진다. 예를 들어 프로그램이 `data.txt`를 열라고 할 때, 별도 경로가 없는 이름은 현재 작업 폴더를 기준으로 해석되는 경우가 일반적이다. 아직 파일 입출력 문법을 배우지 않았으므로 지금은 “상대 경로에는 기준 폴더가 필요하다”는 개념만 익히자.')
h('한글 출력이 깨지는 이유도 이 구분에서 시작한다',2)
p('소스 파일을 UTF-8로 저장하는 일과, 실행 중 한글이 화면에 표시되는 일은 같은 문제가 아니다. 소스 인코딩은 파일 안의 문자를 컴파일러가 어떻게 해석하는지에 영향을 준다. 실행 중 문자열을 어떤 바이트로 표현하고 터미널이 어떻게 해석하는지는 실행 문자 집합과 터미널 설정에 영향을 받는다. 따라서 “한글이 깨졌으니 C 문법이 틀렸다”고 단정하지 않는다.')
p('우선 이번 장의 검증 예제는 영문·숫자 출력으로 통일한다. Windows 콘솔에서 한글이 깨진다면 파일 저장 인코딩, Visual Studio의 소스/실행 문자 집합 설정, 사용하는 콘솔의 문자 인코딩과 글꼴을 구분해 점검한다. 무조건 특정 숫자의 코드 페이지로 바꾸라는 식의 단일 해결책은 다른 환경에서 오히려 문제를 만들 수 있다.')
h('예제 바로 옆 변형 실습과 정답',2)
p('문제 A: v2 소스 파일의 문자열을 revision=3으로 바꾸고 저장만 했다. new_program을 다시 컴파일하지 않고 실행하면 무엇이 나올까? 정답: 기존 new_program이 수정 전의 실행 파일이라면 revision=2가 나온다. 소스 파일 저장과 실행 파일 갱신은 별개다.')
p('문제 B: 새 파일로 빌드했는데도 revision=1이 나온다. 가장 먼저 확인할 것은? 정답: 빌드에 사용한 소스 경로, 빌드 성공 여부, 실행한 프로그램의 실제 경로를 차례로 확인한다. 컴파일러를 재설치하는 것은 첫 번째 조치가 아니다.')
note('오개념 경고  “작업 폴더는 언제나 .c 파일이 있는 폴더다”는 틀리다. IDE나 터미널의 실행 설정에 따라 달라질 수 있다. 상대 경로와 절대 경로의 자세한 사용법은 파일 단원에서 배운다.')

h('2.6 GCC·Clang·MSVC 빌드 명령을 비교하고 경고 읽기',1)
h('왜 배우는가 — 같은 C 코드인데 컴퓨터마다 명령이 다르다',2)
p('앞에서 GCC와 Clang 명령으로 프로그램을 빌드했고, Visual Studio Community에서는 메뉴를 사용했다. 그렇다면 C언어가 세 종류라는 뜻일까? 그렇지 않다. C는 언어이고, GCC·Clang·MSVC는 소스 코드를 번역하는 서로 다른 컴파일러 계열이다. 같은 C11 프로그램이라도 도구가 받아들이는 명령행 옵션과 기본 설정은 다를 수 있다.')
p('같은 목적지에 가더라도 지하철과 버스의 승차 방법은 다르다. C 코드가 목적지라면 컴파일러의 옵션은 각 교통수단의 사용법에 가깝다. 다만 서로 다른 컴파일러가 항상 같은 동작을 보장한다는 뜻은 아니다. C 표준이 정의하지 않은 동작이나 구현별 차이를 가진 코드는 결과가 달라질 수 있다.')
h('첫 번째 비교 실험 — 동일한 C11 코드',2)
code('2_6_01_portable.c')
p('`first`에는 정수 8, `second`에는 3을 처음 넣는다. `total`은 두 변수를 더한 11을 저장한다. `%d` 자리에 total의 값 11이 들어가므로 화면에 total=11이 출력된다. `int`의 크기가 모든 컴퓨터에서 반드시 4바이트라고 가정하지 않았으며, 여기서 사용한 작은 정수의 덧셈은 일반적인 C 구현에서 표현 범위 안에 있다.')
console('total=11')
h('GCC와 Clang에서 빌드하기',2)
for line in ['gcc -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 2_6_01_portable.c -o portable_gcc','clang -std=c11 -Wall -Wextra -Wpedantic -Werror -O0 2_6_01_portable.c -o portable_clang']:p(line,'Code')
p('명령 첫 단어는 사용할 컴파일러 이름이다. `-std=c11`은 C11 모드를 요청한다. `-Wall`과 `-Wextra`는 여러 종류의 경고를 켠다. `-Wpedantic`은 선택한 언어 표준을 기준으로 추가 진단을 요청한다. `-Werror`는 활성화된 경고를 오류로 취급하게 한다. `-O0`은 최적화를 낮춰 입문 실습의 추적을 단순하게 한다. `-o portable_gcc`는 결과 파일 이름을 지정한다. 이 옵션들은 모두 C 코드가 아니라 GCC/Clang에 전달하는 도구 명령이다.')
p('“경고가 하나도 없다”는 사실이 “프로그램이 반드시 올바르다”는 증거는 아니다. 컴파일러는 사람이 의도한 계산까지 알 수 없으므로, 올바르게 번역되었어도 잘못된 합계를 출력하는 프로그램이 있을 수 있다. 그래서 컴파일 성공과 출력 비교를 둘 다 수행한다.')
h('MSVC에서는 왜 다른 옵션을 쓰는가',2)
p('Visual Studio의 “x64 Native Tools Command Prompt for VS”처럼 개발 환경이 설정된 명령 프롬프트에서 cl 명령을 사용할 수 있다. 다음은 MSVC의 한 가지 명령 예시다. `cl`이 일반 명령 프롬프트에서 인식되지 않는다면 C 소스 오류보다 먼저 개발자 명령 프롬프트 사용 여부를 확인한다.')
p('cl /nologo /W4 /WX /Od /std:c11 2_6_01_portable.c /Fe:portable_msvc.exe','Code')
p('`/W4`는 높은 경고 수준을, `/WX`는 경고를 오류로 취급하도록 요청한다. `/Od`는 최적화를 끄도록 요청한다. `/std:c11`은 해당 MSVC 도구 집합에서 C11 언어 모드를 요청한다. `/Fe:`는 실행 파일 이름을 지정한다. GCC의 `-std=c11`을 그대로 cl에 입력하면 안 된다. 또한 MSVC 버전과 설치된 C 라이브러리에 따라 C11 기능 지원 범위가 다르므로, 이 명령이 모든 구버전에서 통한다고 단정하지 않는다.')
h('경고를 이해하고 고치는 연습',2)
p('다음은 일부러 잘못 작성한 코드 조각이다. `unused`라는 변수를 선언하고 99를 넣지만 프로그램의 나머지 부분에서 사용하지 않는다. 컴파일러와 경고 설정에 따라 사용되지 않은 변수에 관한 경고가 발생할 수 있다. 이 조각은 교육용으로 보여 주는 것이며, 모든 컴파일러에서 똑같은 문구가 나온다고 주장하지 않는다.')
for line in ['int main(void)', '{','    int unused = 99;','    return 0;','}']:p(line,'Code')
p('왜 경고가 도움이 될까? 실제 프로그램에서 사용하지 않는 변수는 계산하다가 빼먹은 부분의 흔적일 수 있다. 그렇다고 경고를 없애기 위해 아무 의미 없는 출력을 추가하는 것이 항상 좋은 해결은 아니다. 변수 자체가 불필요하다면 삭제하고, 필요하다면 실제 계산에 사용해야 한다. 다음은 값을 의미 있게 사용하는 완전한 프로그램이다.')
code('2_6_02_warning_fixed.c')
p('`expected`는 예상 수량 10, `actual`은 실제 수량 7을 저장한다. `expected - actual`은 차이 3을 계산하고, printf가 그 결과를 출력한다. 따라서 두 변수를 모두 실제 계산에 사용한다. 변수를 단순히 선언해 놓고 방치하지 않는다는 점이 핵심이다.')
console('difference=3')
h('비교표 대신 말로 정리하는 컴파일러 차이',2)
p('GCC·Clang은 이 장의 실습에서 `-std=c11 -Wall -Wextra -Wpedantic -Werror -O0`을 공통적으로 사용할 수 있다. MSVC는 대체로 `/std:c11 /W4 /WX /Od`와 같은 형식을 사용한다. 두 명령은 경고를 강화하고 최적화를 낮추려는 목적은 비슷하지만, 옵션의 범위가 완전히 일치하지 않는다. 동일한 경고 개수나 동일한 진단 문구를 기대하지 말고 각각의 도구에서 실제 출력된 진단을 읽어야 한다.')
p('Debug/Release는 “정답/오답” 구분이 아니다. Debug 설정은 일반적으로 디버깅을 돕는 정보를 포함하고 최적화를 줄이는 방향으로 구성되고, Release 설정은 배포·성능을 고려한 설정을 사용하는 경우가 많다. 그러나 실제 옵션은 프로젝트 속성에 따라 달라지므로, 설정 이름만 보고 컴파일러 동작을 확정하지 않는다.')
h('바로 옆 변형 실습 — 빌드 명령과 계산을 동시에 바꾸기',2)
p('문제: 2_6_01_portable.c의 second를 3에서 6으로 바꾸고 GCC·Clang에서 각각 새로 빌드한다. 예상 출력은? 정답은 total=14다. first=8은 그대로이고 second=6이므로 total=8+6이다. 결과가 11로 나온다면 계산식부터 의심하지 말고 새로 빌드한 실행 파일을 실행했는지 2.5절 절차로 확인한다.')
p('독립 실습: 정수형 변수 box=15, used=4를 선언하고 remaining=11을 출력하는 C11 프로그램을 처음부터 작성한다. GCC와 Clang에서 각각 빌드하고 결과를 비교한다. 정답의 핵심 문장은 `int remaining = box - used;`와 `printf("remaining=%d\\n", remaining);`이다. 나머지 main 함수, 헤더, 반환문은 2.1절 구조를 참고해 스스로 완성한다.')

h('제2장 2.1~2.6 누적 학습 점검',1)
p('이제 한 파일을 컴파일하는 것(2.1), 여러 파일을 연결하는 것(2.2), 오류 종류를 나누는 것(2.3), Visual Studio 프로젝트에 파일을 등록하는 것(2.4), 편집·빌드·실행 경로를 구별하는 것(2.5), 컴파일러마다 다른 옵션을 사용하는 것(2.6)을 하나의 작업 흐름으로 연결할 수 있어야 한다.')
steps('빈 프로젝트에 .c 파일을 추가해 실행하고, 프로젝트에 등록된 파일과 디스크에 있는 파일의 차이를 설명한다.',
'헤더·함수 정의·main 파일을 각각 어떤 목적으로 사용하는지 말하고, 구현 .c 파일을 빼면 왜 링크 오류가 나는지 설명한다.',
'소스에서 revision=2로 고쳤는데 revision=1이 출력되는 상황의 점검 순서를 설명한다.',
'GCC/Clang의 -std=c11과 MSVC의 /std:c11을 혼동하지 않고 각각의 도구에서 사용한다.',
'컴파일 경고가 없더라도 예상 출력과 실제 출력을 비교한다.')
h('원서 학습목표 대응과 1차 초안 상태',1)
p('이 누적 원고는 Express의 개발 환경·첫 C 프로그램·Visual Studio 프로젝트·오류 수정 학습목표와 K&R의 첫 프로그램·프로그램 구성·여러 파일과 선언 관련 개념을 바탕으로 독립적인 설명과 실습을 구성했다. Express의 개별 LAB/Exercise/Programming/Mini Project/Advanced Topic 및 K&R의 모든 개별 Exercise를 이번 원고만으로 전수 대조했다고 주장하지 않는다.')
p('2.1~2.3절은 이전 원고를 그대로 보존해 재활용했고, 2.4~2.6절은 이번 회차에 새로 작성했다. 제2장의 최종 장 종료 판정은 원서 목표 감사 및 사용자의 1차 초안 검토 이후 수행한다. 다음 회차에서는 새 장을 계속 채우거나 기존 임시 절을 실제 내용에 맞게 정식 장으로 이관한다.')

# Update cover to reflect cumulative status
for pp in d.paragraphs[:12]:
    if pp.text.startswith('2.1~2.3절 |'):
        pp.text='2.1~2.6절 | 기존 2.1~2.3절 원고 재활용 + 신규 2.4~2.6절 초안 · 2026년 10월 17일'
    if pp.text.startswith('이 장은 제1장'):
        pp.text='이 장은 제1장 「프로그래밍과 알고리즘의 첫 원리」에서 경험한 출력·변수·조건을 실제 개발 도구로 실행하는 과정을 학습한다. 자료형·포인터·동적 메모리·POSIX 시스템 호출의 전문 내용은 별도 장에서 다룬다. 새 2.4~2.6절에서는 프로젝트·파일 경로·컴파일러 옵션을 직접 확인한다.'
    if pp.text.startswith('독자는 각 절에서'):
        pp.text='독자는 필요성 → 이미 배운 내용과 연결 → 일상 비유와 한계 → 코드와 한 줄 해설 → 실제 추적과 출력 → 예제 바로 옆 변형 실습의 흐름으로 학습한다. C 코드는 편집 가능한 텍스트다. 검증한 코드와 검증하지 못한 환경별 UI 동작을 구분해 표시한다.'
# Rebuild front learning targets to include new sections
for pp in d.paragraphs[:16]:
    if pp.text.startswith('2.3 문법 오류'):
        from docx.text.paragraph import Paragraph
        anchor=pp._element
        for t in ['2.4 Visual Studio에서 .c/.h 파일을 프로젝트에 등록하고 빌드한다.',
                  '2.5 소스·실행 파일·작업 폴더를 구별해 이전 실행 결과가 나오는 문제를 진단한다.',
                  '2.6 GCC·Clang·MSVC 옵션 차이와 경고의 의미를 설명한다.']:
            new_elm=OxmlElement('w:p');anchor.addnext(new_elm);Paragraph(new_elm,pp._parent).text=t;anchor=new_elm
# 문서에는 Markdown 굵게 표시 기호를 남기지 않는다.
for para in d.paragraphs:
    if '**' in para.text:
        para.text=para.text.replace('**','')
# Metadata
props=d.core_properties;props.title='C언어 종합교과서 제2장 2.1~2.6 누적 1차 초안';props.subject='기존 원고 재활용 및 2.4~2.6 신규 집필'
d.save(OUT)
print(json.dumps({'output':str(OUT),'bytes':OUT.stat().st_size,'source_files':list(codes)},ensure_ascii=False,indent=2))